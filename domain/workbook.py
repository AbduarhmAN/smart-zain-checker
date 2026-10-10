"""Domain Workbook Engine for Smart Zain Checker.
Dedicated to Excel file parsing, header schema analysis, and row extraction.
Pure domain logic: completely decoupled from UI and browser processes.
"""
from __future__ import annotations

import re
import threading
from pathlib import Path
from typing import Any, Optional
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from domain.models import Customer
from domain.money import parse_money_to_halalas


class CalamineCell:
    __slots__ = ("value",)

    def __init__(self, value: Any):
        self.value = value

    def __str__(self) -> str:
        return str(self.value) if self.value is not None else ""


class CalamineSheetAdapter:
    """High-performance in-memory adapter that mimics openpyxl Worksheet for read operations."""

    def __init__(self, data: list[list[Any]], title: str = "Sheet1"):
        self.data = data
        self.title = title
        self._max_col = max((len(r) for r in data), default=0)
        self._max_row = len(data)

    @property
    def max_column(self) -> int:
        return self._max_col

    @property
    def max_row(self) -> int:
        return self._max_row

    def cell(self, row: int, column: int) -> CalamineCell:
        r_idx = row - 1
        c_idx = column - 1
        if 0 <= r_idx < len(self.data) and 0 <= c_idx < len(self.data[r_idx]):
            return CalamineCell(self.data[r_idx][c_idx])
        return CalamineCell(None)

    def iter_rows(
        self,
        min_row: int = 1,
        max_row: Optional[int] = None,
        min_col: int = 1,
        max_col: Optional[int] = None,
        values_only: bool = True,
    ):
        start_r = max(0, min_row - 1)
        end_r = min(len(self.data), max_row) if max_row is not None else len(self.data)
        start_c = max(0, min_col - 1)

        for r_idx in range(start_r, end_r):
            row = self.data[r_idx]
            end_c = min(len(row), max_col) if max_col is not None else len(row)
            if start_c >= len(row):
                row_slice = ()
            else:
                row_slice = row[start_c:end_c]

            if values_only:
                yield tuple(row_slice)
            else:
                yield tuple(CalamineCell(v) for v in row_slice)


class FastWorkbook:
    """Fast in-memory workbook wrapper backed by Rust python-calamine with lazy worksheet loading."""

    def __init__(self, calamine_wb: Any):
        self._calamine_wb = calamine_wb
        self.sheetnames = list(calamine_wb.sheet_names)
        self._sheets: dict[int, CalamineSheetAdapter] = {}
        self._lock = threading.Lock()

    @property
    def worksheets(self) -> list[CalamineSheetAdapter]:
        return [self.get_sheet(i) for i in range(len(self.sheetnames))]

    def get_sheet(self, index: int) -> CalamineSheetAdapter:
        with self._lock:
            if index not in self._sheets:
                title = self.sheetnames[index] if 0 <= index < len(self.sheetnames) else f"Sheet{index+1}"
                raw_sheet = self._calamine_wb.get_sheet_by_index(index)
                # Retain physical 1:1 row coordinates with Excel
                data = raw_sheet.to_python(skip_empty_area=False)
                # Fast trim of trailing phantom rows (e.g. 1M ghost rows created by Excel formatting or overflow)
                # Trimming from the end never alters 1-based coordinates of any preceding data rows.
                last_idx = len(data)
                while last_idx > 0 and not any(data[last_idx - 1]):
                    last_idx -= 1
                if last_idx < len(data):
                    data = data[:last_idx]
                self._sheets[index] = CalamineSheetAdapter(data, title=title)
            return self._sheets[index]

    def __getitem__(self, name_or_idx: int | str) -> CalamineSheetAdapter:
        if isinstance(name_or_idx, int):
            return self.get_sheet(name_or_idx)
        idx = self.sheetnames.index(name_or_idx)
        return self.get_sheet(idx)

    def close(self) -> None:
        pass


def get_workbook_sheet_names(workbook_path: Path | str) -> list[str]:
    """Retrieves sheet names instantly via Calamine Rust engine with openpyxl fallback."""
    path_str = str(workbook_path)
    try:
        import python_calamine
        wb = python_calamine.CalamineWorkbook.from_path(path_str)
        return list(wb.sheet_names)
    except Exception:
        from openpyxl import load_workbook
        wb = load_workbook(path_str, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()


_FAST_WB_CACHE: dict[tuple[str, float], FastWorkbook] = {}
_CACHE_LOCK = threading.Lock()


def clear_fast_workbook_cache() -> None:
    with _CACHE_LOCK:
        _FAST_WB_CACHE.clear()


def load_fast_workbook(workbook_path: Path | str) -> Any:
    """Opens an Excel file with ultra-fast Calamine engine, falling back to openpyxl.
    Maintains a thread-safe, mtime-sensitive in-memory cache of parsed FastWorkbook instances
    to prevent repetitive 16MB file unzipping across consecutive API calls (inspect, validate, start).
    """
    path_obj = Path(workbook_path).resolve()
    path_str = str(path_obj)
    try:
        mtime = path_obj.stat().st_mtime
    except Exception:
        mtime = 0.0

    cache_key = (path_str, mtime)
    with _CACHE_LOCK:
        cached = _FAST_WB_CACHE.get(cache_key)
        if cached is not None:
            return cached

    try:
        import python_calamine
        cal_wb = python_calamine.CalamineWorkbook.from_path(path_str)
        wb = FastWorkbook(cal_wb)
        with _CACHE_LOCK:
            if len(_FAST_WB_CACHE) >= 8:
                _FAST_WB_CACHE.clear()
            _FAST_WB_CACHE[cache_key] = wb
        return wb
    except Exception:
        from openpyxl import load_workbook
        return load_workbook(path_str, read_only=True, data_only=True)



HEADER_ROW_NUMBER = 1

AMOUNT_EXCLUSIONS = (
    "نوع", "رقم", "عمر", "تاريخ", "كود", "ايام", "أيام", "هاتف", "جوال",
    "مفضل", "حاوية", "جهة", "فرع", "مشرف", "مستخدم", "سسس", "يي", "تصنيف",
    "ملاحظات", "متابعة"
)


def is_empty_input_row(values, mapped_columns=()) -> bool:
    """Ignore empty rows and helper links with no account/contract identifier.

    Keep incomplete customer rows: any mapped value, name, note, amount (including
    zero), or populated identifier makes a row meaningful and subject to validation.
    """
    def present(value):
        return value is not None and bool(str(value).strip())

    if any(present(values[col-1]) for col in mapped_columns if col and 0 < col <= len(values)):
        return False
    populated = [value for value in values if present(value)]
    if not populated:
        return True

    def empty_helper_link(value):
        if not isinstance(value, str):
            return False
        try:
            parsed = urlsplit(value.strip())
            host = (parsed.hostname or '').lower()
            if parsed.scheme not in {'http', 'https'} or not (host == 'zain.app' or host.endswith('.zain.sa') or host.endswith('.zain.com')):
                return False
            if parsed.path.rstrip('/').split('/')[-1] not in {'contract-payment', 'quickpay', 'quick-pay'}:
                return False
            query = parse_qs(parsed.query, keep_blank_values=True)
            keys = [key for key in ('contract', 'account') if key in query]
            return bool(keys) and all(not item.strip() for key in keys for item in query[key])
        except ValueError:
            return False

    return all(empty_helper_link(value) for value in populated)


def normalize_header_text(val: Any) -> str:
    """Normalizes Arabic/English header text for fuzzy synonym matching."""
    if val is None:
        return ""
    text = str(val).strip().lower()
    text = re.sub(r"[أإآ]", "ا", text)
    text = re.sub(r"ة", "ه", text)
    text = re.sub(r"ى", "ي", text)
    text = re.sub(r"[_\-–—/\\.]+", " ", text)
    return " ".join(text.split())


def parse_excel_column(col_str: str) -> Optional[int]:
    """Converts column letter ('A', 'AW') or 1-based index number to integer index."""
    if not col_str:
        return None
    s = str(col_str).strip()
    if s.isdigit():
        val = int(s)
        return val if val > 0 else None
    
    # Column letters (e.g. 'L', 'AW')
    res = 0
    for char in s.upper():
        if "A" <= char <= "Z":
            res = res * 26 + (ord(char) - ord("A") + 1)
        else:
            return None
    return res if res > 0 else None


def parse_account_and_service(text: Any) -> tuple[str, str]:
    """Extracts (contract/account, service) numbers from freeform text or compound cells.
    Handles formats like:
      - 'حساب: 1006659013 | خدمة: 2006695911'
      - 'عقد: 1010624047 | جوال: 0599662119'
      - '1001485949 / 2008084328'
      - '1010624047' or '2008084328'
    """
    if not text:
        return "", ""
    s = str(text).strip()

    acc = ""
    srv = ""

    # 1. Regex search for explicit labels (Arabic and English)
    m_acc = re.search(r'(?:حساب|عقد|account|contract)[\s:]*([0-9]{8,12})', s, re.IGNORECASE)
    if m_acc:
        acc = m_acc.group(1)

    m_srv = re.search(r'(?:خدمة|خدمه|جوال|هاتف|محفظة|محفظه|service|phone|wallet)[\s:]*([0-9]{8,12})', s, re.IGNORECASE)
    if m_srv:
        srv = m_srv.group(1)

    # If both found via labels, return immediately
    if acc and srv:
        return acc, srv

    # 2. Extract all digit clusters (min 8 digits for valid Zain accounts/services)
    digit_blocks = re.findall(r'\b[0-9]{8,12}\b', s)
    if not digit_blocks:
        digit_blocks = re.findall(r'[0-9]{6,14}', s)

    for block in digit_blocks:
        # Zain contracts start with 1 (usually 10...)
        if block.startswith("10") or (block.startswith("1") and len(block) >= 9):
            if not acc:
                acc = block
        # Zain services/wallets start with 2 or 5 or 05 or 966
        elif block.startswith("2") or block.startswith("5") or block.startswith("05") or block.startswith("966"):
            if not srv:
                srv = block
        else:
            if not acc:
                acc = block
            elif not srv:
                srv = block

    # If only one was found and no explicit label
    if not acc and not srv and digit_blocks:
        b0 = digit_blocks[0]
        if b0.startswith("2"):
            srv = b0
        else:
            acc = b0

    return acc, srv


def get_sheet_header_columns(sheet, has_headers: bool = True) -> list[dict[str, Any]]:
    """Returns a list of all header columns in Row 1 with their index and letters."""
    headers = []
    first_row = next(
        sheet.iter_rows(min_row=HEADER_ROW_NUMBER, max_row=HEADER_ROW_NUMBER, values_only=True),
        (),
    )
    for idx, val in enumerate(first_row, start=1):
        headers.append({
            "col_index": idx,
            "col_letter": get_column_letter(idx),
            "header": str(val).strip() if has_headers and val is not None else "",
        })
    return headers


def detect_smart_sheet_columns(sheet: Any, has_headers: bool = True) -> dict[str, Any]:
    """Analyzes Row 1 of an Excel sheet and discovers all critical columns:
    - رقم الحساب (lookup_col)
    - رقم الخدمة (service_col)
    - مبلغ العقد (contract_col) -> AW in 2.xlsx
    - متبقي سداد موثق (remaining_col) -> P in 2.xlsx
    - اسم العميل (customer_col)
    - المحصل (collector_col)
    - رقم الصف الأصلي (source_row_col)
    """
    if isinstance(sheet, (str, Path)):
        wb = load_fast_workbook(sheet)
        try:
            ws = wb.worksheets[0]
            return detect_smart_sheet_columns(ws, has_headers=has_headers)
        finally:
            wb.close()

    header_row = next(
        sheet.iter_rows(min_row=HEADER_ROW_NUMBER, max_row=HEADER_ROW_NUMBER, values_only=True),
        (),
    )
    if not has_headers:
        header_row = tuple(None for _ in header_row)

    detected: dict[str, Optional[int]] = {
        "customer_col": None,
        "lookup_col": None,
        "contract_col": None,
        "remaining_col": None,
        "debt_col": None,
        "amount_col": None,
        "amount_col_2": None,
        "service_col": None,
        "collector_col": None,
        "case_status_col": None,
        "main_status_col": None,
        "sub_status_col": None,
        "notes_col": None,
        "source_row_col": None,
    }

    contract_col: Optional[int] = None
    remaining_col: Optional[int] = None
    debt_col: Optional[int] = None
    sheet_amount_col: Optional[int] = None
    live_amount_col: Optional[int] = None
    general_amount_col: Optional[int] = None
    source_row_col: Optional[int] = None

    def _matches_any(text: str, synonyms: tuple[str, ...]) -> bool:
        return any(normalize_header_text(s) in text for s in synonyms)

    def _has_negative_keywords(text: str, neg_words: tuple[str, ...]) -> bool:
        return any(w in text for w in neg_words)

    for idx, raw_val in enumerate(header_row, start=1):
        norm = normalize_header_text(raw_val)
        if not norm:
            continue

        # 1. اسم العميل
        if (
            not detected["customer_col"]
            and not _has_negative_keywords(norm, ("ارقام", "تصنيف", "متابعة", "تاريخ", "رقم"))
            and _matches_any(norm, ("اسم العميل", "العميل", "الاسم", "اسم المشترك", "المشترك", "customer", "name"))
        ):
            detected["customer_col"] = idx

        # 2. رقم الصف في الشيت (Master Row Number from previous audit or raw file)
        elif (
            not source_row_col
            and _matches_any(norm, (
                "رقم الصف في الشيت", "رقم الصف بالشيت", "رقم الصف", "الصف في الشيت",
                "الصف", "source row", "row number", "row no"
            ))
        ):
            source_row_col = idx

        # 3. رقم الحساب / العقد / الخدمة المزدوج
        elif (
            not detected["lookup_col"]
            and not _has_negative_keywords(norm, ("تاريخ", "سدادات", "متبقي", "مبلغ", "قيمة", "نوع", "صف"))
            and _matches_any(norm, (
                "رقم الحساب ورقم الخدمة", "رقم الحساب ورقم الخدمه", "الحساب / الخدمة", "الحساب الخدمه",
                "الحساب/الخدمة", "رقم الحساب والخدمة", "رقم الحساب أو الخدمة", "الحساب أو الخدمة",
                "رقم الحساب", "رقم العقد", "الحساب", "العقد", "رقم الفاتورة", "رقم البحث", "search number", "lookup number", "account", "contract"
            ))
            or (not detected["lookup_col"] and norm in {"number", "رقم", "الرقم"})
        ):
            detected["lookup_col"] = idx

        # 4. المبلغ في الشيت / المبلغ المسجل بالشيت (Expected Amount in Results/Errors sheet)
        elif (
            not sheet_amount_col
            and not _has_negative_keywords(norm, ("نوع", "رقم", "عمر", "تاريخ", "كود", "هاتف", "جوال", "ملاحظات", "متابعة", "زين"))
            and _matches_any(norm, (
                "المبلغ المسجل بالشيت", "المبلغ في الشيت", "المبلغ بالشيت", "مبلغ الشيت",
                "المبلغ المسجل", "مبلغ في الشيت", "قيمة الشيت", "المبلغ المتوقع", "expected amount"
            ))
        ):
            sheet_amount_col = idx

        # 5. مبلغ موقع زين المباشر (Live Zain amount in results sheets)
        elif (
            not live_amount_col
            and _matches_any(norm, (
                "المبلغ الحالي في موقع زين", "المبلغ في موقع زين", "المبلغ بموقع زين",
                "مبلغ زين", "رصيد زين", "المبلغ الحالي", "live amount"
            ))
        ):
            live_amount_col = idx

        # 6. مبلغ العقد (المطابق لموقع زين عند البحث بالحساب في 2.xlsx)
        elif (
            not contract_col
            and not _has_negative_keywords(norm, AMOUNT_EXCLUSIONS)
            and _matches_any(norm, (
                "مبلغ العقد", "قيمة العقد", "إجمالي العقد", "اجمالي العقد",
                "مبلغ عقد", "قيمة عقد", "total contract", "contract amount"
            ))
        ):
            contract_col = idx

        # 7. متبقي سداد موثق (المبلغ المتبقي)
        elif (
            not remaining_col
            and not _has_negative_keywords(norm, AMOUNT_EXCLUSIONS)
            and _matches_any(norm, (
                "متبقي سداد موثق", "باقي السداد الموثق", "باقي السداد", "متبقي السداد",
                "المبلغ المتبقي", "الرصيد المتبقي", "المديونية المتبقية", "صافي المديونية", "remaining"
            ))
        ):
            remaining_col = idx

        # 8. مبلغ المديونية الإجمالية
        elif (
            not debt_col
            and not _has_negative_keywords(norm, AMOUNT_EXCLUSIONS)
            and _matches_any(norm, (
                "مبلغ المديونية", "مبلغ الميدونية", "مبلغ المطالبة", "المبلغ الإجمالي", "الإجمالي", "total amount"
            ))
        ):
            debt_col = idx

        # 9. مبلغ عام بديل
        elif (
            not general_amount_col
            and not _has_negative_keywords(norm, AMOUNT_EXCLUSIONS)
            and _matches_any(norm, ("المبلغ", "المبلغ المطلوب", "المبلغ المستحق", "قيمة الفاتورة", "amount"))
        ):
            general_amount_col = idx

        # 10. رقم الخدمة المنفصل
        elif (
            not detected["service_col"]
            and not _has_negative_keywords(norm, ("تاريخ", "نوع", "مبلغ"))
            and _matches_any(norm, (
                "رقم الخدمة", "رقم الخدمه", "الخدمة", "الخدمه", "رقم المحفظة", "رقم المحفظه",
                "المحفظة", "رقم الجوال", "الجوال", "رقم الهاتف", "الهاتف", "service", "phone"
            ))
        ):
            detected["service_col"] = idx

        # 11. المحصل
        elif (
            not detected["collector_col"]
            and not _has_negative_keywords(norm, ("تاريخ", "اسناد", "اشراف", "مشرف", "مستخدم", "فرع"))
            and _matches_any(norm, ("المحصل", "اسم المحصل", "الموظف", "collector"))
        ):
            detected["collector_col"] = idx

        # 12. الحالات
        elif not detected["main_status_col"] and _matches_any(norm, ("الحالة الرئيسية", "الحاله الرئيسيه", "main status")):
            detected["main_status_col"] = idx
        elif not detected["sub_status_col"] and _matches_any(norm, ("الحالة الفرعية", "الحاله الفرعيه", "sub status")):
            detected["sub_status_col"] = idx
        elif not detected["case_status_col"] and norm in ("الحالة", "الحاله", "حالة", "حاله", "status", "case status"):
            detected["case_status_col"] = idx
        elif not detected["notes_col"] and _matches_any(norm, (
            "المشكلة باختصار", "المشكله باختصار", "المشكلة", "المشكله", "سبب الخطا", "سبب الخطأ",
            "الخطا", "الخطأ", "أخر متابعة للمحصل", "اخر متابعة للمحصل", "المتابعة", "المتابعه", "ملاحظات", "notes"
        )):
            detected["notes_col"] = idx

    # If sheet has a combined lookup column (e.g. 'رقم الحساب ورقم الخدمة' or 'الحساب / الخدمة') and service_col wasn't separate
    if detected["lookup_col"] and not detected["service_col"]:
        detected["service_col"] = detected["lookup_col"]

    # Reconcile Amounts
    if sheet_amount_col:
        contract_col = contract_col or sheet_amount_col
        remaining_col = remaining_col or sheet_amount_col
    elif general_amount_col:
        contract_col = contract_col or general_amount_col
        remaining_col = remaining_col or general_amount_col

    detected["contract_col"] = contract_col
    detected["remaining_col"] = remaining_col
    detected["debt_col"] = debt_col
    detected["source_row_col"] = source_row_col

    # Primary & Secondary amounts
    detected["amount_col"] = sheet_amount_col or remaining_col or debt_col or contract_col or general_amount_col
    detected["amount_col_2"] = live_amount_col or contract_col or debt_col or detected["amount_col"]

    if not detected["amount_col"] and detected["amount_col_2"]:
        detected["amount_col"] = detected["amount_col_2"]
    elif not detected["amount_col_2"] and detected["amount_col"]:
        detected["amount_col_2"] = detected["amount_col"]

    # Infer missing fields only from an unambiguous sample. Never guess among
    # multiple candidate number or amount columns; the user can select them.
    from itertools import islice
    sample = list(islice(sheet.iter_rows(min_row=2 if has_headers else 1, values_only=True), 5))
    candidates = []
    for col in range(1, (sheet.max_column or 0) + 1):
        values = [row[col-1] for row in sample if col <= len(row) and row[col-1] is not None and str(row[col-1]).strip()]
        if values and all(any(n.startswith(("1", "2")) for n in parse_account_and_service(v) if n) for v in values):
            candidates.append(col)
    if not detected["lookup_col"] and not detected["service_col"] and len(candidates) == 1:
        detected["lookup_col"] = detected["service_col"] = candidates[0]
    if not detected["amount_col"]:
        excluded = set(candidates) | {detected["lookup_col"], detected["service_col"]}
        amounts = []
        for col in range(1, (sheet.max_column or 0) + 1):
            if col in excluded:
                continue
            values = [row[col-1] for row in sample if col <= len(row) and row[col-1] is not None and str(row[col-1]).strip()]
            if values and all(parse_money_to_halalas(v) is not None for v in values):
                amounts.append(col)
        if len(amounts) == 1:
            for key in ("amount_col", "amount_col_2", "remaining_col", "contract_col"):
                detected[key] = amounts[0]

    letters = {
        k: (get_column_letter(v) if v else None)
        for k, v in detected.items()
    }
    return {"indices": detected, "letters": letters}


def inspect_sheet_schema(sheet: Any, has_headers: bool = True, count_rows: bool = True) -> dict[str, Any]:
    """Provides complete schema analysis, detected columns, and mode compatibility."""
    if isinstance(sheet, (str, Path)):
        wb = load_fast_workbook(sheet)
        try:
            ws = wb.worksheets[0]
            return inspect_sheet_schema(ws, has_headers=has_headers, count_rows=count_rows)
        finally:
            wb.close()

    header_cols = get_sheet_header_columns(sheet, has_headers)
    smart_info = detect_smart_sheet_columns(sheet, has_headers)
    indices = smart_info["indices"]
    letters = smart_info["letters"]

    has_acc = bool(indices.get("lookup_col"))
    has_srv = bool(indices.get("service_col"))
    has_amt1 = bool(indices.get("amount_col"))
    has_amt2 = bool(indices.get("amount_col_2"))
    contract_c = indices.get("contract_col") or indices.get("amount_col_2")
    remaining_c = indices.get("remaining_col") or indices.get("amount_col")

    # Document type detection
    sheet_title = getattr(sheet, "title", "") or ""
    norm_title = normalize_header_text(sheet_title)
    is_errors_tab = any(k in norm_title for k in ("اخطاء", "ملاحظات", "errors", "notes"))
    is_results_tab = any(k in norm_title for k in ("فروقات", "نتائج", "results", "diff"))

    has_error_header = any(
        normalize_header_text(c.get("header")) in ("المشكله باختصار", "المشكلة باختصار", "سبب الخطا", "سبب الخطأ")
        for c in header_cols
    )
    has_live_header = any(
        "زين" in normalize_header_text(c.get("header")) or "فرق" in normalize_header_text(c.get("header"))
        for c in header_cols
    )

    if is_errors_tab or has_error_header:
        doc_type = "🛠️ شيت أخطاء وملاحظات فحص سابقة (Errors Sheet - جاهز لإعادة التدقيق ومعالجة الأخطاء)"
        sheet_kind = "errors_sheet"
    elif is_results_tab or has_live_header:
        doc_type = "📊 شيت نتائج وفروقات فحص سابقة (Results Sheet - جاهز لإعادة التدقيق والتحديث)"
        sheet_kind = "results_sheet"
    elif indices.get("lookup_col") and indices.get("lookup_col") == indices.get("service_col") and has_amt1:
        doc_type = "ملف أرقام ومبالغ"
        sheet_kind = "raw_portfolio"
    elif has_acc and has_srv and bool(contract_c):
        doc_type = "محفظة تحصيل مديونيات زين (شاملة أرقام الحسابات وعقود الخدمات ومبالغ العقود)"
        sheet_kind = "raw_portfolio"
    elif has_acc and has_srv:
        doc_type = "محفظة مديونيات هجينة (أرقام حسابات وأرقام خدمات)"
        sheet_kind = "raw_portfolio"
    elif has_acc:
        doc_type = "شيت حسابات زين (Account Numbers Sheet)"
        sheet_kind = "raw_portfolio"
    elif has_srv:
        doc_type = "شيت محافظ/خدمات زين (Wallets/Services Sheet)"
        sheet_kind = "raw_portfolio"
    else:
        doc_type = "شيت بيانات عام / بحاجة لضبط الأعمدة"
        sheet_kind = "unknown"

    is_acceptable = bool((has_acc or has_srv) and (has_amt1 or has_amt2))

    if is_acceptable:
        if sheet_kind == "errors_sheet":
            ver_status = "✓ شيت أخطاء صالح ومكتمل - جاهز لإعادة الفحص ومعالجة الأخطاء فوراً"
        elif sheet_kind == "results_sheet":
            ver_status = "✓ شيت نتائج صالح ومكتمل - جاهز لإعادة التدقيق وتحديث الفروقات فوراً"
        else:
            ver_status = "✓ الشيت صالح ومقبول للتدقيق"
    else:
        ver_status = "اختر عمود الرقم وعمود المبلغ لإكمال التحقق"

    supported = {
        "smart_hybrid": bool((has_acc or has_srv) and (has_amt1 or has_amt2)),
        "account_only": bool(has_acc and (has_amt1 or has_amt2)),
        "service_only": bool(has_srv and (has_amt1 or has_amt2)),
    }

    if supported["smart_hybrid"]:
        rec_mode = "smart_hybrid"
    elif supported["account_only"]:
        rec_mode = "account_only"
    elif supported["service_only"]:
        rec_mode = "service_only"
    else:
        rec_mode = "custom"

    amt_options = []
    if sheet_kind in ("errors_sheet", "results_sheet"):
        if remaining_c:
            rem_l = get_column_letter(remaining_c)
            amt_options.append({
                "id": "remaining",
                "col_index": remaining_c,
                "letter": rem_l,
                "label": f"المبلغ المسجل بالشيت [{rem_l}] (المعتمد للفحص)",
            })
        if contract_c and contract_c != remaining_c:
            contract_l = get_column_letter(contract_c)
            amt_options.append({
                "id": "contract",
                "col_index": contract_c,
                "letter": contract_l,
                "label": f"المبلغ السابق بموقع زين [{contract_l}]",
            })
    else:
        if contract_c:
            contract_l = get_column_letter(contract_c)
            amt_options.append({
                "id": "contract",
                "col_index": contract_c,
                "letter": contract_l,
                "label": f"مبلغ العقد [{contract_l}] (المطابق لموقع زين عند البحث برقم الحساب)",
            })
        if remaining_c and remaining_c != contract_c:
            rem_l = get_column_letter(remaining_c)
            amt_options.append({
                "id": "remaining",
                "col_index": remaining_c,
                "letter": rem_l,
                "label": f"باقي السداد الموثق [{rem_l}]",
            })
        if contract_c and remaining_c and remaining_c != contract_c:
            contract_l = get_column_letter(contract_c)
            rem_l = get_column_letter(remaining_c)
            amt_options.append({
                "id": "smart_dual",
                "col_index": contract_c,
                "letter": f"{contract_l}+{rem_l}",
                "label": f"كلاهما (فحص ذكي مزدوج: {contract_l} أو {rem_l} أيهما أقرب)",
            })

    # Excel's max_row includes formatting and autofilled helper columns. Count
    # meaningful input rows; validation callers already obtain this count themselves.
    first_data_row = 2 if has_headers else 1
    mapped_columns = [value for value in indices.values() if value]
    est_rows = (sum(not is_empty_input_row(values, mapped_columns)
                    for values in sheet.iter_rows(min_row=first_data_row, values_only=True))
                if count_rows else None)
    from itertools import islice
    preview_rows = [{"row_number": row_no, "values": [str(v) if v is not None else "" for v in values]}
                    for row_no, values in enumerate(islice(sheet.iter_rows(min_row=first_data_row, values_only=True), 5), first_data_row)]

    return {
        "columns": header_cols,
        "has_headers": has_headers,
        "preview_rows": preview_rows,
        "indices": indices,
        "letters": letters,
        "document_type": doc_type,
        "is_acceptable": is_acceptable,
        "verification_status": ver_status,
        "estimated_rows": est_rows,
        "supported_modes": supported,
        "recommended_mode": rec_mode,
        "amount_options": amt_options,
        "has_customer_col": bool(indices.get("customer_col")),
        "has_collector_col": bool(indices.get("collector_col")),
        "contract_amount_col": contract_c,
        "contract_amount_letter": get_column_letter(contract_c) if contract_c else None,
        "remaining_amount_col": remaining_c,
        "remaining_amount_letter": get_column_letter(remaining_c) if remaining_c else None,
        "source_row_col": indices.get("source_row_col"),
        "source_row_letter": letters.get("source_row_col"),
    }


def extract_customer_records(
    sheet: Any,
    lookup_col: int,
    amount_col: int,
    amount_col_2: Optional[int] = None,
    service_col: Optional[int] = None,
    customer_col: Optional[int] = None,
    collector_col: Optional[int] = None,
    case_status_col: Optional[int] = None,
    main_status_col: Optional[int] = None,
    sub_status_col: Optional[int] = None,
    notes_col: Optional[int] = None,
    source_row_col: Optional[int] = None,
    target_collector: str = "",
    record_type: str = "mixed",
    has_headers: bool = True,
    clean_dedup: bool = True,
) -> list[Customer]:
    """Constructs Customer objects, preserving row 1 in files without headers.
    Supports raw portfolios, previously exported results sheets, and errors sheets.
    """
    if isinstance(sheet, (str, Path)):
        wb = load_fast_workbook(sheet)
        try:
            ws = wb.worksheets[0]
            return extract_customer_records(
                sheet=ws,
                lookup_col=lookup_col,
                amount_col=amount_col,
                amount_col_2=amount_col_2,
                service_col=service_col,
                customer_col=customer_col,
                collector_col=collector_col,
                case_status_col=case_status_col,
                main_status_col=main_status_col,
                sub_status_col=sub_status_col,
                notes_col=notes_col,
                source_row_col=source_row_col,
                target_collector=target_collector,
                record_type=record_type,
                has_headers=has_headers,
                clean_dedup=clean_dedup,
            )
        finally:
            wb.close()

    cols_to_read = []
    for c in (lookup_col, amount_col, amount_col_2, service_col, customer_col, collector_col,
              case_status_col, main_status_col, sub_status_col, notes_col, source_row_col):
        if c:
            try:
                cols_to_read.append(int(c))
            except (ValueError, TypeError):
                pass
    max_c = max(cols_to_read) if cols_to_read else 30
    max_c = max(max_c, 30)

    rows_iter = sheet.iter_rows(
        min_row=2 if has_headers else 1,
        min_col=1,
        max_col=max_c,
        values_only=True,
    )

    norm_target_collector = normalize_header_text(target_collector) if target_collector else ""
    customers: list[Customer] = []
    seen_non2_account_row: set[str] = set()
    seen_service_row: set[str] = set()
    seen_account_row: set[str] = set()

    for row_num, values in enumerate(rows_iter, start=2 if has_headers else 1):
        if is_empty_input_row(values, cols_to_read):
            continue

        # Optional collector filter
        raw_coll = str(values[collector_col - 1] or "").strip() if (collector_col and len(values) >= collector_col) else ""
        if norm_target_collector and normalize_header_text(raw_coll) != norm_target_collector:
            continue

        # Extract Original Master Row Number if present in results/errors sheet
        final_row_num = row_num
        if source_row_col and len(values) >= source_row_col:
            raw_sr = values[source_row_col - 1]
            if raw_sr is not None:
                clean_sr = re.sub(r"\D", "", str(raw_sr))
                if clean_sr:
                    final_row_num = int(clean_sr)

        raw_lookup = str(values[lookup_col - 1] or "").strip() if len(values) >= lookup_col else ""
        raw_serv = str(values[service_col - 1] or "").strip() if (service_col and len(values) >= service_col) else ""

        if not raw_lookup and not raw_serv:
            continue

        # Intelligently parse compound strings like 'حساب: 1006659013 | خدمة: 2006695911'
        extracted_acc, extracted_srv = parse_account_and_service(raw_lookup)
        if raw_serv and raw_serv != raw_lookup:
            s_acc, s_srv = parse_account_and_service(raw_serv)
            if s_acc and not extracted_acc:
                extracted_acc = s_acc
            if s_srv and not extracted_srv:
                extracted_srv = s_srv

        clean_lookup = extracted_acc or re.sub(r"\D", "", raw_lookup)
        clean_serv = extracted_srv or re.sub(r"\D", "", raw_serv)

        if not clean_lookup and not clean_serv:
            continue

        # Parse Amounts into integer Halalas
        raw_amt1 = values[amount_col - 1] if len(values) >= amount_col else None
        amt1_halalas = parse_money_to_halalas(raw_amt1) or 0

        amt2_halalas = None
        if amount_col_2 and len(values) >= amount_col_2:
            raw_amt2 = values[amount_col_2 - 1]
            amt2_halalas = parse_money_to_halalas(raw_amt2)
        if amt2_halalas is None:
            amt2_halalas = amt1_halalas

        cust_name = str(values[customer_col - 1] or "").strip() if (customer_col and len(values) >= customer_col) else "عميل غير محدد"
        c_stat = str(values[case_status_col - 1] or "").strip() if (case_status_col and len(values) >= case_status_col) else ""
        m_stat = str(values[main_status_col - 1] or "").strip() if (main_status_col and len(values) >= main_status_col) else ""
        s_stat = str(values[sub_status_col - 1] or "").strip() if (sub_status_col and len(values) >= sub_status_col) else ""
        notes = str(values[notes_col - 1] or "").strip() if (notes_col and len(values) >= notes_col) else ""

        # Rich Metadata Extraction from Sheet (if in master 2.xlsx format)
        nat_id = str(values[7] or "") if len(values) >= 8 else ""
        ph_nums = str(values[8] or "") if len(values) >= 9 else ""
        br_name = str(values[19] or "") if len(values) >= 20 else ""
        sup_name = str(values[20] or "") if len(values) >= 21 else ""
        f_date = str(values[24] or "") if len(values) >= 25 else ""

        # Determine effective search number and record type based on requested mode/record_type
        acc_num = extracted_acc
        if not acc_num and clean_lookup and not clean_lookup.startswith("2"):
            acc_num = clean_lookup

        srv_num = extracted_srv
        if not srv_num and clean_serv and clean_serv.startswith("2"):
            srv_num = clean_serv
        if not srv_num and clean_lookup and clean_lookup.startswith("2"):
            srv_num = clean_lookup

        if record_type == "account":
            if not acc_num:
                continue
            primary_search = acc_num
            effective_rec_type = "account"
        elif record_type == "wallet":
            if not srv_num:
                continue
            primary_search = srv_num
            effective_rec_type = "wallet"
        else:
            if clean_serv.startswith("2"):
                primary_search = clean_serv
                rec_type = "wallet"
            elif clean_lookup:
                primary_search = clean_lookup
                rec_type = "account"
            else:
                primary_search = clean_serv
                rec_type = "wallet" if clean_serv.startswith("2") else "account"
            effective_rec_type = rec_type

        if clean_dedup:
            if record_type == "account":
                if acc_num and acc_num in seen_account_row:
                    continue
                if acc_num:
                    seen_account_row.add(acc_num)
            elif record_type == "wallet":
                if srv_num and srv_num in seen_service_row:
                    continue
                if srv_num:
                    seen_service_row.add(srv_num)
            else:
                starts_with_2 = bool(srv_num and srv_num.startswith("2"))
                if starts_with_2:
                    if srv_num in seen_service_row:
                        continue
                    seen_service_row.add(srv_num)
                else:
                    if acc_num and acc_num in seen_non2_account_row:
                        continue
                    if acc_num:
                        seen_non2_account_row.add(acc_num)
                    if srv_num:
                        seen_service_row.add(srv_num)

        customers.append(Customer(
            row_number=final_row_num,
            record_type=effective_rec_type,
            lookup_number=primary_search,
            contract=clean_lookup or primary_search,
            expected_amount=amt1_halalas,
            expected_amount_2=amt2_halalas,
            original_account_number=clean_lookup or primary_search,
            service_number=clean_serv,
            customer_name=cust_name or "عميل غير محدد",
            collector_name=raw_coll,
            case_status=c_stat,
            main_status=m_stat,
            sub_status=s_stat,
            notes=notes,
            national_id=nat_id,
            phones=ph_nums,
            supervisor_name=sup_name,
            branch_name=br_name,
            followup_date=f_date,
            error_or_review_details=notes if notes else None,
        ))

    return customers

