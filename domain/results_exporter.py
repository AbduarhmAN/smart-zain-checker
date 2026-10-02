"""Domain Results Exporter for Smart Zain Checker.
Creates, formats, and safely appends records to 'نتائج فحص زين.xlsx'.
Provides the official three-sheet audit deliverable.
"""
from __future__ import annotations

import threading
from datetime import datetime
from pathlib import Path
from typing import Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from domain.models import CheckError, Mismatch
from domain.money import halalas_to_sar_float

FILE_LOCK = threading.Lock()

FONT_NAME = "Segoe UI"
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name=FONT_NAME, size=10)
DIFF_FONT = Font(name=FONT_NAME, size=10, bold=True, color="991B1B")

HEADER_FILL_MISMATCH = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
HEADER_FILL_ERROR = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
HEADER_FILL_SUMMARY = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
ROW_ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

MISMATCH_HEADERS = (
    "رقم السطر",
    "اسم العميل",
    "رقم الهوية",
    "رقم العقد",
    "رقم الحساب",
    "رقم الخدمة",
    "أرقام جوال العميل",
    "المحصل",
    "المشرف",
    "الفرع",
    "المبلغ المعتمد بالملف (ر.س)",
    "المبلغ بموقع زين (ر.س)",
    "صافي الفرق (ر.س)",
    "حالة السداد بموقع زين",
    "الحالة الرئيسية بالملف",
    "الحالة الفرعية بالملف",
    "ملاحظات المتابعة",
    "تاريخ المتابعة",
    "وقت الرصد",
)

FILL_PAID = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
FILL_PARTIAL = PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid")
FILL_INCREASE = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
DIFF_FONT_GREEN = Font(name=FONT_NAME, size=10, bold=True, color="166534")
DIFF_FONT_BLUE = Font(name=FONT_NAME, size=10, bold=True, color="1E40AF")
DIFF_FONT_AMBER = Font(name=FONT_NAME, size=10, bold=True, color="92400E")


ERROR_HEADERS = (
    "أرقام السطور",
    "نوع السجل",
    "رقم البحث",
    "اسم العميل",
    "المحصل",
    "رمز الخطأ",
    "تفاصيل وسبب الخطأ",
    "المبلغ المتوقع بالملف (ر.س)",
    "وقت الرصد",
)


def ensure_audit_workbook_exists(file_path: Path) -> None:
    """Initializes the Excel audit workbook if it does not already exist."""
    with FILE_LOCK:
        if file_path.exists():
            return
        
        file_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        
        # 1. Sheet: الفروقات الصافية للمحصلين
        ws_mismatch = wb.active
        ws_mismatch.title = "الفروقات الصافية للمحصلين"
        ws_mismatch.sheet_view.rightToLeft = True
        ws_mismatch.append(MISMATCH_HEADERS)
        for col_idx in range(1, len(MISMATCH_HEADERS) + 1):
            cell = ws_mismatch.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL_MISMATCH
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER

        # 2. Sheet: المهلات والأخطاء
        ws_err = wb.create_sheet(title="المهلات والأخطاء")
        ws_err.sheet_view.rightToLeft = True
        ws_err.append(ERROR_HEADERS)
        for col_idx in range(1, len(ERROR_HEADERS) + 1):
            cell = ws_err.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL_ERROR
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER

        # 3. Sheet: ملخص الفحص
        ws_summary = wb.create_sheet(title="ملخص الفحص")
        ws_summary.sheet_view.rightToLeft = True
        ws_summary.append(["المؤشر الرقابي", "القيمة"])
        for col_idx in (1, 2):
            cell = ws_summary.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL_SUMMARY
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER

        wb.save(file_path)


def append_mismatch_record(file_path: Path, mismatch: Mismatch) -> None:
    """Appends a single verified mismatch row to the workbook."""
    ensure_audit_workbook_exists(file_path)
    cust = mismatch.customer

    # Pick effective expected amount (closest to Zain)
    eff_expected = cust.expected_amount
    if cust.expected_amount_2 is not None:
        d1 = abs(mismatch.website_amount - cust.expected_amount)
        d2 = abs(mismatch.website_amount - cust.expected_amount_2)
        eff_expected = cust.expected_amount if d1 <= d2 else cust.expected_amount_2

    file_sar = halalas_to_sar_float(eff_expected)
    zain_sar = halalas_to_sar_float(mismatch.website_amount)
    diff_sar = round(zain_sar - file_sar, 2)

    if zain_sar == 0.0:
        pay_status = "✔ مسدد بالكامل (0.00 ر.س)"
        status_fill = FILL_PAID
        diff_font = DIFF_FONT_GREEN
    elif diff_sar < 0:
        pay_status = f"سداد جزئي (فرق {abs(diff_sar):.2f} ر.س)"
        status_fill = FILL_PARTIAL
        diff_font = DIFF_FONT_BLUE
    else:
        pay_status = f"زيادة مطالبة (+{diff_sar:.2f} ر.س)"
        status_fill = FILL_INCREASE
        diff_font = DIFF_FONT_AMBER

    row_data = (
        cust.row_number,
        cust.customer_name or "غير محدد",
        getattr(cust, "national_id", "") or "-",
        cust.contract or cust.lookup_number or "-",
        cust.original_account_number or cust.lookup_number or "-",
        cust.service_number or "-",
        getattr(cust, "phones", "") or "-",
        cust.collector_name or "-",
        getattr(cust, "supervisor_name", "") or "-",
        getattr(cust, "branch_name", "") or "-",
        file_sar,
        zain_sar,
        diff_sar,
        pay_status,
        cust.main_status or "-",
        cust.sub_status or "-",
        cust.notes or "-",
        getattr(cust, "followup_date", "") or "-",
        mismatch.detected_at.strftime("%Y-%m-%d %H:%M:%S"),
    )

    with FILE_LOCK:
        for attempt in range(4):
            try:
                wb = load_workbook(file_path)
                try:
                    ws = wb["الفروقات الصافية للمحصلين"] if "الفروقات الصافية للمحصلين" in wb.sheetnames else wb.worksheets[0]
                    new_row_idx = ws.max_row + 1
                    ws.append(row_data)

                    for col_idx in range(1, len(row_data) + 1):
                        cell = ws.cell(row=new_row_idx, column=col_idx)
                        cell.font = DATA_FONT
                        cell.border = THIN_BORDER
                        cell.alignment = ALIGN_CENTER if col_idx in (1, 3, 4, 5, 6, 14, 18, 19) else ALIGN_RIGHT
                        if col_idx in (11, 12, 13):
                            cell.number_format = "#,##0.00"
                            if col_idx == 13:
                                cell.font = diff_font
                        if col_idx == 14:
                            cell.fill = status_fill
                        elif new_row_idx % 2 == 1:
                            cell.fill = ROW_ALT_FILL

                    wb.save(file_path)
                    break
                finally:
                    wb.close()
            except PermissionError:
                if attempt < 3:
                    import time
                    time.sleep(1.0)
                else:
                    import logging
                    logging.getLogger("ResultsExporter").warning(f"Could not save {file_path} because file is locked by user (attempt {attempt+1}).")
            except Exception as e:
                import logging
                logging.getLogger("ResultsExporter").error(f"Error saving to {file_path}: {e}")
                break


def append_error_record(file_path: Path, error: CheckError) -> None:
    """Appends an error or timeout record to the error sheet."""
    ensure_audit_workbook_exists(file_path)
    amt_sar = halalas_to_sar_float(error.expected_amount) if error.expected_amount is not None else 0.0

    row_data = (
        ", ".join(str(r) for r in error.row_numbers),
        "محفظة" if error.record_type == "wallet" else "حساب",
        error.lookup_number,
        error.customer_name or "-",
        error.collector_name or "-",
        error.error_code,
        error.error_details,
        amt_sar,
        error.detected_at.strftime("%Y-%m-%d %H:%M:%S"),
    )

    with FILE_LOCK:
        wb = load_workbook(file_path)
        try:
            ws = wb["المهلات والأخطاء"] if "المهلات والأخطاء" in wb.sheetnames else wb.create_sheet("المهلات والأخطاء")
            new_row_idx = ws.max_row + 1
            ws.append(row_data)

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=new_row_idx, column=col_idx)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.alignment = ALIGN_CENTER if col_idx in (1, 2, 6, 9) else ALIGN_RIGHT
                if col_idx == 8:
                    cell.number_format = "#,##0.00"
                if new_row_idx % 2 == 0:
                    cell.fill = ROW_ALT_FILL

            wb.save(file_path)
        finally:
            wb.close()
