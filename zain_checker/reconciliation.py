"""Stage 16: Scope-Safe Post-Zain Result Binding, Reconciliation & Audit Workbook Exporter.

Implements:
- Section 1.C: The Four Unbreakable Laws:
  1. Strict Scope Symmetry (Service-to-Service vs Account-to-Account when Complete)
  2. Zero Residual Deduction (Never subtract Service_2 from Account_Total to infer Service_Non2)
  3. No Unfounded Root-Cause Inference (Strict variance descriptions without 'Paid' or 'Settlement' claims)
  4. Error != Zero (Search errors/timeouts/redirects record None, never 0.00)
- Section 9 (Task 9): Post-Zain Result Binding across the 4 outcome categories
  (Service Result, Account Result, Search Error, Ambiguous Result)
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Dict, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .domain_models import (
    AccountGroup,
    PreZainSearchPlan,
    ReconciliationStatus,
    SearchTask,
    SourceRow,
)


@dataclass
class ZainTaskObservation:
    """Represents the raw observation returned from Zain for a specific SearchTask."""

    task_id: str
    outcome_type: str  # "SUCCESS" | "SEARCH_ERROR" | "AMBIGUOUS"
    observed_amount: Optional[Decimal] = None
    raw_screen_text: Optional[str] = None
    error_or_ambiguity_reason: Optional[str] = None


def format_variance_description(file_amount: Decimal, zain_amount: Decimal) -> str:
    """Law #3 (No Unfounded Root-Cause Inference): Formats variance without guessing 'Paid' or 'Settlement'."""
    if file_amount == zain_amount:
        return "تطابق مالي في نفس النطاق (Same-Scope Match)"
    if file_amount == Decimal("0.00") and zain_amount > Decimal("0.00"):
        return "فرق مالي (File = 0, Zain > 0)"
    if file_amount > zain_amount:
        return "فرق مالي (File > Zain)"
    return "فرق مالي (File < Zain)"


def bind_zain_observations_to_plan(
    plan: PreZainSearchPlan,
    observations: Dict[str, ZainTaskObservation],
) -> PreZainSearchPlan:
    """Binds post-Zain search observations back to tasks, account groups, and source rows
    while strictly enforcing the Four Unbreakable Laws (Section 1.C & Section 9).
    """
    row_map: Dict[str, SourceRow] = {r.row_uid: r for r in plan.source_rows}
    acc_map: Dict[str, AccountGroup] = {g.account_number: g for g in plan.account_groups}

    # 1. Process Service Search Tasks (1:1 Service-Level Scope)
    for task in plan.service_search_tasks:
        obs = observations.get(task.task_id)
        if obs is None:
            continue

        # Law #4: Error != Zero
        if obs.outcome_type == "SEARCH_ERROR":
            task.search_status = "SEARCH_ERROR"
            task.zain_observed_amount = None
            task.raw_observed_screen_text = obs.raw_screen_text
            task.error_or_review_details = obs.error_or_ambiguity_reason or "Search / DOM Timeout Error"
            task.reconciliation_status = ReconciliationStatus.SEARCH_ERROR_NEEDS_REVIEW
            task.reconciliation_variance = None
            task.reconciliation_note = (
                f"SEARCH_ERROR_NEEDS_REVIEW ({task.error_or_review_details}) — Error != Zero"
            )
        elif obs.outcome_type == "AMBIGUOUS" or obs.observed_amount is None:
            task.search_status = "AMBIGUOUS"
            task.zain_observed_amount = None
            task.raw_observed_screen_text = obs.raw_screen_text
            task.error_or_review_details = obs.error_or_ambiguity_reason or "Ambiguous screen or identity mismatch"
            task.reconciliation_status = ReconciliationStatus.AMBIGUOUS_RESULT_NEEDS_REVIEW
            task.reconciliation_variance = None
            task.reconciliation_note = f"AMBIGUOUS_RESULT_NEEDS_REVIEW ({task.error_or_review_details})"
        else:
            # Confirmed Service-Level Amount (including explicit 0.00!)
            zain_amt = obs.observed_amount.quantize(Decimal("0.01"))
            file_amt = (task.file_comparison_amount or Decimal("0.00")).quantize(Decimal("0.01"))
            variance = file_amt - zain_amt

            task.search_status = "SUCCESS"
            task.zain_observed_amount = zain_amt
            task.raw_observed_screen_text = obs.raw_screen_text
            task.reconciliation_variance = variance

            if variance == Decimal("0.00"):
                task.reconciliation_status = ReconciliationStatus.SERVICE_MATCH
                task.reconciliation_note = "SERVICE_MATCH (Service-to-Service Exact Match)"
            else:
                task.reconciliation_status = ReconciliationStatus.SERVICE_AMOUNT_DIFFERENCE
                task.reconciliation_note = (
                    f"SERVICE_AMOUNT_DIFFERENCE: {format_variance_description(file_amt, zain_amt)}"
                )

        # Propagate 1:1 to the linked SourceRow
        for uid in task.linked_row_uids:
            row = row_map.get(uid)
            if not row:
                continue
            row.reconciliation_status = task.reconciliation_status
            row.zain_service_amount = task.zain_observed_amount
            row.service_variance = task.reconciliation_variance
            row.reconciliation_note = task.reconciliation_note

    # 2. Process Account Search Tasks (N:1 Account-Level Scope — Zero Residual Deduction!)
    for task in plan.account_search_tasks:
        obs = observations.get(task.task_id)
        if obs is None:
            continue

        grp = acc_map.get(task.account_number or "")

        # Law #4: Error != Zero
        if obs.outcome_type == "SEARCH_ERROR":
            task.search_status = "SEARCH_ERROR"
            task.zain_observed_amount = None
            task.raw_observed_screen_text = obs.raw_screen_text
            task.error_or_review_details = obs.error_or_ambiguity_reason or "Search / DOM Timeout Error"
            task.reconciliation_status = ReconciliationStatus.SEARCH_ERROR_NEEDS_REVIEW
            task.reconciliation_variance = None
            task.reconciliation_note = (
                f"SEARCH_ERROR_NEEDS_REVIEW ({task.error_or_review_details}) — Error != Zero"
            )
            if grp:
                grp.zain_account_total = None
                grp.account_reconciliation_status = ReconciliationStatus.SEARCH_ERROR_NEEDS_REVIEW
                grp.account_variance = None
                grp.account_reconciliation_note = task.reconciliation_note

        elif obs.outcome_type == "AMBIGUOUS" or obs.observed_amount is None:
            task.search_status = "AMBIGUOUS"
            task.zain_observed_amount = None
            task.raw_observed_screen_text = obs.raw_screen_text
            task.error_or_review_details = obs.error_or_ambiguity_reason or "Ambiguous Account Result"
            task.reconciliation_status = ReconciliationStatus.AMBIGUOUS_RESULT_NEEDS_REVIEW
            task.reconciliation_variance = None
            task.reconciliation_note = f"AMBIGUOUS_RESULT_NEEDS_REVIEW ({task.error_or_review_details})"
            if grp:
                grp.zain_account_total = None
                grp.account_reconciliation_status = ReconciliationStatus.AMBIGUOUS_RESULT_NEEDS_REVIEW
                grp.account_variance = None
                grp.account_reconciliation_note = task.reconciliation_note

        else:
            zain_acc_total = obs.observed_amount.quantize(Decimal("0.01"))
            task.search_status = "SUCCESS"
            task.zain_observed_amount = zain_acc_total
            task.raw_observed_screen_text = obs.raw_screen_text

            if grp:
                grp.zain_account_total = zain_acc_total

                # Law #1 (Strict Scope Symmetry) & Law #2 (Zero Residual Deduction):
                if grp.file_scope_status.is_complete:
                    file_acc_total = grp.included_total_sum.quantize(Decimal("0.01"))
                    acc_variance = file_acc_total - zain_acc_total
                    task.reconciliation_variance = acc_variance
                    grp.account_variance = acc_variance

                    if acc_variance == Decimal("0.00"):
                        task.reconciliation_status = ReconciliationStatus.ACCOUNT_TOTAL_MATCH
                        grp.account_reconciliation_status = ReconciliationStatus.ACCOUNT_TOTAL_MATCH
                        note = (
                            f"ACCOUNT_TOTAL_MATCH (Complete File Scope): "
                            f"File Account Total ({file_acc_total:.2f}) == Zain Account Total ({zain_acc_total:.2f}). "
                            f"[Zero Residual Deduction Enforced]"
                        )
                    else:
                        task.reconciliation_status = ReconciliationStatus.ACCOUNT_TOTAL_DIFFERENCE
                        grp.account_reconciliation_status = ReconciliationStatus.ACCOUNT_TOTAL_DIFFERENCE
                        note = (
                            f"ACCOUNT_TOTAL_DIFFERENCE: {format_variance_description(file_acc_total, zain_acc_total)} "
                            f"(File Account Total={file_acc_total:.2f}, Zain Account Total={zain_acc_total:.2f}). "
                            f"[Not Single Service Amount]"
                        )
                    task.reconciliation_note = note
                    grp.account_reconciliation_note = note
                else:
                    # Partial / Incomplete File Scope -> Direct Automatic Comparison Locked!
                    task.reconciliation_status = ReconciliationStatus.SCOPE_MISMATCH_WARNING_LOCKED
                    task.reconciliation_variance = None
                    grp.account_reconciliation_status = ReconciliationStatus.SCOPE_MISMATCH_WARNING_LOCKED
                    grp.account_variance = None

                    zero_diag = (
                        " | Zain Account Total = 0.00 (Zero Account Balance in Zain)"
                        if zain_acc_total == Decimal("0.00")
                        else ""
                    )
                    note = (
                        "SCOPE_MISMATCH_WARNING: Partial/Incomplete File Scope — Direct Comparison Locked | "
                        f"Zain Account Total={zain_acc_total:.2f}, "
                        f"File Included Sum={grp.included_total_sum:.2f}, "
                        f"File All Known Rows Sum={grp.all_known_file_rows_sum:.2f}, "
                        f"Reasons={grp.scope_incompleteness_reasons}{zero_diag}"
                    )
                    task.reconciliation_note = note
                    grp.account_reconciliation_note = note

        # Propagate Account-Level Context to each linked Non-2 SourceRow WITHOUT assigning Zain Account Total as a Service Amount!
        for uid in task.linked_row_uids:
            row = row_map.get(uid)
            if not row:
                continue
            row.zain_service_amount = None  # Law #1 & Law #2: Never assign Account Total as Service Amount!
            row.zain_account_total_context = task.zain_observed_amount
            row.service_variance = None
            row.reconciliation_status = task.reconciliation_status
            row.reconciliation_note = (
                f"Zain Account Total (Not Single Service Amount): "
                f"{task.zain_observed_amount if task.zain_observed_amount is not None else 'N/A'} | "
                f"{task.reconciliation_note}"
            )

    return plan


# ============================================================================
# Multi-Sheet Excel Audit & Reconciliation Report Exporter
# ============================================================================

def _style_sheet_header(sheet, headers: list[str], header_color: str = "1F4E78") -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    fill = PatternFill("solid", fgColor=header_color)
    font = Font(bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, text in enumerate(headers, start=1):
        cell = sheet.cell(1, idx, text)
        cell.fill = fill
        cell.font = font
        cell.alignment = align


def _auto_fit_sheet(sheet, col_count: int) -> None:
    if sheet.max_row >= 1 and col_count >= 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(col_count)}{sheet.max_row}"
    for col_idx in range(1, col_count + 1):
        col_letter = get_column_letter(col_idx)
        sheet.column_dimensions[col_letter].width = 24
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=col_count):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_comprehensive_audit_workbook(plan: PreZainSearchPlan, output_path: Path) -> Path:
    """Exports the complete 7-sheet Audit & Reconciliation Workbook to Excel."""
    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    # ------------------------------------------------------------------------
    # Sheet 1: الفروقات (Confirmed Same-Scope Variances - Primary Sheet)
    # ------------------------------------------------------------------------
    ws_diff = wb.active
    ws_diff.title = "الفروقات"
    diff_headers = [
        "نطاق المقارنة (Scope)",
        "معرف المهمة (Task ID)",
        "رقم البحث في زين",
        "رقم الحساب",
        "الخدمة / الخدمات المرتبطة",
        "موقع الصف (Sheet!Row)",
        "اسم العميل",
        "المبلغ في الملف (نفس النطاق)",
        "المبلغ في موقع زين (نفس النطاق)",
        "الفرق المالي (File - Zain)",
        "حالة المطابقة",
        "التوصيف المحاسبي الصارم (بدون تخمين سبب جذري)",
    ]
    _style_sheet_header(ws_diff, diff_headers, "1F4E78")

    row_map = {r.row_uid: r for r in plan.source_rows}
    acc_map = {g.account_number: g for g in plan.account_groups}

    diff_row_idx = 2
    for task in plan.service_search_tasks:
        if task.reconciliation_status == ReconciliationStatus.SERVICE_AMOUNT_DIFFERENCE:
            r = row_map.get(task.linked_row_uids[0])
            ws_diff.append(
                [
                    "SERVICE_LEVEL (خدمة مقابل خدمة)",
                    task.task_id,
                    task.search_key,
                    task.account_number or "يتيم (Orphan)",
                    task.service_number or "",
                    ", ".join(task.linked_row_uids),
                    r.customer_name_raw if r else "",
                    float(task.file_comparison_amount or 0),
                    float(task.zain_observed_amount or 0),
                    float(task.reconciliation_variance or 0),
                    task.reconciliation_status.value,
                    task.reconciliation_note,
                ]
            )
            ws_diff.cell(diff_row_idx, 3).number_format = "@"
            ws_diff.cell(diff_row_idx, 4).number_format = "@"
            ws_diff.cell(diff_row_idx, 5).number_format = "@"
            diff_row_idx += 1

    for task in plan.account_search_tasks:
        if task.reconciliation_status == ReconciliationStatus.ACCOUNT_TOTAL_DIFFERENCE:
            grp = acc_map.get(task.account_number or "")
            ws_diff.append(
                [
                    "ACCOUNT_LEVEL_TOTAL (إجمالي حساب مكتمل النطاق بالملف)",
                    task.task_id,
                    task.search_key,
                    task.account_number or "",
                    ", ".join(task.all_account_services_in_file),
                    ", ".join(grp.source_row_uids if grp else task.linked_row_uids),
                    ", ".join(grp.distinct_customer_names) if grp else "",
                    float(task.file_all_included_sum or 0),
                    float(task.zain_observed_amount or 0),
                    float(task.reconciliation_variance or 0),
                    task.reconciliation_status.value,
                    task.reconciliation_note,
                ]
            )
            ws_diff.cell(diff_row_idx, 3).number_format = "@"
            ws_diff.cell(diff_row_idx, 4).number_format = "@"
            diff_row_idx += 1
    _auto_fit_sheet(ws_diff, len(diff_headers))

    # ------------------------------------------------------------------------
    # Sheet 2: الأخطاء (Validation Errors, Duplicates, Conflicts, Search Errors & Scope Locks)
    # ------------------------------------------------------------------------
    ws_err = wb.create_sheet("الأخطاء")
    err_headers = [
        "تصنيف الحالة",
        "موقع الصف (Sheet!Row)",
        "رقم الحساب الخام / المطبع",
        "رقم الخدمة الخام / المطبع",
        "اسم العميل",
        "المبلغ الخام",
        "حالة المبلغ (Amount State)",
        "أكواد قواعد التحقق",
        "تفاصيل الخطأ أو المراجعة أو قفل النطاق",
    ]
    _style_sheet_header(ws_err, err_headers, "8B0000")

    for r in plan.source_rows:
        if r.routing_disposition == "BLOCKED_REVIEW" or r.reconciliation_status in (
            ReconciliationStatus.SEARCH_ERROR_NEEDS_REVIEW,
            ReconciliationStatus.AMBIGUOUS_RESULT_NEEDS_REVIEW,
            ReconciliationStatus.SCOPE_MISMATCH_WARNING_LOCKED,
        ):
            details = "; ".join(r.validation_errors) if r.validation_errors else r.reconciliation_note
            ws_err.append(
                [
                    r.validation_status.value
                    if r.routing_disposition == "BLOCKED_REVIEW"
                    else r.reconciliation_status.value,
                    r.row_uid,
                    f"{r.account_number_raw or ''} -> {r.account_number_norm or 'None'}",
                    f"{r.service_number_raw or ''} -> {r.service_number_norm or 'None'}",
                    r.customer_name_raw or "",
                    str(r.file_amount_raw) if r.file_amount_raw is not None else "[Blank]",
                    r.amount_state.value,
                    ", ".join(r.validation_rule_codes),
                    details,
                ]
            )
    _auto_fit_sheet(ws_err, len(err_headers))

    # ------------------------------------------------------------------------
    # Sheet 3: ملخص الجاهزية والاتزان (Pre-Zain Readiness Summary)
    # ------------------------------------------------------------------------
    ws_sum = wb.create_sheet("ملخص الجاهزية والاتزان")
    sum_headers = ["المؤشر الهندسي (Metric)", "القيمة (Value)", "التفسير والحكم (Audit Verification)"]
    _style_sheet_header(ws_sum, sum_headers, "2E5A88")
    summary_rows = [
        ("اسم الملف المصدر", plan.source_file_name, f"SHA-256: {plan.source_file_sha256}"),
        (
            "إجمالي الصفوف المستلمة (Total Input Rows)",
            plan.total_input_rows,
            "جميع صفوف البيانات في كافة الأوراق الداخلة",
        ),
        (
            "صفوف موجهة لبحث الخدمة (Service-Searchable Rows)",
            plan.service_searchable_rows_count,
            f"أنتجت {len(plan.service_search_tasks)} مهمة بحث بالخدمة (1:1 Scope)",
        ),
        (
            "صفوف موجهة لبحث الحساب (Account-Searchable Rows)",
            plan.account_searchable_rows_count,
            f"بعد منع التكرار (Deduplication) أنتجت {len(plan.account_search_tasks)} مهمة بحث حساب فقط",
        ),
        (
            "صفوف مستبعدة محفوظة السياق (Excluded Rows Preserved)",
            plan.excluded_rows_count,
            "لم تُحذف من الذاكرة؛ حُفظت داخل أوعية الحسابات لحماية اكتمال النطاق",
        ),
        (
            "صفوف محجوزة للمراجعة / الأخطاء (Blocked / Review Rows)",
            plan.blocked_or_review_rows_count,
            "تشمل المبالغ الفارغة، الصيغ العلمية التالفة، وتكرارات/تعارضات الخدمات",
        ),
        (
            "معادلة الاتزان (Readiness Invariant Check)",
            "100% BALANCED" if plan.invariant_balanced else "FAILED",
            f"{plan.service_searchable_rows_count} + {plan.account_searchable_rows_count} + "
            f"{plan.excluded_rows_count} + {plan.blocked_or_review_rows_count} = {plan.total_input_rows}",
        ),
        (
            "إجمالي طلبات البحث المخططة لزين (Total Planned Zain Searches)",
            len(plan.service_search_tasks) + len(plan.account_search_tasks),
            f"{len(plan.service_search_tasks)} Service Tasks + {len(plan.account_search_tasks)} Account Tasks",
        ),
        (
            "قفل الجاهزية (READY_FOR_ZAIN_EXECUTION)",
            str(plan.ready_for_zain_execution).upper(),
            "شهادة الجاهزية قبل الاتصال بنظام زين",
        ),
    ]
    for metric, val, desc in summary_rows:
        ws_sum.append([metric, val, desc])
    _auto_fit_sheet(ws_sum, len(sum_headers))

    # ------------------------------------------------------------------------
    # Sheet 4: مهام البحث بالخدمة (Service Search Tasks)
    # ------------------------------------------------------------------------
    ws_srv = wb.create_sheet("مهام البحث بالخدمة")
    srv_headers = [
        "Task ID",
        "Search Type",
        "Search Key (Service)",
        "Account Number",
        "Linked Row",
        "File Amount (Service Scope)",
        "Expected Result Scope",
        "Zain Observed Amount",
        "Variance",
        "Reconciliation Status",
        "Creation / Reconciliation Note",
    ]
    _style_sheet_header(ws_srv, srv_headers, "1F6B48")
    for t in plan.service_search_tasks:
        ws_srv.append(
            [
                t.task_id,
                t.search_type,
                t.search_key,
                t.account_number or "[None - Orphan Service]",
                ", ".join(t.linked_row_uids),
                float(t.file_comparison_amount) if t.file_comparison_amount is not None else None,
                t.expected_result_scope,
                float(t.zain_observed_amount) if t.zain_observed_amount is not None else None,
                float(t.reconciliation_variance) if t.reconciliation_variance is not None else None,
                t.reconciliation_status.value,
                t.reconciliation_note or t.creation_reason,
            ]
        )
    _auto_fit_sheet(ws_srv, len(srv_headers))

    # ------------------------------------------------------------------------
    # Sheet 5: مهام البحث بالحساب (Deduplicated Account Search Tasks)
    # ------------------------------------------------------------------------
    ws_acc = wb.create_sheet("مهام البحث بالحساب")
    acc_headers = [
        "Task ID",
        "Search Type",
        "Search Key (Account)",
        "Triggering Non-2 Services",
        "Linked Rows (Cross-Sheet)",
        "Included Non-2 Sum",
        "Total Account File Sum (All Included)",
        "File All Known Rows Sum (Incl. Excluded)",
        "File Scope Status",
        "Zain Account Total",
        "Account Variance",
        "Reconciliation Status",
        "Audit Note",
    ]
    _style_sheet_header(ws_acc, acc_headers, "5B3A82")
    for t in plan.account_search_tasks:
        ws_acc.append(
            [
                t.task_id,
                t.search_type,
                t.search_key,
                ", ".join(t.triggering_services),
                ", ".join(t.linked_row_uids),
                float(t.file_included_non2_sum or 0),
                float(t.file_all_included_sum or 0),
                float(t.file_all_known_rows_sum or 0),
                t.account_file_scope_status.value if t.account_file_scope_status else "",
                float(t.zain_observed_amount) if t.zain_observed_amount is not None else None,
                float(t.reconciliation_variance) if t.reconciliation_variance is not None else None,
                t.reconciliation_status.value,
                t.reconciliation_note or t.creation_reason,
            ]
        )
    _auto_fit_sheet(ws_acc, len(acc_headers))

    # ------------------------------------------------------------------------
    # Sheet 6: تجميع الحسابات والنطاق (Account Groups & 9 Concepts)
    # ------------------------------------------------------------------------
    ws_grp = wb.create_sheet("تجميع الحسابات والنطاق")
    grp_headers = [
        "Account Number",
        "Sheets",
        "All Linked Rows",
        "Starts-With-2 Rows",
        "Non-2 Rows",
        "Excluded Rows",
        "Invalid / Conflict Rows",
        "Included Starts-w-2 Sum",
        "Included Non-2 Sum",
        "Included Total Sum",
        "All Known File Rows Sum",
        "File Scope Status",
        "9 Engineering Concept Tags",
        "Requires Service Search?",
        "Requires Account Search?",
    ]
    _style_sheet_header(ws_grp, grp_headers, "3A5F6F")
    for g in plan.account_groups:
        ws_grp.append(
            [
                g.account_number,
                ", ".join(g.distinct_sheets),
                ", ".join(g.source_row_uids),
                f"{len(g.included_starts_with_2_rows)} ({', '.join(g.included_starts_with_2_rows)})",
                f"{len(g.included_non_2_rows)} ({', '.join(g.included_non_2_rows)})",
                f"{len(g.excluded_service_rows)} ({', '.join(g.excluded_service_rows)})",
                f"{len(g.invalid_or_conflict_rows)} ({', '.join(g.invalid_or_conflict_rows)})",
                float(g.included_starts_with_2_sum),
                float(g.included_non_2_sum),
                float(g.included_total_sum),
                float(g.all_known_file_rows_sum),
                g.file_scope_status.value,
                ", ".join(c.value for c in g.concept_tags),
                "Yes" if g.requires_service_searches else "No",
                "Yes" if g.requires_account_search else "No",
            ]
        )
    _auto_fit_sheet(ws_grp, len(grp_headers))

    # ------------------------------------------------------------------------
    # Sheet 7: سجل الصفوف والسلالة (Complete Source Row Lineage)
    # ------------------------------------------------------------------------
    ws_rows = wb.create_sheet("سجل الصفوف والسلالة")
    row_headers = [
        "Row UID",
        "Sheet",
        "Excel Row",
        "Account Raw -> Norm",
        "Service Raw -> Norm",
        "Starts w/ 2?",
        "Amount Raw -> Decimal",
        "Amount State",
        "Status Raw -> Norm",
        "Include in Scope?",
        "Validation Status",
        "Routing Disposition",
        "Linked Task ID",
        "Normalization & Audit Notes",
    ]
    _style_sheet_header(ws_rows, row_headers, "444444")
    for r in plan.source_rows:
        ws_rows.append(
            [
                r.row_uid,
                r.source_sheet,
                r.source_row,
                f"{r.account_number_raw or '[Blank]'} -> {r.account_number_norm or 'None'}",
                f"{r.service_number_raw or '[Blank]'} -> {r.service_number_norm or 'None'}",
                str(r.service_starts_with_2),
                f"{r.file_amount_raw if r.file_amount_raw is not None else '[Blank]'} -> "
                f"{f'{r.file_amount_decimal:.2f}' if r.file_amount_decimal is not None else 'None'}",
                r.amount_state.value,
                f"{r.main_status_raw or ''} -> {r.main_status_norm or ''}",
                str(r.include_in_scope),
                r.validation_status.value,
                r.routing_disposition,
                r.linked_search_task_id or "None",
                "; ".join(r.normalization_notes + r.validation_errors + r.exclusion_reasons),
            ]
        )
    _auto_fit_sheet(ws_rows, len(row_headers))

    wb.save(output_path)
    wb.close()
    return output_path
