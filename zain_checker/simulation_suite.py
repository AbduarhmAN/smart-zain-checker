"""Full End-to-End Simulation Suite (Section 11 / Task 10) for Smart Zain Checker.

Builds the exact 20-row, 2-sheet ('Sheet1_Main' & 'Sheet2_Branch') complex simulation
workbook specified in Section 11, executes all 16 pipeline stages, verifies every
mathematical invariant and table assertion, and exports the complete audit package.
"""

from __future__ import annotations

from decimal import Decimal
import io
import json
from pathlib import Path
import sys
from typing import Dict

from openpyxl import Workbook

from .domain_models import (
    AccountFileScopeStatus,
    AmountState,
    PreZainSearchPlan,
    ReconciliationStatus,
    RowValidationStatus,
)
from .pipeline import run_pre_zain_pipeline
from .reconciliation import (
    ZainTaskObservation,
    bind_zain_observations_to_plan,
    export_comprehensive_audit_workbook,
)


def create_section_11_simulation_workbook(target_path: Path) -> Path:
    """Creates the exact 20-row, 2-sheet Excel workbook from Section 11 (Task 10)."""
    target_path = Path(target_path).resolve()
    target_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # ------------------------------------------------------------------------
    # Sheet 1: Sheet1_Main (15 data rows: R2..R16)
    # ------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Sheet1_Main"
    headers = [
        "اسم المحصل",
        "اسم العميل",
        "رقم الحساب",
        "رقم الخدمة",
        "مبلغ المديونية في ملف المحصل",
        "الحالة الرئيسية",
    ]
    ws1.append(headers)

    sheet1_rows = [
        # R2: Account 1001 - single service starting with 2
        ["سعد الحربي", "شركة المدار", "1001", "2001001", 400.00, "نشط"],
        # R3 & R4: Account 1002 - two services both starting with 2
        ["سعد الحربي", "مؤسسة النخيل", "1002", "2002001", 250.00, "نشط"],
        ["سعد الحربي", "مؤسسة النخيل", "1002", "2002002", 350.00, "نشط"],
        # R5 & R6: Account 1003 - two services NOT starting with 2 (plus 3rd in Sheet2!)
        ["سعد الحربي", "خالد العتيبي", "1003", "5503001", 600.00, "نشط"],
        ["سعد الحربي", "خالد العتيبي", "1003", "5503002", 400.00, "نشط"],
        # R7 & R8: Account 1004 (Mixed) - one starting with 2, one non-2 with leading zero
        ["سعد الحربي", "فهد الدوسري", "1004", "2004001", 300.00, "نشط"],
        ["سعد الحربي", "فهد الدوسري", "1004", "0594002", 700.00, "نشط"],
        # R9 & R10: Account 1005 - one active non-2, one excluded by status ('مغلق نهائيا')
        ["سعد الحربي", "ناصر السبيعي", "1005", "5505001", 500.00, "نشط"],
        ["سعد الحربي", "ناصر السبيعي", "1005", "5505002", 200.00, "مغلق نهائيا"],
        # R11 & R12: Account 1006 - one starting with 2 and explicit zero (0), one non-2 (850.00)
        ["سعد الحربي", "عمر الغامدي", "1006", "2006001", 0, "نشط"],
        ["سعد الحربي", "عمر الغامدي", "1006", "5506002", 850.00, "نشط"],
        # R13 & R14: Account 1007 - one valid non-2 (450.00), one with Blank amount
        ["سعد الحربي", "سلطان الشهري", "1007", "5507001", 450.00, "نشط"],
        ["سعد الحربي", "سلطان الشهري", "1007", "5507002", None, "نشط"],
        # R15 & R16: Account 1008 - duplicate service 2008001 with same amount (320.00)
        ["سعد الحربي", "ماجد القحطاني", "1008", "2008001", 320.00, "نشط"],
        ["سعد الحربي", "ماجد القحطاني", "1008", "2008001", 320.00, "نشط"],
    ]
    for r_vals in sheet1_rows:
        ws1.append(r_vals)
        # Ensure identifiers are stored as text to preserve leading zero in R8 ("0594002")
        cur_row = ws1.max_row
        ws1.cell(cur_row, 3).number_format = "@"
        ws1.cell(cur_row, 4).number_format = "@"

    # ------------------------------------------------------------------------
    # Sheet 2: Sheet2_Branch (5 data rows: R2..R6)
    # ------------------------------------------------------------------------
    ws2 = wb.create_sheet("Sheet2_Branch")
    ws2.append(headers)

    sheet2_rows = [
        # R2: Belongs to Account 1003 in Sheet1! (Arabic-Indic digits + .0 suffix + '250.00 SAR' + tatweel)
        ["سعد الحربي", "خالد عبدالله العتيبي", " ١٠٠٣ ", "5503003.0", "250.00 SAR", " نشـط "],
        # R3: Missing account + valid service starting with 2 -> Searchable Orphan Service!
        ["سعد الحربي", "وليد الشمري", None, "2009001", 180.00, "نشط"],
        # R4: Missing account + non-2 service -> Blocked / Error
        ["سعد الحربي", "تركي العنزي", None, "5509002", 290.00, "نشط"],
        # R5: Account 1010 + missing service -> Blocked / Error
        ["سعد الحربي", "بندر المالكي", "1010", None, 510.00, "نشط"],
        # R6: Account 1011 + corrupted scientific notation service -> Blocked / Error
        ["سعد الحربي", "ريان الزهراني", "1011", "2.011E+11", 900.00, "نشط"],
    ]
    for r_vals in sheet2_rows:
        ws2.append(r_vals)
        cur_row = ws2.max_row
        ws2.cell(cur_row, 3).number_format = "@"
        ws2.cell(cur_row, 4).number_format = "@"

    wb.save(target_path)
    wb.close()
    return target_path


def run_and_verify_section_11_simulation(output_dir: Path) -> PreZainSearchPlan:
    """Runs the 20-row simulation end-to-end and asserts every invariant from Section 11."""
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    sim_xlsx = create_section_11_simulation_workbook(output_dir / "simulation_input_20rows.xlsx")
    plan = run_pre_zain_pipeline(
        workbook_path=sim_xlsx,
        target_collector="سعد الحربي",
        all_collectors=False,
    )

    # 1. Verify Stage 9 Summary Invariants
    assert plan.total_input_rows == 20, f"Expected 20 input rows, got {plan.total_input_rows}"
    assert plan.service_searchable_rows_count == 6, (
        f"Expected 6 service-searchable rows, got {plan.service_searchable_rows_count}"
    )
    assert plan.account_searchable_rows_count == 7, (
        f"Expected 7 account-searchable rows, got {plan.account_searchable_rows_count}"
    )
    assert plan.excluded_rows_count == 1, f"Expected 1 excluded row, got {plan.excluded_rows_count}"
    assert plan.blocked_or_review_rows_count == 6, (
        f"Expected 6 blocked/review rows, got {plan.blocked_or_review_rows_count}"
    )
    assert len(plan.service_search_tasks) == 6, (
        f"Expected 6 Service Search Tasks, got {len(plan.service_search_tasks)}"
    )
    assert len(plan.account_search_tasks) == 5, (
        f"Expected 5 Deduplicated Account Search Tasks, got {len(plan.account_search_tasks)}"
    )
    assert plan.invariant_balanced is True
    assert plan.ready_for_zain_execution is True

    # 2. Verify Cross-Sheet Deduplication on Account 1003 (ACC-001)
    acc_1003_task = next(t for t in plan.account_search_tasks if t.search_key == "1003")
    assert acc_1003_task.task_id == "ACC-001"
    assert acc_1003_task.triggering_services == ["5503001", "5503002", "5503003"]
    assert acc_1003_task.linked_row_uids == [
        "Sheet1_Main!R5",
        "Sheet1_Main!R6",
        "Sheet2_Branch!R2",
    ]
    assert acc_1003_task.file_included_non2_sum == Decimal("1250.00")
    assert acc_1003_task.file_all_included_sum == Decimal("1250.00")
    assert acc_1003_task.account_file_scope_status == AccountFileScopeStatus.FILE_SCOPE_COMPLETE

    # 3. Verify Mixed Account 1004 (SRV-004 + ACC-002 with leading zero service '0594002')
    acc_1004_task = next(t for t in plan.account_search_tasks if t.search_key == "1004")
    assert acc_1004_task.triggering_services == ["0594002"]
    assert acc_1004_task.file_included_non2_sum == Decimal("700.00")
    assert acc_1004_task.file_all_included_sum == Decimal("1000.00")
    assert acc_1004_task.account_file_scope_status == AccountFileScopeStatus.FILE_SCOPE_COMPLETE

    # 4. Verify Partial Account 1005 (ACC-003 with excluded row S1!R10 preserved in context)
    acc_1005_task = next(t for t in plan.account_search_tasks if t.search_key == "1005")
    assert acc_1005_task.file_included_non2_sum == Decimal("500.00")
    assert acc_1005_task.file_all_known_rows_sum == Decimal("700.00")
    assert (
        acc_1005_task.account_file_scope_status
        == AccountFileScopeStatus.FILE_SCOPE_INCOMPLETE_PARTIAL_EXCLUSION
    )

    # 5. Verify Explicit Zero Service in Account 1006 (SRV-005 @ 0.00 + ACC-004 @ 850.00)
    srv_005 = next(t for t in plan.service_search_tasks if t.search_key == "2006001")
    assert srv_005.file_comparison_amount == Decimal("0.00")
    acc_1006_task = next(t for t in plan.account_search_tasks if t.search_key == "1006")
    assert acc_1006_task.file_all_included_sum == Decimal("850.00")
    assert acc_1006_task.account_file_scope_status == AccountFileScopeStatus.FILE_SCOPE_COMPLETE

    # 6. Verify Blank Sibling in Account 1007 (ACC-005 marked INCOMPLETE_HAS_ERRORS)
    acc_1007_task = next(t for t in plan.account_search_tasks if t.search_key == "1007")
    assert (
        acc_1007_task.account_file_scope_status
        == AccountFileScopeStatus.FILE_SCOPE_INCOMPLETE_HAS_ERRORS
    )
    row_r14 = next(r for r in plan.source_rows if r.row_uid == "Sheet1_Main!R14")
    assert row_r14.amount_state == AmountState.MISSING_BLANK
    assert row_r14.file_amount_decimal is None  # Never converted to 0.00!

    # 7. Verify Duplicate Service in Account 1008 (Held, never summed to 640.00)
    acc_1008_grp = next(g for g in plan.account_groups if g.account_number == "1008")
    assert acc_1008_grp.included_total_sum == Decimal("0.00")
    assert (
        acc_1008_grp.file_scope_status
        == AccountFileScopeStatus.FILE_SCOPE_INCOMPLETE_DUPLICATE_CONFLICT
    )

    # 8. Verify Orphan Service Starting with 2 in Sheet2_Branch!R3 (SRV-006)
    srv_006 = next(t for t in plan.service_search_tasks if t.search_key == "2009001")
    assert srv_006.account_number is None
    assert srv_006.file_comparison_amount == Decimal("180.00")

    # 9. Simulate Post-Zain Observations & Verify Scope-Safe Reconciliation (Stage 16)
    simulated_observations: Dict[str, ZainTaskObservation] = {
        "SRV-001": ZainTaskObservation("SRV-001", "SUCCESS", Decimal("400.00")),  # Match
        "SRV-002": ZainTaskObservation("SRV-002", "SUCCESS", Decimal("200.00")),  # File(250) > Zain(200)
        "SRV-003": ZainTaskObservation("SRV-003", "SUCCESS", Decimal("350.00")),  # Match
        "SRV-004": ZainTaskObservation("SRV-004", "SUCCESS", Decimal("300.00")),  # Match on Service 2004001
        "SRV-005": ZainTaskObservation("SRV-005", "SUCCESS", Decimal("150.00")),  # File(0) < Zain(150)
        "SRV-006": ZainTaskObservation("SRV-006", "SEARCH_ERROR", None, error_or_ambiguity_reason="DOM Timeout"),
        "ACC-001": ZainTaskObservation("ACC-001", "SUCCESS", Decimal("1250.00")),  # Complete Account Match
        "ACC-002": ZainTaskObservation("ACC-002", "SUCCESS", Decimal("1000.00")),  # Complete Mixed Account Match
        "ACC-003": ZainTaskObservation("ACC-003", "SUCCESS", Decimal("700.00")),   # Partial Scope Locked!
        "ACC-004": ZainTaskObservation("ACC-004", "SUCCESS", Decimal("1000.00")),  # Account Diff (850 vs 1000)
        "ACC-005": ZainTaskObservation("ACC-005", "SUCCESS", Decimal("900.00")),   # Incomplete Blank Sibling Locked!
    }
    bind_zain_observations_to_plan(plan, simulated_observations)

    # Verify Law #4 (Error != Zero on SRV-006)
    assert srv_006.zain_observed_amount is None
    assert srv_006.reconciliation_status == ReconciliationStatus.SEARCH_ERROR_NEEDS_REVIEW

    # Verify Law #1 & #2 on ACC-003 (Partial Account 1005 is locked from false 500 vs 700 variance)
    assert acc_1005_task.reconciliation_status == ReconciliationStatus.SCOPE_MISMATCH_WARNING_LOCKED
    assert acc_1005_task.reconciliation_variance is None

    # Export JSON and Excel Audit Workbook
    json_path = output_dir / "simulation_plan_and_reconciliation.json"
    json_path.write_text(json.dumps(plan.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
    export_comprehensive_audit_workbook(plan, output_dir / "simulation_audit_report.xlsx")

    return plan


if __name__ == "__main__":
    if hasattr(sys.stdout, "buffer"):
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    out_folder = Path(__file__).resolve().parent.parent / "simulation_outputs"
    completed_plan = run_and_verify_section_11_simulation(out_folder)
    print(
        json.dumps(
            {
                "status": "SIMULATION_VERIFIED_100_PERCENT",
                "output_directory": str(out_folder),
                "readiness_summary": completed_plan.to_dict()["readiness_summary"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
