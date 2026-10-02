from __future__ import annotations

import io
import json
import os
import sys
import threading
import time
import webbrowser
import zipfile
import xml.etree.ElementTree as ET
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

# Ensure UTF-8 stdout on Windows safely without closing underlying buffer (User Global Rule)
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
elif hasattr(sys.stdout, "buffer") and getattr(sys.stdout, "encoding", "").lower() != "utf-8":
    _PREV_STDOUT = sys.stdout
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

_PROJECT_ROOT = str(Path(__file__).resolve().parent.parent)
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

# Enforce Multi-Factor Hardware Lock & Anti-Debug Check
from zain_checker.security import enforce_hardware_lock
enforce_hardware_lock()

from zain_checker.config import (
    PROJECT_DIRECTORY as PROJECT_DIR,
    BUNDLE_DIRECTORY,
)
STATIC_DIR = BUNDLE_DIRECTORY / "web_ui_prototype" / "static"
PORT = 5050


def fast_read_sheet_names(xlsx_path: Path) -> list[str]:
    """Ultra-fast (2ms) Excel sheet name extraction without loading rows into RAM."""
    try:
        with zipfile.ZipFile(xlsx_path, "r") as z:
            with z.open("xl/workbook.xml") as f:
                tree = ET.parse(f)
                sheets = [
                    node.attrib["name"]
                    for node in tree.findall(
                        ".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"
                    )
                ]
                return sheets
    except Exception:
        return []


def inspect_smart_columns_for_workbook(xlsx_path: Path) -> list[dict[str, Any]]:
    """Scans Row 1 of each sheet in xlsx_path using detect_smart_row1_columns and returns
    the detected column letters per sheet for automatic UI and backend assignment.
    """
    from openpyxl import load_workbook
    from zain_checker.workbook import detect_smart_row1_columns

    result: list[dict[str, Any]] = []
    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                info = detect_smart_row1_columns(ws)
                result.append(info["letters"])
        finally:
            wb.close()
    except Exception:
        pass
    return result


def inspect_full_schema_for_sheet(xlsx_path: Path, sheet_index: int = 0) -> dict[str, Any]:
    """Inspects Row 1 of the given sheet and provides full schema, column headers,
    and automatic mode support detection (smart_hybrid, account_only, service_only).
    """
    from openpyxl import load_workbook
    from zain_checker.workbook import detect_sheet_schema_and_modes

    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        try:
            if 0 <= sheet_index < len(wb.worksheets):
                ws = wb.worksheets[sheet_index]
            else:
                ws = wb.worksheets[0]
            analysis = detect_sheet_schema_and_modes(ws)
            analysis["sheet_name"] = ws.title
            analysis["total_sheets"] = len(wb.worksheets)
            return analysis
        finally:
            wb.close()
    except Exception as exc:
        return {
            "error": str(exc),
            "columns": [],
            "supported_modes": {"smart_hybrid": False, "account_only": False, "service_only": False},
            "recommended_mode": "custom",
            "missing_for_modes": {},
            "amount_options": [],
        }



def read_checkpoint_info() -> dict[str, Any]:
    """Reads .zain-checkpoint.json in ~1ms without blocking, with auto-fallback to backup."""
    primary = PROJECT_DIR / ".zain-checkpoint.json"
    backup = PROJECT_DIR / ".zain-checkpoint.json.bak"
    for path in (primary, backup):
        if not path.exists():
            continue
        try:
            raw = path.read_text(encoding="utf-8").replace("\x00", "").strip()
            if not raw:
                continue
            data = json.loads(raw)
            sources = data.get("sources", {})
            acc_done = len(sources.get("account", {}).get("completed", []))
            wal_done = len(sources.get("wallet", {}).get("completed", []))
            return {
                "accounts_completed": acc_done,
                "wallets_completed": wal_done,
                "total_completed": acc_done + wal_done,
                "raw": data,
            }
        except Exception:
            continue
    return {"accounts_completed": 0, "wallets_completed": 0, "total_completed": 0}


class WebRunState:
    def __init__(self) -> None:
        self.lock = threading.Lock()
        self.is_running = False
        self.is_paused = False
        self.total = 0
        self.completed = 0
        self.remaining = 0
        self.matches = 0
        self.mismatches = 0
        self.mismatch_total_halalas = 0
        self.errors = 0
        self.current_task: dict[str, Any] | None = None
        self.table_rows: list[dict[str, Any]] = []
        self.logs: list[dict[str, str]] = []
        self.checker_state: Any = None
        self.active_config: dict[str, Any] | None = None

    def reset(self, total: int = 0) -> None:
        with self.lock:
            self.is_running = False
            self.is_paused = False
            self.total = total
            self.completed = 0
            self.remaining = total
            self.matches = 0
            self.mismatches = 0
            self.mismatch_total_halalas = 0
            self.errors = 0
            self.current_task = None
            self.table_rows = []
            self.logs = []

    def update_task(self, task: dict[str, Any]) -> None:
        with self.lock:
            self.current_task = task

    def record_result(self, row: dict[str, Any]) -> None:
        with self.lock:
            self.table_rows.insert(0, row)
            self.completed += 1
            if self.remaining > 0:
                self.remaining -= 1
            if row.get("status") == "match":
                self.matches += 1
            elif row.get("status") == "mismatch":
                self.mismatches += 1
                try:
                    expected_halalas = round(float(row.get("expected_amount") or 0) * 100)
                    live_halalas = round(float(row.get("live_amount") or 0) * 100)
                    self.mismatch_total_halalas += expected_halalas - live_halalas
                except (TypeError, ValueError):
                    pass
            elif row.get("status") == "error":
                self.errors += 1

    def add_log(self, text: str, level: str = "ok") -> None:
        with self.lock:
            now_str = time.strftime("%H:%M:%S")
            self.logs.insert(0, {"time": now_str, "msg": text, "level": level})
            if len(self.logs) > 100:
                self.logs.pop()

    def get_snapshot(self) -> dict[str, Any]:
        with self.lock:
            network_waiting = False
            paused_workers = []
            if self.checker_state is not None:
                network_waiting = bool(getattr(self.checker_state, "network_waiting", False))
                paused_workers = list(getattr(self.checker_state, "paused_workers", []))
            return {
                "running": self.is_running,
                "active_config": getattr(self, "active_config", None),
                "paused": self.is_paused or network_waiting,
                "network_waiting": network_waiting,
                "paused_workers": paused_workers,
                "kpis": {
                    "total": self.total,
                    "completed": self.completed,
                    "remaining": self.remaining,
                    "matches": self.matches,
                    "mismatches": self.mismatches,
                    "mismatch_total": self.mismatch_total_halalas / 100.0,
                    "errors": self.errors,
                },
                "current_task": self.current_task,
                "table_rows": self.table_rows[:50],
                "logs": self.logs[:30],
            }


run_state = WebRunState()


def build_engine_payload_from_queue_job(job: Any) -> dict[str, Any]:
    mapping = job.column_mapping or {}
    url_type = "mixed" if job.mode == "smart_hybrid" else ("account" if job.mode == "account_only" else "wallet")

    has_amt_2 = (job.amount_target == "smart_dual" or bool(mapping.get("amount_col_2")))
    amt_col = mapping.get("amount_col", "O")
    if job.amount_target == "contract" and mapping.get("amount_col_2"):
        amt_col = mapping.get("amount_col_2")
        has_amt_2 = False

    custom_cfg = {
        "sheet": str(job.sheet_index),
        "lookup_col": mapping.get("lookup_col", "L"),
        "amount_col": amt_col,
        "has_amount_col_2": has_amt_2,
        "amount_col_2": mapping.get("amount_col_2", "AS") if has_amt_2 else "",
        "customer_col": mapping.get("customer_col", "G"),
        "no_customer": bool(mapping.get("no_customer", not bool(mapping.get("customer_col")))),
        "collector_col": mapping.get("collector_col", "U"),
        "no_collector": bool(mapping.get("no_collector", True)),
        "service_col": mapping.get("service_col", "AQ"),
        "url_type": url_type,
    }

    return {
        "workbook": job.filename,
        "scope": "custom",
        "collector": "all",
        "start_mode": "continue",
        "custom_config": custom_cfg,
        "queue_job_id": job.id,
        "result_file": job.result_file,
        "checkpoint_file": job.checkpoint_file,
    }


def start_queue_worker_service() -> tuple[bool, str]:
    from zain_checker.queue_manager import QUEUE_MANAGER
    if run_state.is_running:
        return False, "عملية الفحص جارية بالفعل"
    next_job = QUEUE_MANAGER.get_next_pending_job()
    if not next_job:
        return False, "لا توجد شيتات قيد الانتظار في الطابور"
    threading.Thread(target=run_queue_worker_loop, daemon=True).start()
    return True, "تم بدء تشغيل الطابور بنجاح"


def run_queue_worker_loop() -> None:
    """Sequential queue execution loop: processes pending jobs one after another."""
    from zain_checker.queue_manager import QUEUE_MANAGER
    while True:
        next_job = QUEUE_MANAGER.get_next_pending_job()
        if not next_job:
            run_state.add_log("🏁 اكتمل فحص جميع الشيتات في طابور الفحص بنجاح!", "ok")
            try:
                from zain_checker.telegram_controller import notify_user
                notify_user("🏁 **اكتمل فحص جميع الشيتات في طابور الفحص بنجاح!** 🎉")
            except Exception:
                pass
            break

        job_id = next_job.id
        QUEUE_MANAGER.mark_job_running(job_id)
        run_state.add_log(f"🚀 بدء فحص الشيت التالي في الطابور: {next_job.filename} (ورقة: {next_job.sheet_name})", "ok")

        payload = build_engine_payload_from_queue_job(next_job)
        try:
            run_real_checker_engine(payload, queue_job=next_job)
        except Exception as exc:
            QUEUE_MANAGER.mark_job_completed(job_id, error_msg=str(exc))
            run_state.add_log(f"⚠️ خطأ أثناء فحص الشيت {next_job.filename}: {exc}", "err")

        time.sleep(2.0)


class WebBridgeHandler(BaseHTTPRequestHandler):
    def end_headers(self) -> None:
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        super().end_headers()

    def do_OPTIONS(self) -> None:
        self.send_response(204)
        self.end_headers()

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path

        if path == "/api/session-info":
            self.handle_session_info()
            return
        elif path == "/api/live-status":
            self.handle_live_status()
            return
        elif path == "/api/queue":
            self.handle_get_queue()
            return
        elif path == "/api/telegram/status":
            self.handle_telegram_status()
            return

        # Serve Static Files
        if path == "/" or path == "":
            file_path = STATIC_DIR / "index.html"
        else:
            rel = path.lstrip("/")
            file_path = STATIC_DIR / rel

        if file_path.exists() and file_path.is_file():
            content_type = "text/html; charset=utf-8"
            if file_path.suffix == ".css":
                content_type = "text/css; charset=utf-8"
            elif file_path.suffix == ".js":
                content_type = "application/javascript; charset=utf-8"
            elif file_path.suffix == ".json":
                content_type = "application/json; charset=utf-8"
            elif file_path.suffix == ".png":
                content_type = "image/png"
            elif file_path.suffix in (".ico", ".icon"):
                content_type = "image/x-icon"
            elif file_path.suffix in (".jpg", ".jpeg"):
                content_type = "image/jpeg"

            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.end_headers()
            with open(file_path, "rb") as f:
                self.wfile.write(f.read())
        else:
            self.send_error(404, f"File {path} not found.")

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path

        content_length = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(content_length).decode("utf-8") if content_length > 0 else "{}"
        try:
            payload = json.loads(body)
        except Exception:
            payload = {}

        if path == "/api/inspect-sheets":
            self.handle_inspect_sheets(payload)
        elif path == "/api/inspect-sheet-columns":
            self.handle_inspect_sheet_columns(payload)
        elif path == "/api/queue/add":
            self.handle_queue_add(payload)
        elif path == "/api/queue/remove":
            self.handle_queue_remove(payload)
        elif path == "/api/queue/clear-completed":
            self.handle_queue_clear_completed()
        elif path == "/api/queue/start":
            self.handle_queue_start()
        elif path == "/api/start-session":
            self.handle_start_session(payload)
        elif path == "/api/browse-file":
            self.handle_browse_file()
        elif path == "/api/upload-workbook":
            self.handle_upload_workbook(payload)
        elif path == "/api/pause-session":
            self.handle_pause_session()
        elif path == "/api/resume-session":
            self.handle_resume_session()
        elif path == "/api/resume-worker":
            self.handle_resume_worker(payload)
        elif path == "/api/cancel-session":
            self.handle_cancel_session()
        elif path == "/api/pre-zain-audit":
            self.handle_pre_zain_audit(payload)
        elif path == "/api/build-clean-sheet":
            self.handle_build_clean_sheet(payload)
        elif path == "/api/run-simulation":
            self.handle_run_simulation()
        elif path == "/api/telegram/test-ping":
            self.handle_telegram_test_ping()
        else:
            self.send_error(404, "Unknown API endpoint.")


    def handle_build_clean_sheet(self, payload: dict[str, Any]) -> None:
        """Runs the sequential first-seen memory Clean Sheet Builder and exports <name>_الشيت_النظيف.xlsx."""
        try:
            from zain_checker.clean_sheet_builder import build_sequential_clean_sheet

            wb_name = payload.get("workbook", "zain_data.xlsx")
            wb_path = PROJECT_DIR / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
            if not wb_path.exists():
                self.send_json(404, {"status": "error", "message": f"الملف غير موجود: {wb_name}"})
                return

            collector = payload.get("collector", "")
            filter_status = bool(payload.get("filter_status", False))
            res = build_sequential_clean_sheet(
                workbook_path=wb_path,
                target_collector=collector if collector != "all" else "",
                filter_status=filter_status,
            )
            self.send_json(
                200,
                {
                    "status": "ok",
                    "summary": res.to_summary_dict(),
                    "clean_workbook_name": res.clean_workbook_path.name,
                    "clean_workbook_path": str(res.clean_workbook_path),
                },
            )
        except Exception as e:
            self.send_json(500, {"status": "error", "message": str(e)})

    def handle_upload_workbook(self, payload: dict[str, Any]) -> None:
        """Handles instant file selection from browser <input type='file'>."""
        try:
            import base64

            filename = os.path.basename(str(payload.get("filename") or "uploaded_workbook.xlsx"))
            if not filename.lower().endswith((".xlsx", ".xls")):
                filename += ".xlsx"
            content_b64 = payload.get("content_base64") or ""
            target_path = PROJECT_DIR / filename
            if content_b64:
                raw_bytes = base64.b64decode(content_b64)
                target_path.write_bytes(raw_bytes)
            elif not target_path.exists():
                self.send_json(400, {"status": "error", "message": "No file content received."})
                return

            norm_path = str(target_path.resolve())
            sheets = fast_read_sheet_names(target_path)
            smart_columns = inspect_smart_columns_for_workbook(target_path)
            self.send_json(
                200,
                {
                    "status": "ok",
                    "path": norm_path,
                    "filename": filename,
                    "sheets": sheets,
                    "smart_columns": smart_columns,
                },
            )
        except Exception as e:
            self.send_json(500, {"status": "error", "message": str(e)})

    def handle_pre_zain_audit(self, payload: dict[str, Any]) -> None:
        try:
            from zain_checker.pipeline import run_pre_zain_pipeline
            from zain_checker.reconciliation import export_comprehensive_audit_workbook

            wb_name = payload.get("workbook", "zain_data.xlsx")
            wb_path = PROJECT_DIR / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
            if not wb_path.exists():
                self.send_json(404, {"status": "error", "message": f"Workbook not found: {wb_name}"})
                return

            collector = payload.get("collector", "")
            plan = run_pre_zain_pipeline(
                workbook_path=wb_path,
                target_collector=collector if collector != "all" else None,
                all_collectors=(not collector or collector == "all"),
            )
            audit_out = PROJECT_DIR / "تقرير_التدقيق_الشامل_قبل_زين.xlsx"
            export_comprehensive_audit_workbook(plan, audit_out)
            self.send_json(
                200,
                {
                    "status": "ok",
                    "audit_workbook": str(audit_out),
                    "readiness_summary": plan.to_dict()["readiness_summary"],
                    "warnings": plan.warnings,
                },
            )
        except Exception as exc:
            self.send_json(500, {"status": "error", "message": str(exc)})

    def handle_run_simulation(self) -> None:
        try:
            from zain_checker.simulation_suite import run_and_verify_section_11_simulation

            out_dir = PROJECT_DIR / "simulation_outputs"
            plan = run_and_verify_section_11_simulation(out_dir)
            self.send_json(
                200,
                {
                    "status": "ok",
                    "output_directory": str(out_dir),
                    "readiness_summary": plan.to_dict()["readiness_summary"],
                },
            )
        except Exception as exc:
            self.send_json(500, {"status": "error", "message": str(exc)})

    def handle_browse_file(self) -> None:
        try:
            import subprocess

            script = (
                "import tkinter as tk; from tkinter import filedialog; "
                "root = tk.Tk(); root.withdraw(); root.wm_attributes('-topmost', 1); root.focus_force(); "
                "p = filedialog.askopenfilename(parent=root, title='اختر ملف الإكسل (Select Excel File)', "
                "filetypes=[('Excel Files (*.xlsx, *.xls)', '*.xlsx *.xls'), ('All Files (*.*)', '*.*')]); "
                "root.destroy(); print(p, end='')"
            )
            proc = subprocess.run(
                [sys.executable, "-c", script],
                capture_output=True,
                text=True,
                encoding="utf-8",
                timeout=120,
            )
            file_path = (proc.stdout or "").strip()
            if file_path:
                norm_path = os.path.normpath(file_path)
                sheets = fast_read_sheet_names(Path(norm_path))
                smart_columns = inspect_smart_columns_for_workbook(Path(norm_path))
                self.send_json(
                    200,
                    {
                        "status": "ok",
                        "path": norm_path,
                        "filename": os.path.basename(norm_path),
                        "sheets": sheets,
                        "smart_columns": smart_columns,
                    },
                )
            else:
                self.send_json(200, {"status": "cancelled"})
        except Exception as e:
            self.send_json(500, {"status": "error", "message": str(e)})

    def handle_session_info(self) -> None:
        workbooks = []
        for p in PROJECT_DIR.glob("*.xlsx"):
            if not p.name.startswith("~$") and not p.name.startswith("نتائج فحص"):
                workbooks.append(p.name)

        if "zain_data.xlsx" in workbooks:
            workbooks.remove("zain_data.xlsx")
            workbooks.insert(0, "zain_data.xlsx")

        cp_info = read_checkpoint_info()
        first_sheets = []
        if workbooks:
            first_sheets = fast_read_sheet_names(PROJECT_DIR / workbooks[0])

        response_data = {
            "workbooks": workbooks,
            "default_sheets": first_sheets,
            "checkpoint": cp_info,
            "collectors": [
                {"id": "all", "name": "5. فحص كامل الورقة (All the sheet)", "code": "الكل"},
                {"id": "maha", "name": "1. مها أمير حسين محمد", "code": "M02253"},
                {"id": "azza", "name": "2. عزة مغربي محمد عبدالرحمن", "code": "M02254"},
                {"id": "mawada", "name": "3. مودة يحيى تاج الدين فضيل", "code": "M02051"},
                {"id": "custom", "name": "4. اسم محصل مخصص (كتابة يدوية)...", "code": "Custom"},
            ],
        }
        self.send_json(200, response_data)

    def handle_inspect_sheets(self, payload: dict[str, Any]) -> None:
        wb_name = payload.get("workbook", "zain_data.xlsx")
        wb_path = PROJECT_DIR / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        if wb_path.exists():
            sheets = fast_read_sheet_names(wb_path)
            smart_columns = inspect_smart_columns_for_workbook(wb_path)
            self.send_json(200, {"status": "ok", "sheets": sheets, "smart_columns": smart_columns})
        else:
            self.send_json(404, {"status": "error", "message": f"File {wb_name} not found."})

    def handle_get_queue(self) -> None:
        from zain_checker.queue_manager import QUEUE_MANAGER
        jobs = QUEUE_MANAGER.get_jobs()
        self.send_json(200, {"status": "ok", "jobs": jobs})

    def handle_inspect_sheet_columns(self, payload: dict[str, Any]) -> None:
        wb_name = payload.get("workbook", "zain_data.xlsx")
        wb_path = PROJECT_DIR / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        sheet_idx = int(payload.get("sheet_index", 0))
        if not wb_path.exists():
            self.send_json(404, {"status": "error", "message": f"الملف غير موجود: {wb_name}"})
            return
        analysis = inspect_full_schema_for_sheet(wb_path, sheet_idx)
        self.send_json(200, {"status": "ok", "analysis": analysis})

    def handle_queue_add(self, payload: dict[str, Any]) -> None:
        from zain_checker.queue_manager import QUEUE_MANAGER
        wb_name = payload.get("workbook", "")
        wb_path = PROJECT_DIR / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        if not wb_path.exists():
            self.send_json(404, {"status": "error", "message": f"الملف غير موجود: {wb_name}"})
            return

        sheet_idx = int(payload.get("sheet_index", 0))
        sheet_name = str(payload.get("sheet_name", f"ورقة {sheet_idx+1}"))
        mode = str(payload.get("mode", "smart_hybrid"))
        amount_target = str(payload.get("amount_target", "remaining"))
        col_mapping = payload.get("column_mapping") or {}
        total_records = int(payload.get("total_records", 0))

        job = QUEUE_MANAGER.add_job(
            file_path=wb_path,
            sheet_index=sheet_idx,
            sheet_name=sheet_name,
            mode=mode,
            amount_target=amount_target,
            column_mapping=col_mapping,
            total_records=total_records,
        )

        auto_start = bool(payload.get("auto_start", False))
        if auto_start and not run_state.is_running:
            threading.Thread(target=self._run_queue_worker, daemon=True).start()

        self.send_json(200, {"status": "ok", "job": job.to_dict(), "jobs": QUEUE_MANAGER.get_jobs()})

    def handle_queue_remove(self, payload: dict[str, Any]) -> None:
        from zain_checker.queue_manager import QUEUE_MANAGER
        job_id = str(payload.get("job_id", ""))
        removed = QUEUE_MANAGER.remove_job(job_id)
        if removed:
            self.send_json(200, {"status": "ok", "jobs": QUEUE_MANAGER.get_jobs()})
        else:
            self.send_json(400, {"status": "error", "message": "لا يمكن حذف الشيت أثناء تشغيله أو أنه غير موجود."})

    def handle_queue_clear_completed(self) -> None:
        from zain_checker.queue_manager import QUEUE_MANAGER
        QUEUE_MANAGER.clear_completed()
        self.send_json(200, {"status": "ok", "jobs": QUEUE_MANAGER.get_jobs()})

    def handle_telegram_status(self) -> None:
        self.send_json(200, {
            "status": "ok",
            "bot_username": "ZainCheckerbot",
            "authorized_user_id": 1085138908,
            "connected": True,
        })

    def handle_telegram_test_ping(self) -> None:
        try:
            from zain_checker.telegram_controller import notify_user
            notify_user(
                "🔔 **فحص اتصال ناجح!**\n\n"
                "• تم اختبار الاتصال بنجاح من واجهة المتصفح (Web Workstation).\n"
                "• البوت متصل بالكامل وجاهز لاستقبال وتمرير الشيتات للطابور 🚀"
            )
            self.send_json(200, {"status": "ok", "message": "تم إرسال إشعار الاختبار بنجاح"})
        except Exception as exc:
            self.send_json(500, {"status": "error", "message": str(exc)})

    def handle_queue_start(self) -> None:
        from zain_checker.queue_manager import QUEUE_MANAGER
        ok, msg = start_queue_worker_service()
        if not ok:
            self.send_json(400, {"status": "error", "message": msg})
            return
        self.send_json(200, {"status": "started", "jobs": QUEUE_MANAGER.get_jobs()})

    def _run_queue_worker(self) -> None:
        run_queue_worker_loop()

    def _build_engine_payload_from_queue_job(self, job) -> dict[str, Any]:
        return build_engine_payload_from_queue_job(job)

    def handle_start_session(self, payload: dict[str, Any]) -> None:
        if run_state.is_running:
            self.send_json(400, {
                "status": "already_running",
                "message": "تنبيه أمان: جلسة الفحص جارية بالفعل في الخلفية! لا يمكن بدء فحص متزامن لتفادي تضارب البيانات."
            })
            return

        wb_name = payload.get("workbook", "zain_data.xlsx")
        workbook_path = PROJECT_DIR / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        if not workbook_path.exists():
            msg = f"❌ تنبيه: لم يتم العثور على ملف الإكسل ({wb_name}) في مجلد البرنامج! يرجى وضع الملف بجوار التطبيق والمحاولة مجدداً."
            run_state.add_log(msg, "err")
            self.send_json(404, {"status": "error", "message": msg})
            return

        scope = payload.get("scope", "both")
        collector = payload.get("collector", "all")

        run_state.reset(total=1970)
        run_state.is_running = True
        run_state.active_config = payload
        run_state.add_log(f"Starting session on {wb_name} (Scope: {scope}, Collector: {collector}).", "ok")

        try:
            from zain_checker.telegram_controller import notify_user
            notify_user(
                f"🚀 **بدء جلسة فحص جديدة!**\n\n"
                f"• الملف: `{wb_name}`\n"
                f"• النطاق: `{scope}` | المحصل: `{collector}`\n"
                f"• تم تشغيل الجلسة بنجاح من واجهة البرنامج."
            )
        except Exception:
            pass
        
        # Reset live feed and active task for fresh run
        try:
            (PROJECT_DIR / ".live_feed.jsonl").write_text("", encoding="utf-8")
            (PROJECT_DIR / ".active_task.json").write_text("{}", encoding="utf-8")
        except Exception:
            pass

        # Sync with queue manager
        from zain_checker.queue_manager import QUEUE_MANAGER
        active_job = QUEUE_MANAGER.get_active_job()
        if not active_job:
            active_job = QUEUE_MANAGER.add_job(
                file_path=workbook_path,
                sheet_index=int(payload.get("custom_config", {}).get("sheet", 0) if payload.get("custom_config") else 0),
                sheet_name=wb_name,
                mode=scope,
                amount_target="smart_dual" if payload.get("custom_config", {}).get("has_amount_col_2") else "remaining",
                column_mapping=payload.get("custom_config") or {},
                total_records=1970,
            )
            QUEUE_MANAGER.mark_job_running(active_job.id)

        # Launch real engine in background thread
        threading.Thread(target=self._run_real_checker_engine, args=(payload, active_job), daemon=True).start()
        self.send_json(200, {"status": "started", "config": payload})


    @staticmethod
    def _run_real_checker_engine(payload: dict[str, Any], queue_job: Any = None) -> None:
        """Launches the actual Zain Checker pipeline from main.py."""
        try:
            from main import (
                load_records,
                prepare_run,
                ensure_checker_chrome_profile_ready,
                CheckerState,
                BridgeServer,
                launch_initial_incognito_tab,
                launch_initial_tabs_for_workers,
                RESULT_WORKBOOK_PATH,
                write_checkpoint,
                append_mismatch,
                append_error,
                resolve_redirect_error,
                load_unresolved_redirect_keys,
                selected_source_names,
                Mismatch,
                Customer,
                ProgressMismatch,
                CheckError,
            )
            from datetime import datetime

            wb_name = payload.get("workbook", "zain_data.xlsx")
            workbook_path = PROJECT_DIR / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
            
            job_cp_file = payload.get("checkpoint_file") or (queue_job.checkpoint_file if queue_job else None)
            job_cp_path = (PROJECT_DIR / job_cp_file) if job_cp_file else None

            res_filename = payload.get("result_file") or (queue_job.result_file if queue_job else "نتائج فحص زين.xlsx")
            target_result_path = PROJECT_DIR / res_filename

            source_mode = payload.get("scope", "both")
            raw_collector = payload.get("collector", "all")
            
            # Map collector id to real collector name
            collector_map = {
                "all": "",
                "maha": "مها امير حسين محمد-M02253",
                "azza": "عزة مغربى محمد عبدالرحمن-M02254",
                "mawada": "مودة يحي تاج الدين فضيل-M02051",
            }
            target_collector = collector_map.get(raw_collector, raw_collector)

            min_debt = payload.get("min_debt")
            max_debt = payload.get("max_debt")
            min_amount_halalas = int(round(float(min_debt) * 100)) if min_debt is not None else None
            max_amount_halalas = int(round(float(max_debt) * 100)) if max_debt is not None else None

            from main import parse_excel_column
            
            # ALWAYS parse custom_config if it's sent from the frontend to allow overrides in ANY mode.
            cfg = payload.get("custom_config") or {}
            lookup_col = parse_excel_column(str(cfg.get("lookup_col", "L"))) or 12
            amount_col = parse_excel_column(str(cfg.get("amount_col", "O"))) or 15
            
            no_customer = bool(cfg.get("no_customer", False))
            customer_col = parse_excel_column(str(cfg.get("customer_col", "G"))) if not no_customer else None
            
            no_collector = bool(cfg.get("no_collector", False))
            collector_col = parse_excel_column(str(cfg.get("collector_col", "U"))) if not no_collector else None
            
            record_type = str(cfg.get("url_type", "mixed"))
            service_col = parse_excel_column(str(cfg.get("service_col", "AQ"))) or 43
            
            has_amount_col_2 = bool(cfg.get("has_amount_col_2", False))
            amount_col_2_raw = cfg.get("amount_col_2", "")
            amount_col_2 = parse_excel_column(str(amount_col_2_raw)) if (has_amount_col_2 and amount_col_2_raw and str(amount_col_2_raw).strip()) else None

            if amount_col_2:
                col2_letter = str(amount_col_2_raw).upper()
                run_state.add_log(f"تم تفعيل فحص المبلغ البديل (ع2): العمود {col2_letter}", "info")

            custom_configs = None

            if source_mode == "custom":
                # Determine sheet index for custom mode
                sheet_val = str(cfg.get("sheet", "0")).strip()
                if sheet_val.isdigit():
                    sheet_idx = int(sheet_val)
                else:
                    # Match by name
                    detected_sheets = fast_read_sheet_names(workbook_path)
                    sheet_idx = 0
                    for s_i, s_name in enumerate(detected_sheets):
                        if s_name.strip() == sheet_val or s_name in sheet_val:
                            sheet_idx = s_i
                            break

                custom_configs = [{
                    "sheet_index": sheet_idx,
                    "lookup_column": lookup_col,
                    "amount_column": amount_col,
                    "customer_column": customer_col,
                    "collector_column": collector_col,
                    "record_type": record_type,
                    "service_column": service_col,
                    "no_customer": no_customer,
                    "no_collector": no_collector,
                    "amount_column_2": amount_col_2,
                }]
            else:
                # For 'account', 'wallet', 'both', we pass the overrides as well
                # We do this by essentially passing a custom_configs array with the predetermined sheet indices,
                # but with the user's column overrides!
                detected_sheets = fast_read_sheet_names(workbook_path)
                num_sheets = len(detected_sheets) if detected_sheets else 1
                if num_sheets <= 1:
                    # Single-sheet file: Automatically use sheet_index 0 regardless of mode
                    rec_type = "mixed" if source_mode == "both" else ("wallet" if source_mode == "wallet" else "account")
                    custom_configs = [{
                        "sheet_index": 0,
                        "lookup_column": lookup_col,
                        "amount_column": amount_col,
                        "customer_column": customer_col,
                        "collector_column": collector_col,
                        "record_type": rec_type,
                        "service_column": service_col,
                        "no_customer": no_customer,
                        "no_collector": no_collector,
                        "amount_column_2": amount_col_2,
                    }]
                else:
                    custom_configs = []
                    if source_mode in ("account", "both"):
                        custom_configs.append({
                            "sheet_index": 1, # ACCOUNTS_WORKSHEET_INDEX
                            "lookup_column": lookup_col,
                            "amount_column": amount_col,
                            "customer_column": customer_col,
                            "collector_column": collector_col,
                            "record_type": "account", # Force account mode for this sheet
                            "service_column": service_col,
                            "no_customer": no_customer,
                            "no_collector": no_collector,
                            "amount_column_2": amount_col_2,
                        })
                    if source_mode in ("wallet", "both"):
                        custom_configs.append({
                            "sheet_index": 0, # WALLET_WORKSHEET_INDEX
                            "lookup_column": lookup_col,
                            "amount_column": amount_col,
                            "customer_column": customer_col,
                            "collector_column": collector_col,
                            "record_type": "wallet", # Force wallet mode for this sheet
                            "service_column": service_col,
                            "no_customer": no_customer,
                            "no_collector": no_collector,
                            "amount_column_2": amount_col_2,
                        })
                
                # Now we force source_mode to 'custom' so workbook.py handles it with the overrides!
                source_mode = "custom"
            filter_status = payload.get("filter_status")
            status_col_raw = payload.get("status_col")
            status_col = parse_excel_column(str(status_col_raw)) if (status_col_raw and str(status_col_raw).strip()) else None

            if filter_status:
                col_name = str(status_col_raw).upper() if status_col_raw else "Y"
                run_state.add_log(f"تصفية الحالة مفعلة على العمود {col_name} (فارغ / شرطات / عدم توصل)", "info")
            elif filter_status is False:
                run_state.add_log("تصفية الحالة معطلة (فحص كافة السجلات بدون تصفية بالحالة)", "info")

            run_state.add_log("Loading records from Excel...", "ok")
            all_customers, all_data_errors = load_records(
                workbook_path, 
                source_mode,
                target_collector,
                custom_configs,
                filter_status=filter_status,
                status_column=status_col,
            )

            if str(payload.get("start_mode", "continue")) in ("2", "new"):
                try:
                    if RESULT_WORKBOOK_PATH.exists():
                        RESULT_WORKBOOK_PATH.unlink()
                except Exception:
                    pass

            # Record any initial data/configuration errors
            if all_data_errors:
                for data_err in all_data_errors:
                    try:
                        append_error(RESULT_WORKBOOK_PATH, data_err)
                        run_state.add_log(f"⚠️ خطأ بيانات في الصف {data_err.row_numbers_text}: {data_err.details}", "warn")
                    except Exception:
                        pass

            if source_mode == "custom":
                records_by_source = {"custom": all_customers}
            else:
                records_by_source = {
                    "account": [c for c in all_customers if c.record_type == "account"],
                    "wallet": [c for c in all_customers if c.record_type == "wallet"],
                }

            # Apply debt filters
            if min_amount_halalas is not None or max_amount_halalas is not None:
                for source in records_by_source:
                    filtered = []
                    for c in records_by_source[source]:
                        valid = True
                        if min_amount_halalas is not None and c.expected_amount < min_amount_halalas:
                            valid = False
                        if max_amount_halalas is not None and c.expected_amount > max_amount_halalas:
                            valid = False
                        if valid:
                            filtered.append(c)
                    records_by_source[source] = filtered

            selected_sources = selected_source_names(source_mode)
            selected_customer_count = sum(len(records_by_source[s]) for s in selected_sources if s in records_by_source)
            run_state.total = selected_customer_count
            run_state.remaining = selected_customer_count

            if selected_customer_count == 0:
                run_state.add_log("No records found for the chosen criteria.", "warn")
                run_state.is_running = False
                return

            raw_start_mode = str(payload.get("start_mode", "continue"))
            raw_start_val = payload.get("start_value")
            start_val_int = int(raw_start_val) if raw_start_val and str(raw_start_val).isdigit() else None
            
            mode_map = {
                "1": "continue", "continue": "continue",
                "2": "new", "new": "new",
                "6": "end", "end": "end",
                "4": "row", "row": "row",
                "3": "search", "search": "search",
                "5": "customer", "customer": "customer",
            }
            mapped_mode = mode_map.get(raw_start_mode, "continue")
            start_choice = (mapped_mode, start_val_int)

            customers, mismatches, progress, direction = prepare_run(
                workbook_path.name,
                records_by_source,
                source_mode,
                start_choice=start_choice,
                checkpoint_path=job_cp_path,
            )

            run_state.sheet_total = selected_customer_count
            run_state.previously_completed = max(0, selected_customer_count - len(customers))
            run_state.session_target = len(customers)
            run_state.total = len(customers)
            run_state.remaining = len(customers)
            unresolved_redirects = load_unresolved_redirect_keys(target_result_path)

            def save_progress(
                next_index: int,
                current_mismatches: list[ProgressMismatch],
                completed: bool,
                completed_customer: Customer | None = None,
            ) -> None:
                del completed
                if completed_customer is not None:
                    source_state = progress["sources"][completed_customer.progress_source]
                    if completed_customer.lookup_number not in source_state["completed"]:
                        source_state["completed"].append(completed_customer.lookup_number)
                elif next_index > 0 and next_index <= len(customers):
                    checked_customer = customers[next_index - 1]
                    source_state = progress["sources"][checked_customer.progress_source]
                    if checked_customer.lookup_number not in source_state["completed"]:
                        source_state["completed"].append(checked_customer.lookup_number)
                write_checkpoint(progress, current_mismatches, checkpoint_path=job_cp_path)
                run_state.completed = len(state.completed_indices)
                run_state.matches = max(0, run_state.completed - len(state.mismatches))
                run_state.remaining = max(0, run_state.total - run_state.completed)
                
                if queue_job:
                    from zain_checker.queue_manager import QUEUE_MANAGER
                    QUEUE_MANAGER.update_active_progress(
                        queue_job.id,
                        completed=run_state.completed,
                        remaining=run_state.remaining,
                        matches=run_state.matches,
                        mismatches=len(state.mismatches),
                        errors=run_state.errors,
                    )

            def save_mismatch(customer: Customer, website_amount: int) -> None:
                append_mismatch(target_result_path, Mismatch(customer, website_amount, datetime.now()))
                if target_result_path.name != RESULT_WORKBOOK_PATH.name:
                    try:
                        append_mismatch(RESULT_WORKBOOK_PATH, Mismatch(customer, website_amount, datetime.now()))
                    except Exception:
                        pass

                effective_exp = customer.expected_amount
                if getattr(customer, "expected_amount_2", None) is not None:
                    diff1 = abs(website_amount - customer.expected_amount)
                    diff2 = abs(website_amount - customer.expected_amount_2)
                    effective_exp = customer.expected_amount if diff1 <= diff2 else customer.expected_amount_2

                run_state.record_result({
                    "row": customer.row_number,
                    "record_type": customer.record_type,
                    "lookup_number": customer.lookup_number,
                    "customer_name": customer.customer_name,
                    "expected_amount": effective_exp / 100.0,
                    "live_amount": website_amount / 100.0,
                    "status": "mismatch",
                    "time": time.strftime("%H:%M:%S")
                })
                if getattr(customer, "expected_amount_2", None) is not None:
                    run_state.add_log(f"Mismatch at row {customer.row_number}: File=[{customer.expected_amount/100:.2f}, {customer.expected_amount_2/100:.2f}] (Closest: {effective_exp/100:.2f}), Zain={website_amount/100:.2f} SAR", "warn")
                else:
                    run_state.add_log(f"Mismatch at row {customer.row_number}: File={customer.expected_amount/100:.2f} SAR, Zain={website_amount/100:.2f} SAR", "warn")

            def save_error(error: CheckError) -> None:
                append_error(target_result_path, error)
                if target_result_path.name != RESULT_WORKBOOK_PATH.name:
                    try:
                        append_error(RESULT_WORKBOOK_PATH, error)
                    except Exception:
                        pass

                error_amount = (
                    error.expected_amount / 100.0
                    if error.expected_amount is not None
                    else 0.0
                )
                run_state.record_result({
                    "row": error.row_numbers[0],
                    "record_type": error.record_type,
                    "lookup_number": error.lookup_number,
                    "customer_name": error.customer_name,
                    "expected_amount": error_amount,
                    "live_amount": 0.0,
                    "status": "error",
                    "time": time.strftime("%H:%M:%S")
                })
                run_state.add_log(f"Error at row {error.row_numbers[0]}: {error.error_type} - {error.details[:80]}", "err")

            def resolve_prior_redirect(customer: Customer) -> None:
                key = (customer.record_type, customer.lookup_number)
                if key in unresolved_redirects:
                    if resolve_redirect_error(target_result_path, customer):
                        unresolved_redirects.discard(key)
                    if target_result_path.name != RESULT_WORKBOOK_PATH.name:
                        resolve_redirect_error(RESULT_WORKBOOK_PATH, customer)

            from zain_checker.chrome_manager import (
                terminate_checker_chrome_processes,
                ensure_checker_chrome_profile_ready,
                clean_extension_storage,
                get_worker_profile_dir,
            )
            from zain_checker.config import WORKERS_CONFIG
            terminate_checker_chrome_processes()
            ensure_checker_chrome_profile_ready(interactive=False)
            for w_cfg in WORKERS_CONFIG:
                clean_extension_storage(get_worker_profile_dir(w_cfg.get("worker_id", "worker_1")))

            state = CheckerState(
                customers,
                0,
                mismatches,
                save_progress,
                save_mismatch,
                save_error,
                resolve_prior_redirect,
            )
            run_state.checker_state = state

            if getattr(run_state, "bridge_server", None) is not None:
                try:
                    run_state.bridge_server.shutdown()
                    run_state.bridge_server.server_close()
                except Exception:
                    pass
                run_state.bridge_server = None

            server = BridgeServer(state)
            run_state.bridge_server = server
            threading.Thread(target=server.serve_forever, daemon=True).start()

            state.note_initial_incognito_launch()
            launch_initial_tabs_for_workers(customers, state)
            run_state.add_log("Incognito Chrome launched and checking started.", "ok")

            from zain_checker.chrome_manager import is_worker_chrome_running, clean_extension_storage, get_worker_profile_dir, terminate_worker_chrome_process
            from main import normalize_worker_id, resolve_worker_proxy, launch_initial_incognito_tab
            from zain_checker.config import WORKERS_CONFIG

            last_worker_relaunch: dict[str, float] = {}
            last_watchdog_check: float = 0.0

            while not state.finished.wait(0.5):
                if not run_state.is_running:
                    break
                if getattr(state, "network_waiting", False):
                    if not run_state.is_paused:
                        run_state.is_paused = True
                        run_state.add_log(
                            "تم إيقاف الفحص مؤقتاً لحماية الاتصال من الحظر. يمكنك تغيير الشبكة/IP والضغط على زر استئناف في اللوحة للمتابعة.",
                            "warn",
                        )
                    continue

                now = time.monotonic()

                # Chrome Liveness & Auto-Reopening Watchdog for all workers (runs every 3 seconds)
                if now - last_watchdog_check > 3.0:
                    last_watchdog_check = now
                    for w_cfg in WORKERS_CONFIG:
                        w_id = w_cfg.get("worker_id", "worker_1")
                        canonical_name, worker_tag = normalize_worker_id(w_id)

                        # Check if worker is cooling down after a temporary block
                        cooldown_expiry = state.worker_cooldown_until.get(canonical_name, 0.0)
                        if cooldown_expiry > 0.0 and now >= cooldown_expiry:
                            state.paused_workers.discard(canonical_name)
                            state.paused_workers.discard(worker_tag)
                            state.worker_cooldown_until.pop(canonical_name, None)
                            run_state.add_log(f"انتهاء فترة التهدئة لـ {canonical_name}. جاري استئناف المتصفح تلقائياً...", "ok")

                        # If worker is intentionally paused by user, skip auto-reopen
                        if canonical_name in state.paused_workers or worker_tag in state.paused_workers:
                            continue

                        # Check if Chrome is currently running for this worker
                        chrome_alive = is_worker_chrome_running(worker_tag)
                        relaunch_delay = now - last_worker_relaunch.get(worker_tag, 0.0)

                        # If Chrome is NOT running and at least 6s passed since last relaunch attempt:
                        if not chrome_alive and relaunch_delay > 6.0 and not state.completed:
                            last_worker_relaunch[worker_tag] = now
                            run_state.add_log(
                                f"⚠️ متصفح {canonical_name} تم إغلاقه أو توقف! جاري إعادة الفتح والتشغيل التلقائي فوراً...",
                                "warn",
                            )
                            # Pick next customer for this worker
                            with state.lock:
                                target_cust = None
                                target_task_id = None
                                for tid, tinfo in list(state.in_flight.items()):
                                    t_c, t_t = normalize_worker_id(tinfo.get("worker_id", ""))
                                    if t_c == canonical_name or t_t == worker_tag:
                                        target_cust = tinfo["customer"]
                                        target_task_id = tid
                                        break
                                if target_cust is None:
                                    if state.retry_queue:
                                        t_idx = state.retry_queue.pop(0)
                                        target_cust = state.customers[t_idx]
                                    elif state.assigned_index < len(state.customers):
                                        t_idx = state.assigned_index
                                        state.assigned_index += 1
                                        target_cust = state.customers[t_idx]
                                    else:
                                        target_cust = state.customers[-1]
                                    target_task_id = state._task_id(target_cust, target_cust.row_number)
                                    state.in_flight[target_task_id] = {
                                        "index": target_cust.row_number,
                                        "customer": target_cust,
                                        "worker_id": canonical_name,
                                        "leased_at": now,
                                        "recheck_stage": 0,
                                        "redirect_retry": False,
                                        "incognito_retry": False,
                                    }

                            w_proxy = resolve_worker_proxy(worker_tag)
                            clean_extension_storage(get_worker_profile_dir(worker_tag))
                            launch_initial_incognito_tab(
                                target_cust,
                                target_task_id,
                                worker_id=worker_tag,
                                proxy_server=w_proxy,
                            )
                            run_state.add_log(
                                f"✅ تم إعادة تشغيل متصفح {canonical_name} بنجاح على الصف {target_cust.row_number} ({target_cust.lookup_number}).",
                                "ok",
                            )

                stalled_workers = state.get_stalled_worker_handoffs()
                for worker_tag, stalled_cust, stalled_tid in stalled_workers:
                    canonical_name, _ = normalize_worker_id(worker_tag)
                    run_state.add_log(
                        f"⚠️ متصفح {canonical_name} تأخر في تسليم الجلسة (> {state.incognito_handoff_timeout_seconds}ث). جاري إعادة تشغيله منفرداً على الصف {stalled_cust.row_number}...",
                        "warn",
                    )
                    terminate_worker_chrome_process(worker_tag)
                    time.sleep(0.8)
                    clean_extension_storage(get_worker_profile_dir(worker_tag))
                    w_proxy = resolve_worker_proxy(worker_tag)
                    launch_initial_incognito_tab(
                        stalled_cust,
                        stalled_tid,
                        worker_id=worker_tag,
                        proxy_server=w_proxy,
                    )
                    run_state.add_log(
                        f"✅ تم إعادة تشغيل متصفح {canonical_name} منفرداً بنجاح على الصف {stalled_cust.row_number} ({stalled_cust.lookup_number}).",
                        "ok",
                    )

            run_state.is_running = False
            run_state.add_log(f"Session checking completed for {wb_name}.", "ok")

            if queue_job:
                from zain_checker.queue_manager import QUEUE_MANAGER
                from zain_checker.telegram_controller import send_queue_completion_report
                QUEUE_MANAGER.mark_job_completed(queue_job.id)
                run_state.add_log(f"✅ تم إكمال فحص الشيت {queue_job.filename} بنجاح وتحديث الطابور.", "ok")
                try:
                    send_queue_completion_report(queue_job.to_dict(), target_result_path)
                    run_state.add_log(f"📥 تم إرسال ملف النتائج {target_result_path.name} لتلقرام مع احتساب المجاميع.", "ok")
                except Exception as tg_err:
                    run_state.add_log(f"ملاحظة: تعذر إرسال النتائج لتلقرام: {tg_err}", "warn")
            else:
                try:
                    from zain_checker.telegram_controller import send_excel_report
                    send_excel_report(file_path=target_result_path, caption_header=f"✅ اكتمل فحص الجلسة: {wb_name}")
                except Exception:
                    pass

        except Exception as e:
            run_state.is_running = False
            run_state.add_log(f"Fatal execution error: {e}", "err")

    def handle_pause_session(self) -> None:
        run_state.is_paused = True
        run_state.add_log("Session paused by user.", "warn")
        self.send_json(200, {"status": "paused"})

    def handle_resume_session(self) -> None:
        run_state.is_paused = False
        if getattr(run_state, "checker_state", None) is not None:
            checker_state = run_state.checker_state
            from main import launch_initial_incognito_tab

            # 1. Resume any specific workers that were paused due to IP block (e.g. Worker 2)
            paused_list = list(getattr(checker_state, "paused_workers", set()))
            if paused_list:
                for wid in paused_list:
                    resumed_cust = checker_state.resume_worker(wid)
                    use_proxy = bool("proxy" in wid.lower() or "worker 2" in wid.lower())
                    if resumed_cust:
                        launch_initial_incognito_tab(
                            resumed_cust,
                            checker_state.current_task_id(),
                            worker_id=wid,
                            use_proxy=use_proxy,
                        )
                        run_state.add_log(
                            f"تم استئناف وتشغيل المتصفح لـ {wid} بنجاح.",
                            "ok",
                        )

            # 2. Resume global network waiting if applicable
            if getattr(checker_state, "network_waiting", False):
                try:
                    resumed_customer = checker_state.resume_after_network_change()
                    launch_initial_incognito_tab(
                        resumed_customer,
                        checker_state.current_task_id(),
                        worker_id="worker_1",
                        use_proxy=False,
                    )
                    run_state.add_log(
                        f"تم استئناف الفحص بنجاح. جاري إعادة فحص الصف {resumed_customer.row_number} ({resumed_customer.lookup_number}).",
                        "ok",
                    )
                except Exception as e:
                    run_state.add_log(f"خطأ أثناء استئناف الفحص: {e}", "err")
            else:
                run_state.add_log("Session resumed.", "ok")
        else:
            run_state.add_log("Session resumed.", "ok")
        self.send_json(200, {"status": "resumed"})

    def handle_resume_worker(self, payload: dict[str, Any]) -> None:
        worker_id = str(payload.get("worker_id", "Worker 2 (Proxy)"))
        if getattr(run_state, "checker_state", None) is not None:
            checker_state = run_state.checker_state
            resumed_cust = checker_state.resume_worker(worker_id)
            use_proxy = bool("proxy" in worker_id.lower() or "worker 2" in worker_id.lower())
            if resumed_cust:
                from main import launch_initial_incognito_tab
                launch_initial_incognito_tab(
                    resumed_cust,
                    checker_state.current_task_id(),
                    worker_id=worker_id,
                    use_proxy=use_proxy,
                )
                run_state.add_log(f"تم إعادة إطلاق متصفح {worker_id} بنجاح.", "ok")
                self.send_json(200, {"status": "ok", "worker_id": worker_id})
                return
        self.send_json(400, {"status": "error", "message": "Checker is not running or no customer found."})

    def handle_cancel_session(self) -> None:
        run_state.is_running = False
        run_state.is_paused = False
        if getattr(run_state, "checker_state", None) is not None:
            try:
                if getattr(run_state.checker_state, "network_waiting", False):
                    run_state.checker_state.stop_during_network_wait()
                else:
                    run_state.checker_state.finished.set()
            except Exception:
                pass
        if getattr(run_state, "bridge_server", None) is not None:
            try:
                run_state.bridge_server.shutdown()
                run_state.bridge_server.server_close()
            except Exception:
                pass
            run_state.bridge_server = None
        try:
            from zain_checker.chrome_manager import terminate_checker_chrome_processes
            terminate_checker_chrome_processes()
        except Exception:
            pass
        run_state.add_log("Session cancelled by user.", "err")
        self.send_json(200, {"status": "cancelled", "message": "تم إلغاء الجلسة بنجاح"})

    def handle_live_status(self) -> None:
        snapshot = run_state.get_snapshot()

        # 1. Read active task from .active_task.json (zero-overhead decoupled file)
        active_task_file = PROJECT_DIR / ".active_task.json"
        if active_task_file.exists():
            try:
                with open(active_task_file, "r", encoding="utf-8") as f:
                    snapshot["current_task"] = json.load(f)
            except Exception:
                pass

        # 2. Read live checks from .live_feed.jsonl (last 50 rows)
        feed_file = PROJECT_DIR / ".live_feed.jsonl"
        if feed_file.exists():
            try:
                with open(feed_file, "r", encoding="utf-8") as f:
                    lines = f.readlines()
                    rows = []
                    matches = 0
                    mismatches = 0
                    mismatch_total_halalas = 0
                    errors = 0
                    for line in reversed(lines):
                        line = line.strip()
                        if not line:
                            continue
                        try:
                            item = json.loads(line)
                            if len(rows) < 50:
                                rows.append(item)
                            st = item.get("status")
                            if st == "match":
                                matches += 1
                            elif st == "mismatch":
                                mismatches += 1
                                try:
                                    expected_halalas = round(float(item.get("expected_amount") or 0) * 100)
                                    live_halalas = round(float(item.get("live_amount") or 0) * 100)
                                    mismatch_total_halalas += expected_halalas - live_halalas
                                except (TypeError, ValueError):
                                    pass
                            elif st == "error":
                                errors += 1
                        except Exception:
                            continue
                    if rows:
                        snapshot["table_rows"] = rows
                        session_target = getattr(run_state, "session_target", run_state.total)
                        session_done = len(lines)
                        snapshot["kpis"]["total"] = session_target
                        snapshot["kpis"]["completed"] = session_done
                        snapshot["kpis"]["matches"] = matches
                        snapshot["kpis"]["mismatches"] = mismatches
                        snapshot["kpis"]["mismatch_total"] = mismatch_total_halalas / 100.0
                        snapshot["kpis"]["errors"] = errors
                        snapshot["kpis"]["remaining"] = max(0, session_target - session_done)
                        snapshot["kpis"]["previously_completed"] = getattr(run_state, "previously_completed", 0)
                        snapshot["kpis"]["sheet_total"] = getattr(run_state, "sheet_total", 0)
            except Exception:
                pass

        self.send_json(200, snapshot)

    def send_json(self, status_code: int, data: Any) -> None:
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status_code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: Any) -> None:
        return


run_real_checker_engine = WebBridgeHandler._run_real_checker_engine


def launch_app(open_browser: bool = True) -> None:
    """Unified entrypoint: starts Web Workstation server and opens default browser."""
    print("=" * 64)
    print("  منظومة تدقيق زين | Zain Retail Audit Workstation v2.4")
    print(f"  جاري تشغيل المنظومة على: http://localhost:{PORT}")
    print("=" * 64)

    # Automatically start Telegram Remote Control for Abdurahman with the program
    try:
        from zain_checker.telegram_controller import init_telegram_controller
        init_telegram_controller(get_stats_callback=read_checkpoint_info)
    except Exception as e:
        print(f"[Telegram] Note: could not start bot: {e}")

    if open_browser:
        threading.Timer(0.8, lambda: webbrowser.open(f"http://localhost:{PORT}")).start()

    server = ThreadingHTTPServer(("127.0.0.1", PORT), WebBridgeHandler)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n[Zain Checker] تم إيقاف الخادم.")
        server.shutdown()


def export_extension(target_dir: Path | None = None) -> Path:
    """Exports the embedded Chrome extension folder directly to the user's project directory."""
    import shutil
    if target_dir is None:
        target_dir = PROJECT_DIR / "chrome_extension"
    
    target_dir.mkdir(parents=True, exist_ok=True)
    
    # Method 1: Write directly from embedded data (100% reliable, no temp paths needed)
    try:
        from zain_checker.extension_data import EXTENSION_FILES
        for filename, content in EXTENSION_FILES.items():
            (target_dir / filename).write_text(content, encoding="utf-8")
        return target_dir
    except Exception:
        pass

    # Method 2: Copy from bundle directory as fallback
    src_dir = BUNDLE_DIRECTORY / "chrome_extension"
    if not src_dir.exists():
        alt_src = Path(__file__).resolve().parent.parent / "chrome_extension"
        if alt_src.exists():
            src_dir = alt_src

    if src_dir.exists() and src_dir.resolve() != target_dir.resolve():
        for item in src_dir.iterdir():
            dest = target_dir / item.name
            if item.is_dir():
                shutil.copytree(item, dest, dirs_exist_ok=True)
            else:
                shutil.copy2(item, dest)
            
    return target_dir


if __name__ == "__main__":
    if "-e" in sys.argv or "--export-extension" in sys.argv:
        try:
            out_dir = export_extension()
            print("=" * 64)
            print("  تصدير إضافة متصفح كروم | Chrome Extension Export")
            print("=" * 64)
            print(f"[OK] تم تصدير واستخراج إضافة كروم بنجاح إلى المجلد:")
            print(f"     {out_dir}")
            print()
            print("يمكنك الآن تثبيتها يدوياً في كروم في حال الرغبة:")
            print("  1. افتح متصفح كروم وانتقل إلى: chrome://extensions")
            print("  2. قم بتفعيل (Developer mode / وضع مطوّر البرامج).")
            print("  3. اضغط على (Load unpacked / تحميل إضافة تم فك حزمتها).")
            print(f"  4. حدد المجلد: {out_dir}")
            print("=" * 64)
        except Exception as err:
            print(f"[ERROR] فشل تصدير الإضافة: {err}")
            sys.exit(1)
        sys.exit(0)

    launch_app()
