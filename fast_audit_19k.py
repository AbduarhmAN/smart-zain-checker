# -*- coding: utf-8 -*-
"""
عنوان الاسكربت: fast_audit_19k.py
المهمة: فحص مباشر ولحظي لملف العقود (19,890 عقد) عبر 3 مسارات Business API مباشرة.
الميزات:
- طباعة النتائج لحظياً سطر بسطر (Real-time stdout with flush=True).
- حفظ تدريجي وتلقائي كل 100 عميل في ملف Excel وملف JSONL منفصل.
- استئناف تلقائي من آخر عميل محفوظ في حال الإيقاف أو الانقطاع.
- نتائج نقية ومباشرة من موقع زين مع مقارنة الفروقات.
"""
from __future__ import annotations

import io
import json
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

# قاعدة ترميز ويندوز الصارمة
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
if hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from workers.zain_api import query_contract_due_amount

# مسارات الملفات
SOURCE_FILE = Path("رقم عقد03-10-2026_من_الصف_26022.xlsx")
OUTPUT_EXCEL = Path("نتائج_فحص_زين_المباشرة_19ألف.xlsx")
OUTPUT_JSONL = Path("نتائج_فحص_زين_المباشرة_19ألف.jsonl")

# تنسيقات الإكسل
FONT_NAME = "Segoe UI"
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name=FONT_NAME, size=10)
HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
FILL_MATCH = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")    # أخضر فاتح
FILL_MISMATCH = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid") # أصفر فاتح
FILL_ERROR = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")    # أحمر فاتح
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

EXCEL_HEADERS = (
    "رقم السطر بالملف",
    "اسم العميل",
    "رقم العقد المستهدف",
    "رقم الحساب",
    "رقم الخدمة",
    "مبلغ العقد بالملف (ر.س)",
    "مبلغ المديونية بالملف (ر.س)",
    "المبلغ المعتمد للمطابقة (ر.س)",
    "المبلغ المستحق بموقع زين (ر.س)",
    "فرق المبلغ (ر.س)",
    "حالة الفحص",
    "تفاصيل الرد من زين",
    "وقت الرصد",
)


def load_completed_keys() -> set[str]:
    """تحميل الأرقام المفحوصة مسبقاً لاستئناف الفحص دون تكرار."""
    completed = set()
    if OUTPUT_JSONL.exists():
        try:
            with open(OUTPUT_JSONL, "r", encoding="utf-8") as f:
                for line in f:
                    if line.strip():
                        try:
                            data = json.loads(line)
                            contract = str(data.get("contract", "")).strip()
                            if contract:
                                completed.add(contract)
                        except Exception:
                            continue
        except Exception as e:
            print(f"⚠️ تنبيه أثناء قراءة السجل السابق: {e}", flush=True)
    return completed


def init_output_excel():
    """إنشاء ملف الإكسل وترويسته إذا لم يكن موجوداً."""
    if not OUTPUT_EXCEL.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "نتائج فحص موقع زين"
        ws.views.sheetView[0].rightToLeft = True
        ws.append(list(EXCEL_HEADERS))
        for col_idx in range(1, len(EXCEL_HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        wb.save(OUTPUT_EXCEL)
        wb.close()


def append_batch_to_excel(batch: List[Dict[str, Any]]):
    """إضافة دفعة من 100 عميل إلى ملف الإكسل وحفظه."""
    try:
        wb = openpyxl.load_workbook(OUTPUT_EXCEL)
        ws = wb.active
        for item in batch:
            row_cells = [
                item["source_row"],
                item["name"],
                item["contract"],
                item["account"],
                item["service"],
                item["contract_amount"],
                item["debt_amount"],
                item["expected_amount"],
                item["live_amount"],
                item["diff"],
                item["status_label"],
                item["message"],
                item["timestamp"],
            ]
            ws.append(row_cells)
            curr_row = ws.max_row
            st = item["status_label"]
            fill_style = FILL_MATCH if "تطابق" in st else (FILL_ERROR if "خطأ" in st or "غير موجود" in st else FILL_MISMATCH)
            for col_idx in range(1, len(row_cells) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.font = DATA_FONT
                c.border = THIN_BORDER
                c.alignment = ALIGN_CENTER
                c.fill = fill_style
        wb.save(OUTPUT_EXCEL)
        wb.close()
    except Exception as e:
        print(f"❌ خطأ أثناء حفظ الدفعة في الإكسل: {e}", flush=True)


def append_batch_to_jsonl(batch: List[Dict[str, Any]]):
    """إضافة دفعة من 100 عميل إلى ملف JSONL للحماية القصوى من انقطاع الكهرباء."""
    try:
        with open(OUTPUT_JSONL, "a", encoding="utf-8") as f:
            for item in batch:
                f.write(json.dumps(item, ensure_ascii=False) + "\n")
    except Exception as e:
        print(f"❌ خطأ أثناء كتابة الـ JSONL: {e}", flush=True)


def check_single_contract(item: Dict[str, Any]) -> Dict[str, Any]:
    """فحص عقد واحد عبر API زين المباشر مع محاولة إعادة تلقائية عند أخطاء الشبكة العابرة."""
    contract = item["contract"]
    status, live_sar, msg = "network_error", None, "لم يتم الرد"
    
    for attempt in range(2):
        try:
            status, live_sar, msg = query_contract_due_amount(contract, timeout=12)
            if status != "network_error":
                break
            time.sleep(0.3)
        except Exception as exc:
            status, live_sar, msg = "network_error", None, str(exc)
            time.sleep(0.3)

    diff = None
    if status == "ok" and live_sar is not None:
        diff = round(item["expected_amount"] - live_sar, 2)
        if abs(diff) < 0.01:
            status_label = "✅ تطابق تام"
        else:
            status_label = f"⚠️ فرق ({diff:+.2f} ر.س)"
    elif status == "not_found":
        status_label = "🔍 غير موجود بموقع زين"
    else:
        status_label = f"❌ خطأ ({status})"

    return {
        "source_row": item["source_row"],
        "name": item["name"],
        "contract": contract,
        "account": item["account"],
        "service": item["service"],
        "contract_amount": item["contract_amount"],
        "debt_amount": item["debt_amount"],
        "expected_amount": item["expected_amount"],
        "live_amount": live_sar,
        "diff": diff,
        "status": status,
        "status_label": status_label,
        "message": msg,
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def main():
    print("=" * 80, flush=True)
    print("🚀 المحرك السريع: فحص ملف العقود (19,890 عقد) - نتائج موقع زين اللحظية", flush=True)
    print(f"📁 ملف المصدر : {SOURCE_FILE}", flush=True)
    print(f"📊 ملف الإكسل : {OUTPUT_EXCEL} (يُحدّث كل 100 عميل)", flush=True)
    print(f"🔒 سجل الأمان : {OUTPUT_JSONL}", flush=True)
    print("⚡ 3 مسارات فحص متزامنة ومباشرة (Direct Business API)", flush=True)
    print("=" * 80, flush=True)

    init_output_excel()
    completed_contracts = load_completed_keys()
    print(f"ℹ️ عقود مفحوصة ومحفوظة مسبقاً: {len(completed_contracts)} عقد (سيتم تخطيها).", flush=True)

    print("⏳ جاري قراءة بيانات العقود من الشيت...", flush=True)
    t_load = time.time()
    wb_src = openpyxl.load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    ws_src = wb_src.active

    items_to_check: List[Dict[str, Any]] = []
    # Col 6: Name, Col 11: Account, Col 12: Debt, Col 40: Contract, Col 42: Service, Col 45: ContractAmt
    for idx, row in enumerate(ws_src.iter_rows(min_row=2, values_only=True), start=2):
        if not row or len(row) < 46:
            continue
        contract_raw = str(row[40] or row[11] or "").strip()
        if not contract_raw or contract_raw.lower() in ("none", "", "0"):
            continue

        if contract_raw in completed_contracts:
            continue

        name = str(row[6] or "").strip()
        account = str(row[11] or "").strip()
        service = str(row[42] or "").strip()

        try:
            debt = float(str(row[12] or "0").replace(",", "").strip())
        except Exception:
            debt = 0.0

        try:
            contract_amt = float(str(row[45] or "0").replace(",", "").strip())
        except Exception:
            contract_amt = 0.0

        expected = contract_amt if contract_amt > 0 else debt

        items_to_check.append({
            "source_row": idx,
            "name": name,
            "contract": contract_raw,
            "account": account,
            "service": service,
            "debt_amount": debt,
            "contract_amount": contract_amt,
            "expected_amount": expected,
        })

    wb_src.close()
    print(f"✅ تم تحميل {len(items_to_check)} عقد متبقي للفحص في {time.time()-t_load:.1f} ثانية.", flush=True)
    if not items_to_check:
        print("🎉 جميع العقود في الملف مفحوصة ومحفوظة بالكامل!", flush=True)
        return

    total_target = len(items_to_check)
    print(f"🏁 انطلاق الفحص لـ {total_target} عقد الآن...", flush=True)
    print("-" * 80, flush=True)

    start_time = time.time()
    batch_buffer: List[Dict[str, Any]] = []
    saved_total = len(completed_contracts)
    processed_this_run = 0
    matches_count = 0
    mismatches_count = 0
    errors_count = 0

    # التنفيذ عبر 3 عمال متزامنين (3 Parallel Workers)
    with ThreadPoolExecutor(max_workers=3) as executor:
        # Submit in rolling chunks of 150 to keep RAM ultra light
        CHUNK_SIZE = 150
        for chunk_start in range(0, total_target, CHUNK_SIZE):
            chunk = items_to_check[chunk_start:chunk_start + CHUNK_SIZE]
            futures = {executor.submit(check_single_contract, item): item for item in chunk}

            for fut in as_completed(futures):
                res = fut.result()
                processed_this_run += 1
                batch_buffer.append(res)

                if "تطابق" in res["status_label"]:
                    matches_count += 1
                elif "فرق" in res["status_label"]:
                    mismatches_count += 1
                else:
                    errors_count += 1

                elapsed = max(0.001, time.time() - start_time)
                rate = processed_this_run / elapsed

                # طباعة لحظية لكل سطر
                live_str = f"{res['live_amount']:.2f}" if res['live_amount'] is not None else "—"
                print(
                    f"[{processed_this_run:5d}/{total_target}] "
                    f"السطر {res['source_row']:5d} | "
                    f"عقد: {res['contract']:10s} | "
                    f"المتوقع: {res['expected_amount']:8.2f} | "
                    f"زين: {live_str:8s} | "
                    f"{res['status_label']:22s} | "
                    f"⚡ {rate:4.1f} سجل/ث",
                    flush=True,
                )

                # حفظ كل 100 عميل في ملف الإكسل وملف السجل
                if len(batch_buffer) >= 100:
                    t_save = time.time()
                    append_batch_to_jsonl(batch_buffer)
                    append_batch_to_excel(batch_buffer)
                    saved_total += len(batch_buffer)
                    print(
                        f"\n💾 [تم حفظ دفعة 100 عميل بنجاح] "
                        f"الإجمالي المحفوظ بالملف: {saved_total:,} عقد | "
                        f"مطابقات: {matches_count:,} | فروقات: {mismatches_count:,} | "
                        f"وقت الحفظ: {time.time()-t_save:.2f} ثانية\n"
                        f"{'-'*80}",
                        flush=True,
                    )
                    batch_buffer.clear()

    # حفظ أي متبقي في نهاية الفحص
    if batch_buffer:
        append_batch_to_jsonl(batch_buffer)
        append_batch_to_excel(batch_buffer)
        saved_total += len(batch_buffer)
        print(f"\n💾 [تم حفظ الدفعة الأخيرة ({len(batch_buffer)} عميل)] الإجمالي المحفوظ: {saved_total:,} عقد", flush=True)
        batch_buffer.clear()

    total_duration = time.time() - start_time
    overall_rate = processed_this_run / max(0.001, total_duration)

    print("=" * 80, flush=True)
    print("🏆 اكتمل فحص الملف بنجاح تام!", flush=True)
    print(f"⏱️ إجمالي الوقت المستغرق : {total_duration/60:.2f} دقيقة ({total_duration:.1f} ثانية)", flush=True)
    print(f"⚡ معدل السرعة الفعلي    : {overall_rate:.2f} عقد / ثانية", flush=True)
    print(f"📁 ملف الإكسل النهائي    : {OUTPUT_EXCEL}", flush=True)
    print(f"✅ إجمالي المطابقات التامة: {matches_count:,}", flush=True)
    print(f"⚠️ إجمالي الفروقات       : {mismatches_count:,}", flush=True)
    print(f"❌ إجمالي غير الموجود/خطأ: {errors_count:,}", flush=True)
    print("=" * 80, flush=True)


if __name__ == "__main__":
    main()
