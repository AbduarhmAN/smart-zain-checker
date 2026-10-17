# -*- coding: utf-8 -*-
"""
repair_network_errors.py
المهمة: إعادة فحص العقود التي واجهت انقطاع شبكة مؤقت (1,848 عقد) وتحديث النتائج مباشرة
في ملف الإكسل (نتائج_فحص_زين_المباشرة_19ألف.xlsx) وسجل الـ JSONL.
"""
from __future__ import annotations

import io
import json
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

# قاعدة ترميز ويندوز الصارمة
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
if hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from workers.zain_api import query_contract_due_amount

OUTPUT_EXCEL = Path("نتائج_فحص_زين_المباشرة_19ألف.xlsx")
OUTPUT_JSONL = Path("نتائج_فحص_زين_المباشرة_19ألف.jsonl")

# التنسيقات
FONT_NAME = "Segoe UI"
DATA_FONT = Font(name=FONT_NAME, size=10)
FILL_MATCH = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")    # أخضر فاتح
FILL_MISMATCH = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid") # أصفر فاتح
FILL_ERROR = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")    # أحمر فاتح
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def check_single_item(item: Dict[str, Any]) -> Dict[str, Any]:
    contract = item["contract"]
    status, live_sar, msg = "network_error", None, "لم يتم الرد"
    
    for attempt in range(3):
        try:
            status, live_sar, msg = query_contract_due_amount(contract, timeout=12)
            if status != "network_error":
                break
            time.sleep(0.5)
        except Exception as exc:
            status, live_sar, msg = "network_error", None, str(exc)
            time.sleep(0.5)

    diff = None
    if status == "ok" and live_sar is not None:
        diff = round(item["expected_amount"] - live_sar, 2)
        if abs(diff) < 0.01:
            status_label = "✅ تطابق تام"
        else:
            status_label = f"⚠️ فرق ({diff:+.2f} ر.س)"
    elif status == "not_found":
        status_label = "🔍 غير موجود بموقع زين"
    else:
        status_label = f"❌ خطأ ({status})"

    updated = dict(item)
    updated["live_amount"] = live_sar
    updated["diff"] = diff
    updated["status"] = status
    updated["status_label"] = status_label
    updated["message"] = msg
    updated["timestamp"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return updated


def main():
    print("=" * 80, flush=True)
    print("🛠️  بدء الفحص التكميلي لأخطاء الشبكة العابرة", flush=True)
    print(f"📊 ملف الإكسل : {OUTPUT_EXCEL}", flush=True)
    print(f"🔒 سجل الأمان : {OUTPUT_JSONL}", flush=True)
    print("=" * 80, flush=True)

    if not OUTPUT_JSONL.exists():
        print("❌ لم يتم العثور على ملف السجل JSONL!", flush=True)
        return

    # 1. قراءة جميع السجلات وتحديد أخطاء الشبكة
    print("⏳ قراءة السجل وتحديد العقود المطلوب إعادة فحصها...", flush=True)
    with open(OUTPUT_JSONL, "r", encoding="utf-8") as f:
        all_records = [json.loads(line) for line in f if line.strip()]

    target_items = [x for x in all_records if x.get("status") == "network_error"]
    total_targets = len(target_items)
    print(f"🎯 تم العثور على {total_targets:,} عقد بحالة network_error جاهزة للإصلاح.", flush=True)

    if total_targets == 0:
        print("🎉 لا توجد أي أخطاء شبكة! جميع العقود مفحوصة بنجاح.", flush=True)
        return

    # 2. الفحص عبر 3 عمال متوازيين
    print(f"🏁 انطلاق فحص الـ {total_targets} عقد الآن...", flush=True)
    print("-" * 80, flush=True)

    start_time = time.time()
    repaired_map: Dict[str, Dict[str, Any]] = {}
    processed_count = 0
    matches_count = 0
    diffs_count = 0
    errors_count = 0

    with ThreadPoolExecutor(max_workers=3) as executor:
        futures = {executor.submit(check_single_item, item): item for item in target_items}

        for fut in as_completed(futures):
            res = fut.result()
            processed_count += 1
            repaired_map[res["contract"]] = res

            if "تطابق" in res["status_label"]:
                matches_count += 1
            elif "فرق" in res["status_label"]:
                diffs_count += 1
            else:
                errors_count += 1

            elapsed = max(0.001, time.time() - start_time)
            rate = processed_count / elapsed
            live_str = f"{res['live_amount']:.2f}" if res['live_amount'] is not None else "—"

            print(
                f"[{processed_count:4d}/{total_targets}] "
                f"السطر {res['source_row']:5d} | "
                f"عقد: {res['contract']:10s} | "
                f"المتوقع: {res['expected_amount']:8.2f} | "
                f"زين: {live_str:8s} | "
                f"{res['status_label']:22s} | "
                f"⚡ {rate:4.1f} سجل/ث",
                flush=True,
            )

    duration = time.time() - start_time
    print("=" * 80, flush=True)
    print(f"✅ اكتمل الفحص التكميلي في {duration/60:.2f} دقيقة ({duration:.1f} ثانية)!", flush=True)
    print(f"⚡ متوسط السرعة: {processed_count/max(0.001, duration):.2f} عقد/ثانية", flush=True)
    print(f"   - المطابقات التامة: {matches_count:,}", flush=True)
    print(f"   - الفروقات المالية : {diffs_count:,}", flush=True)
    print(f"   - المتبقي بأخطاء  : {errors_count:,}", flush=True)
    print("=" * 80, flush=True)

    # 3. تحديث ملف الـ JSONL
    print("💾 جاري تحديث ملف السجل JSONL...", flush=True)
    updated_records = []
    for rec in all_records:
        contract = rec.get("contract")
        if contract in repaired_map:
            updated_records.append(repaired_map[contract])
        else:
            updated_records.append(rec)

    with open(OUTPUT_JSONL, "w", encoding="utf-8") as f:
        for rec in updated_records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print("✅ تم تحديث ملف الـ JSONL بالكامل بنجاح.", flush=True)

    # 4. تحديث ملف الإكسل
    print("📊 جاري تحديث ملف الإكسل (نتائج_فحص_زين_المباشرة_19ألف.xlsx)...", flush=True)
    t_excel = time.time()
    wb = openpyxl.load_workbook(OUTPUT_EXCEL)
    ws = wb.active

    # بناء فهرس أرقام الأسطر في الإكسل بالاعتماد على رقم العقد
    # الصف الأول ترويسة، السطور تبدأ من صف 2
    row_count = ws.max_row
    excel_contract_row = {}
    for r in range(2, row_count + 1):
        contract_cell_val = str(ws.cell(row=r, column=3).value or "").strip()
        if contract_cell_val:
            excel_contract_row[contract_cell_val] = r

    updated_excel_cells = 0
    for contract, res in repaired_map.items():
        if contract in excel_contract_row:
            r = excel_contract_row[contract]
            st = res["status_label"]
            fill_style = FILL_MATCH if "تطابق" in st else (FILL_ERROR if "خطأ" in st or "غير موجود" in st else FILL_MISMATCH)

            ws.cell(row=r, column=9).value = res["live_amount"]
            ws.cell(row=r, column=10).value = res["diff"]
            ws.cell(row=r, column=11).value = res["status_label"]
            ws.cell(row=r, column=12).value = res["message"]
            ws.cell(row=r, column=13).value = res["timestamp"]

            for col in range(1, 14):
                c = ws.cell(row=r, column=col)
                c.fill = fill_style

            updated_excel_cells += 1

    wb.save(OUTPUT_EXCEL)
    wb.close()
    print(f"🎉 تم تحديث {updated_excel_cells:,} صف في ملف الإكسل وحفظه بنجاح في {time.time()-t_excel:.2f} ثانية!", flush=True)
    print("=" * 80, flush=True)


if __name__ == "__main__":
    main()
