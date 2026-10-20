"""Offline recovery/isolation tests: no sockets, browsers, files, or deletion."""
import json
import queue
import threading
import time
import unittest
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook
from domain.models import Customer, QueueJob
from manager.orchestrator import Orchestrator
from manager.queue import QueueService
from manager.recovery import input_signature, restore_records, read_journal
from manager.checkpoint import CheckpointManager
from server.web_api import WebApiHandler
from workers.actor import WorkerActor
import workers.service_transport as transport
import workers.stealth_service_engine as browser
from zain_checker.executive_reporter import ExactTemplateReporter


def customer(index=0, service=False):
    number = ('2' if service else '1') + str(10000+index)
    return Customer(index+2, 'wallet' if service else 'account', number, number, 1000,
                    service_number=number if service else '')


def memory_queue():
    service = QueueService.__new__(QueueService)
    service._lock = threading.RLock()
    service._jobs = []
    service.storage_path = Path('UNUSED_QUEUE.json')
    service._save_to_disk = Mock()
    service._load_from_disk = Mock()
    service.pricing_settings = Mock()
    service.pricing_settings.quote.return_value = {'mode':'sheet','unit_minor':2000,'currency':'EGP'}
    return service


class RecoveryAndIsolation(unittest.TestCase):
    def setUp(self):
        self.stack = []
        for name in ('socket.create_connection','pathlib.Path.unlink'):
            guard = patch(name, side_effect=AssertionError('Forbidden in offline tests'))
            guard.start()
            self.addCleanup(guard.stop)
        with transport._IP_LOCKS_GUARD:
            transport._IP_SERVICE_LOCKS.clear()
            transport._IP_LAST_QUERY_TIME.clear()
            transport._IP_RETRY_UNTIL.clear()

    def record(self, c, **extra):
        return {'row':c.row_number,'account':c.lookup_number,'expected_sar':10,
                'live_sar':10,'status':'match', **extra}

    def test_signature_changes_with_amount_number_and_mapping(self):
        c=customer()
        original=input_signature([c],{'lookup_col':'A'},'smart_hybrid','remaining',0)
        self.assertEqual(original,input_signature([c],{'lookup_col':'A'},'smart_hybrid','remaining',0))
        self.assertNotEqual(original,input_signature([c],{'lookup_col':'B'},'smart_hybrid','remaining',0))
        c=replace(c,expected_amount=2000)
        self.assertNotEqual(original,input_signature([c],{'lookup_col':'A'},'smart_hybrid','remaining',0))

    def test_unordered_completed_rows_restore_without_rechecking(self):
        customers=[customer(i) for i in range(4)]
        saved={'completed_indices':[3,0], 'input_signature':'signature'}
        records={c.row_number:self.record(c,input_signature='signature') for c in (customers[3],customers[0])}
        restored=restore_records(customers,saved,records)
        self.assertEqual(set(restored),{0,3})

    def test_journal_ahead_of_checkpoint_is_recovered(self):
        c=customer()
        restored=restore_records([c],{'completed_indices':[], 'input_signature':'signature'},
            {c.row_number:self.record(c,input_signature='signature')})
        self.assertEqual(set(restored),{0})

    def test_legacy_transient_failures_are_retried_but_not_found_stays_complete(self):
        customers = [customer(i, service=True) for i in range(5)]
        records = {c.row_number: self.record(c, live_sar=None, status='needs_review',
            error=message) for c, message in zip(customers, (
                'رفض الطلب (HTTP 200)', 'تحتاج جلسة الويب إلى تجديد',
                'انتهت مهلة استعلام المتصفح', 'تعذر استكمال المتصفح: TargetClosedError',
                'رقم العقد لا يطابق أي سجل (EZB5029)',
            ))}
        records[customers[-1].row_number]['status'] = 'not_found'
        restored = restore_records(customers, {'completed_indices': list(range(5))}, records)
        self.assertEqual(set(restored), {4})
        self.assertIsNone(restored[4]['live_sar'])

    def test_missing_or_changed_durable_result_is_not_skipped(self):
        c=customer(); saved={'completed_indices':[0,999,-1]}
        self.assertEqual(restore_records([c],saved,{}),{})
        self.assertEqual(restore_records([c],saved,{c.row_number:self.record(c,account='9999')}),{})
        self.assertEqual(restore_records([c],saved,{c.row_number:self.record(c,expected_sar=999)}),{})
        self.assertEqual(restore_records([c],saved,{c.row_number:self.record(c,expected_sar='NaN')}),{})

    def test_malformed_checkpoint_falls_back_to_backup(self):
        manager=CheckpointManager(Path('UNUSED.json'))
        with patch('pathlib.Path.exists', return_value=True), patch('pathlib.Path.read_text',
                side_effect=['[]', '{"completed_indices":[2],"mismatches":[]}']):
            with self.assertLogs('CheckpointManager', level='WARNING'):
                self.assertEqual(manager.load()['completed_indices'],[2])

    def handler(self, service):
        handler=WebApiHandler.__new__(WebApiHandler)
        handler.project_root=Path('.')
        handler.queue_service=service
        handler.send_json=Mock()
        handler.orchestrator=SimpleNamespace(lock=threading.RLock(),is_running=False,
                                             start_queue=Mock(return_value=True))
        handler._validate_input_sheet=Mock(return_value=({'lookup_col':'A','amount_col':'B','has_headers':True,'collector_col':''},'Sheet',10))
        return handler

    def test_api_resume_retains_price_and_job_with_legacy_optional_mapping(self):
        service=memory_queue()
        job=service.add_job(Path('sample.xlsx'),0,'Sheet',total_records=10,
            column_mapping={'lookup_col':'A','amount_col':'B','has_headers':True})
        job.completed=4; job.status='failed'
        handler=self.handler(service)
        handler.handle_queue_add({'workbook':'sample.xlsx','recovery_job_id':job.id,
            'recovery_choice':'resume','amount_target':job.amount_target,'mode':job.mode,
            'column_mapping':job.column_mapping,'pricing':{'mode':'customer','unit_price':'999'}})
        self.assertEqual(handler.send_json.call_args.args[0],200)
        self.assertTrue(handler.send_json.call_args.args[1]['resumed'])
        self.assertEqual(len(service._jobs),1)
        self.assertEqual(job.pricing['unit_minor'],2000)
        self.assertEqual(job.completed,4)

    def test_api_fresh_round_archives_previous_and_preserves_its_progress(self):
        service=memory_queue(); job=service.add_job(Path('sample.xlsx'),0,'Sheet',total_records=10)
        job.completed=4
        handler=self.handler(service)
        handler.handle_queue_add({'workbook':'sample.xlsx','recovery_job_id':job.id,
            'recovery_choice':'new'})
        self.assertEqual(handler.send_json.call_args.args[0],200)
        self.assertEqual(job.status,'archived')
        self.assertEqual(job.completed,4)
        self.assertEqual(len(service._jobs),2)
        self.assertNotEqual(service._jobs[0].result_file,service._jobs[1].result_file)

    def test_final_export_failure_is_recoverable_and_not_completed(self):
        service=memory_queue(); job=service.add_job(Path('sample.xlsx'),0,'Sheet',total_records=1)
        job.completed=1; job.status='active'
        obj=Orchestrator(Mock(),service,Path('.')); obj.current_job_id=job.id
        obj.is_finalizing=True; obj.all_completed_records={0:self.record(customer())}
        with patch('zain_checker.executive_reporter.export_executive_workbook',side_effect=OSError('disk full')):
            with self.assertLogs('Orchestrator',level='ERROR'):
                obj._finalize_job()
        self.assertFalse(obj.result_ready)
        self.assertFalse(obj.is_finalizing)
        self.assertEqual(job.status,'failed')
        self.assertEqual(job.completed,1)

    def test_specific_failed_job_is_validated_and_resumed(self):
        service=memory_queue(); job=service.add_job(Path('sample.xlsx'),0,'Sheet',total_records=10)
        job.completed=4; job.status='failed'
        handler=self.handler(service)
        handler.handle_queue_start({'job_id':job.id})
        handler.orchestrator.start_queue.assert_called_once_with(job_id=job.id)
        self.assertEqual(job.status,'pending')

    def test_broken_journal_line_does_not_discard_valid_results(self):
        c=customer(); text=json.dumps(self.record(c))+'\n{broken\n'+json.dumps(self.record(c,status='mismatch'))+'\n'
        path=Mock();path.exists.return_value=True
        path.open.return_value.__enter__=Mock(return_value=text.splitlines())
        path.open.return_value.__exit__=Mock(return_value=False)
        self.assertEqual(read_journal(path)[c.row_number]['status'],'mismatch')

    def test_new_round_retains_old_progress_and_has_distinct_paths(self):
        service=memory_queue(); source=service.add_job(Path('sample.xlsx'),0,'Sheet',total_records=10)
        source.completed=4
        fresh=service.fork_job(source.id)
        self.assertEqual(source.completed,4)
        self.assertEqual(source.status,'archived')
        self.assertEqual(fresh.completed,0)
        self.assertNotEqual(source.checkpoint_file,fresh.checkpoint_file)
        self.assertNotEqual(source.result_file,fresh.result_file)
        self.assertEqual(service.get_next_pending_job().id,fresh.id)

    def test_specific_resume_claim_does_not_restart_another_sheet(self):
        service=memory_queue()
        first=service.add_job(Path('one.xlsx'),0,'Sheet',total_records=10)
        second=service.add_job(Path('two.xlsx'),0,'Sheet',total_records=10)
        claimed=service.claim_next_pending_job(job_id=second.id)
        self.assertEqual(claimed.id,second.id)
        self.assertEqual(first.status,'pending')

    def test_stale_zero_queue_count_never_forces_checkpoint_reset(self):
        service=memory_queue(); job=service.add_job(Path('sample.xlsx'),0,'Sheet',total_records=10)
        obj=Orchestrator(Mock(),service,Path('.'))
        obj.start_session_from_config=Mock(return_value=True)
        self.assertTrue(obj.start_queue())
        self.assertFalse(obj.start_session_from_config.call_args.kwargs['force_restart'])

    def test_resume_keeps_original_start_and_durable_results(self):
        service=memory_queue(); job=service.add_job(Path('sample.xlsx'),0,'Sheet',total_records=2)
        supervisor=Mock();supervisor.get_all_workers.return_value=[]
        obj=Orchestrator(supervisor,service,Path('.'))
        customers=[customer(0),customer(1)]
        mapping={'lookup_col':'A','amount_col':'B'}
        signature=input_signature(customers,mapping,'smart_hybrid','contract',0)
        saved={'completed_indices':[0], 'input_signature':signature,'started_at':1234}
        sheet=Workbook();sheet.active.append(['رقم العقد','مبلغ العقد']);sheet.active.append(['10000',10])
        checkpoint=Mock();checkpoint.load.return_value=saved
        with patch('domain.workbook.load_fast_workbook',return_value=sheet), \
             patch('domain.workbook.extract_session_customers',return_value=customers) as extract, \
             patch('manager.orchestrator.CheckpointManager',return_value=checkpoint), \
             patch('domain.workbook.inspect_sheet_schema',return_value={'letters':{},'indices':{}}), \
             patch('manager.orchestrator.ensure_audit_workbook_exists'), \
             patch('manager.orchestrator.read_journal',return_value={customers[0].row_number:self.record(customers[0],input_signature=signature)}), \
             patch('zain_checker.async_pipeline_engine.AsyncExcelWriterThread'):
            self.assertTrue(obj.start_session_from_config(Path('sample.xlsx'),0,mapping,job_id=job.id,result_file=job.result_file))
        self.assertEqual(extract.call_args.args[1],mapping)
        self.assertEqual(obj.input_signature,signature)
        self.assertEqual(obj.session_started_at,1234)
        self.assertEqual(obj.completed_indices,{0})
        self.assertEqual(set(obj.all_completed_records),{0})
        checkpoint.reset.assert_not_called()

    def actors(self, customers):
        actors={name:WorkerActor(name,name,Path('.'),Path('.'),use_proxy=name=='proxy',
            proxy_url='socks5://proxy.invalid:1080' if name=='proxy' else None) for name in ('one','two','proxy')}
        for actor in actors.values():actor.status='ready'
        obj=Orchestrator(SimpleNamespace(get_worker=actors.get),Mock(),Path('.'))
        obj.is_running=True;obj.customers=customers;obj.result_queue=queue.Queue();obj.checkpoint_manager=Mock()
        return obj,actors

    def test_challenge_does_not_stop_contracts_on_same_ip_or_service_on_proxy(self):
        obj,actors=self.actors([customer(0,True),customer(1,True),customer(2)])
        task=obj.lease_next_task_for_worker('one')
        obj.notify_pending_verification(task['task_id'],'verify','one')
        self.assertFalse(obj.is_paused)
        self.assertEqual(obj.lease_next_task_for_worker('two')['search_number'][0],'1')
        self.assertEqual(obj.lease_next_task_for_worker('proxy')['search_number'][0],'2')
        self.assertFalse(obj.resume_service_route('one'),'Live verification cannot be bypassed')

    def test_disk_failure_cannot_mark_row_completed_or_bill_it(self):
        obj,actors=self.actors([customer()]);obj.wal_path=Path('UNUSED_WAL')
        task=obj.lease_next_task_for_worker('one')
        with patch('manager.orchestrator.append_journal',side_effect=OSError('disk full')):
            obj.record_task_outcome(task['task_id'],10,'match',worker_id='one')
        self.assertFalse(obj.completed_indices)
        self.assertEqual(obj.matches_count,0)
        self.assertTrue(obj.is_paused)
        self.assertTrue(obj.result_queue.empty())

    def test_slow_browser_start_never_holds_registry_lock_for_other_route(self):
        entered=threading.Event();release=threading.Event()
        class Ready:
            def __init__(self, slow):self.slow=slow
            def wait(self, timeout):
                if self.slow:entered.set();release.wait(2)
                return True
        def factory(proxy_url=None, **options):
            return SimpleNamespace(start=lambda:None,is_alive=lambda:True,is_ready_event=Ready(not proxy_url))
        with patch.object(browser,'_GLOBAL_STEALTH_WORKER',None), \
             patch.object(browser,'_PROXY_BROWSER_WORKERS',{}),patch.object(browser,'StealthServiceWorker',side_effect=factory):
            thread=threading.Thread(target=browser.get_stealth_service_worker);thread.start()
            self.assertTrue(entered.wait(1))
            start=time.monotonic()
            browser.get_stealth_service_worker('socks5://proxy.invalid:1080')
            elapsed=time.monotonic()-start
            release.set();thread.join(2)
            self.assertLess(elapsed,.5)

    def test_mawarid_export_omits_time_and_keeps_collector_amounts_and_formulas(self):
        record={'row':2,'name':'Client','collector':'Collector','account':'10001',
                'expected_sar':10,'live_sar':5,'diff_sar':5,'status':'mismatch','timestamp':'TEST_TIMESTAMP'}
        reporter=ExactTemplateReporter([record],Path('UNUSED.xlsx'),mawarid_format=True)
        wb=Workbook()
        reporter._populate_master_sheet(wb.active,[record],'Total','rows')
        self.assertEqual(wb.active.max_column,12)
        self.assertEqual(wb.active.cell(2,4).value,'Collector')
        self.assertEqual(wb.active.cell(2,6).value,10)
        self.assertEqual(wb.active.cell(3,6).value,'=SUM(F2:F2)')
        diffs=wb.create_sheet('diffs');reporter._populate_diff_sheet(diffs,[record],'Total','rows')
        self.assertEqual(diffs.max_column,11)
        self.assertFalse(any(c.value=='TEST_TIMESTAMP' for sheet in wb for row in sheet for c in row))


if __name__=='__main__':unittest.main()
