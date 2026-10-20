"""User deliverables hide diagnostics without changing recorded outcomes."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from domain.check_messages import public_check_details
from domain.models import CheckError, CheckResult
from domain.results_exporter import append_error_record
from zain_checker.executive_reporter import export_executive_workbook


class PublicCheckMessages(unittest.TestCase):
    def test_public_result_preserves_internal_diagnostic_and_unverified_amount(self):
        for raw, public in (
            ('رفض الطلب (HTTP 200)', 'خطأ نظام 403'),
            ('تحتاج جلسة الويب إلى تجديد', 'خطأ نظام 40'),
            ('تعذر استكمال المتصفح: TargetClosedError', 'خطأ نظام 503'),
        ):
            with self.subTest(raw=raw):
                result = CheckResult(2, 'wallet', '20001', 'Test', 10, None, 'needs_review', details=raw)
                exposed = result.to_dict()
                self.assertEqual(exposed['details'], public)
                self.assertEqual(exposed['status'], 'needs_review')
                self.assertIsNone(exposed['live_amount'])
                self.assertEqual(result.details, raw)

    def test_real_not_found_and_customer_stop_stay_understandable(self):
        self.assertEqual(public_check_details('not_found', 'رقم العقد لا يطابق أي سجل (EZB5029)'),
                         'رقم العقد لا يطابق أي سجل')
        self.assertEqual(public_check_details('cancelled', 'لم يُفحص بطلب العميل'), 'لم يُفحص بطلب العميل')

    def test_workbook_round_trip_hides_technical_errors_and_retains_financial_results(self):
        records = [
            {'row': 2, 'account': '10001', 'expected_sar': 10, 'live_sar': None,
             'status': 'needs_review', 'error': 'رفض الطلب (HTTP 200)'},
            {'row': 3, 'account': '10002', 'expected_sar': 20, 'live_sar': None,
             'status': 'needs_review', 'error': 'تحتاج جلسة الويب إلى تجديد'},
            {'row': 4, 'account': '10003', 'expected_sar': 30, 'live_sar': None,
             'status': 'not_found', 'error': 'رقم العقد لا يطابق أي سجل (EZB5029)'},
            {'row': 5, 'account': '10004', 'expected_sar': 40, 'live_sar': 15,
             'status': 'mismatch', 'diff_sar': -25},
        ]
        original = deepcopy(records)
        with tempfile.TemporaryDirectory() as root, patch('services.file_client.remote_files_enabled', return_value=False):
            path = Path(root) / 'report.xlsx'
            export_executive_workbook(records, path)
            wb = load_workbook(path)
            try:
                values = [cell.value for ws in wb for row in ws for cell in row]
                self.assertIn('خطأ نظام 403', values)
                self.assertIn('خطأ نظام 40', values)
                self.assertIn('رقم العقد لا يطابق أي سجل', values)
                self.assertNotIn('رفض الطلب (HTTP 200)', values)
                self.assertNotIn('تحتاج جلسة الويب إلى تجديد', values)
                positive = wb['الفروقات الايجابيه']
                self.assertEqual(positive.cell(2, 6).value, 40)
                self.assertEqual(positive.cell(2, 7).value, 15)
                self.assertEqual(positive.cell(2, 8).value, '=F2-G2')
            finally:
                wb.close()
            audit = Path(root) / 'audit.xlsx'
            append_error_record(audit, CheckError('20001', 'wallet', 'session_expired',
                'تحتاج جلسة الويب إلى تجديد', [2], expected_amount=1000))
            wb = load_workbook(audit)
            try:
                ws = wb['المهلات والأخطاء']
                self.assertEqual(ws.cell(2, 6).value, '40')
                self.assertEqual(ws.cell(2, 7).value, 'خطأ نظام 40')
            finally:
                wb.close()
        self.assertEqual(records, original)


if __name__ == '__main__':
    unittest.main()
