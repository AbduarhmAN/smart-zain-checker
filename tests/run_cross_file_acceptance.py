"""Real 20-workbook/checkpoint/WAL/export acceptance with synthetic local lookups.

The scheduler, Excel parsing, leasing, result commits and exports are real.
Zain lookup responses are controlled fixtures, never claimed as live checks.
"""
from datetime import datetime, timezone
import hashlib
import json
import logging
from pathlib import Path
import sys
import time
from types import SimpleNamespace
from uuid import uuid4
from unittest.mock import Mock

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT/'tests'))
from retained_runtime import install
retained = install()
from openpyxl import Workbook, load_workbook
from manager.event_bus import EVENT_BUS
from manager.queue import QueueService
from manager.scheduler import QueueScheduler
from workers.actor import WorkerActor

logging.getLogger().setLevel(logging.ERROR)

def emit(stage, **values):
    print(json.dumps({'stage': stage, **values}), flush=True)

def supervisor(root):
    actors = {}
    for scope, count in [('account', 50), ('service', 3)]:
        for i in range(count):
            key = scope+'-'+str(i)
            actor = WorkerActor(key,key,root/key,root,task_scope=scope,use_proxy=scope=='service',
                proxy_url='http://acceptance-route-'+str(i)+'.invalid:8080' if scope=='service' else None)
            actor.status = 'ready'
            actors[key] = actor
    return SimpleNamespace(get_worker=actors.get,get_all_workers=lambda:list(actors.values()),
        get_status_summary=lambda:[a.to_dict() for a in actors.values()],start_all=Mock(),stop_all=Mock()),actors

def wait_for(predicate, seconds=120):
    deadline=time.monotonic()+seconds
    while not predicate():
        if time.monotonic()>deadline:raise AssertionError('Publication did not reach the expected state')
        time.sleep(.05)

def main():
    source_paths=[p for folder in ('manager','domain','services','workers','server','telegram_bot','zain_checker')
                  for p in (ROOT/folder).rglob('*.py')]
    source_paths.append(ROOT/'app.py')
    source_hashes={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in source_paths}
    root=ROOT/'retained-test-runs'/('cross-file-'+uuid4().hex[:10])
    root.mkdir(parents=True,exist_ok=False)
    jobs=QueueService(root/'.zain-queue.json')
    expected={};ids=[];balances={};sizes=[(1000,1)]*10+[(1,100)]*10
    mapping={'lookup_col':'A','service_col':'B','amount_col':'C','customer_col':'D','has_headers':True}
    events=[]
    EVENT_BUS.subscribe('session_completed',lambda _,payload:events.append(dict(payload)))
    for index,(accounts,services) in enumerate(sizes,1):
        file=root/('fixture-'+str(index)+'.xlsx');book=Workbook();sheet=book.active
        sheet.append(['رقم الحساب','رقم الخدمة','المبلغ','اسم العميل'])
        names=[]
        for row in range(accounts+services):
            service=row>=accounts;number=('2' if service else '1')+str(index).zfill(4)+str(row).zfill(6)
            name='LOCAL-FILE-'+str(index)+'-ROW-'+str(row+2)
            names.append(name);balances[number]=str(index)
            sheet.append(['' if service else number,number if service else '',index,name])
        book.save(file);book.close()
        job=jobs.add_job(file,0,'Sheet',column_mapping={**mapping,'telegram_chat_id':990000+index},total_records=accounts+services)
        ids.append(job.id);expected[job.id]={'names':set(names),'count':accounts+services,'index':index}
    pool,actors=supervisor(root)
    scheduler=QueueScheduler(pool,jobs,root,start_workers=False)
    for index,key in enumerate(ids,1):
        assert scheduler.start_queue(key)
        if index%5==0:emit('real_workbooks_admitted',files=index)
    assert len({s.engine.wal_path for s in scheduler._sessions.values()})==20
    assert len({s.engine.checkpoint_manager.db_path for s in scheduler._sessions.values()})==20
    processed=set();old_task=None
    # Accounts independently reach file 20 with every service still unfinished.
    for worker_type,target in [('account',10010),('service',1009)]:
        finished=0
        while finished<target:
            progress=False
            for worker,actor in actors.items():
                if actor.task_scope!=worker_type:continue
                task=scheduler.lease_next_task_for_worker(worker)
                if task is None:continue
                if worker_type=='service' and task['job_id']==ids[14] and old_task is None:
                    old_task=task
                    scheduler.defer_task_to_end(task['task_id'],'local simulated retry',worker_id=worker,
                        retry_after=1000,category='network_error')
                    progress=True
                    continue
                identity=(task['job_id'],task['row_number'])
                assert identity not in processed
                processed.add(identity)
                scheduler.record_task_outcome(task['task_id'],balances[task['search_number']],'match',worker_id=worker)
                finished+=1;progress=True
            assert progress,'Available fixture work unexpectedly stalled'
            if finished and finished%2000<53:emit('durable_rows_committed',kind=worker_type,rows=finished)
        emit('type_phase_complete',kind=worker_type,rows=finished)
    assert jobs.get_job(ids[18]).completed==101
    assert jobs.get_job(ids[14]).completed==100
    scheduler.maintain_sessions()
    wait_for(lambda:jobs.get_job(ids[18]).status=='completed')
    assert jobs.get_job(ids[14]).status=='active'
    emit('file19_published_while_file15_waits',file19_completed=True,file15_remaining=1)
    wait_for(lambda:all(jobs.get_job(key).status=='completed' for key in ids if key!=ids[14]))
    remaining=scheduler._sessions[ids[14]].engine
    assert len(remaining.completed_indices)==100 and len(remaining.deferred_indices)==1
    # Graceful process replacement retains durable state; no previous row is queried again.
    scheduler.cancel_session()
    recovered_queue=QueueService(root/'.zain-queue.json')
    recovered_pool,recovered_actors=supervisor(root)
    recovered=QueueScheduler(recovered_pool,recovered_queue,root,start_workers=False)
    assert recovered.start_queue(ids[14])
    engine=recovered._sessions[ids[14]].engine
    assert len(engine.completed_indices)==100
    recovered.record_task_outcome(old_task['task_id'],'999','match',worker_id='service-0')
    assert len(engine.completed_indices)==100
    engine.deferred_retry_at[engine.deferred_indices[0]]=0
    task=recovered.lease_next_task_for_worker('service-0')
    assert task is not None and task['job_id']==ids[14]
    assert (task['job_id'],task['row_number']) not in processed
    recovered.record_task_outcome(task['task_id'],balances[task['search_number']],'match',worker_id='service-0')
    processed.add((task['job_id'],task['row_number']))
    recovered.maintain_sessions()
    wait_for(lambda:recovered_queue.get_job(ids[14]).status=='completed')
    reports=[]
    for key in ids:
        job=recovered_queue.get_job(key);assert job.status=='completed'
        path=root/job.result_file;book=load_workbook(path,read_only=True,data_only=False)
        cells={str(value) for sheet in book for row in sheet.iter_rows(values_only=True) for value in row if value is not None}
        other_names=set().union(*(v['names'] for k,v in expected.items() if k!=key))
        assert not (cells&other_names)
        book.close()
        # The executive report intentionally lists differences/errors only.
        # The updated master contains all rows, including matching balances.
        master=load_workbook(root/('fixture-'+str(expected[key]['index'])+'_محدث_بالفحص.xlsx'),read_only=True)
        master_cells={str(value) for row in master.worksheets[0].iter_rows(values_only=True) for value in row if value is not None}
        assert expected[key]['names']<=master_cells
        assert not (master_cells&other_names)
        master.close()
        reports.append({'file_number':expected[key]['index'],'rows':expected[key]['count'],'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    assert len(processed)==11020
    assert len(events)==20
    assert {p['job_id'] for p in events}==set(ids)
    assert pool.stop_all.call_count==1  # Explicit process replacement only.
    assert recovered_pool.stop_all.call_count==0
    assert source_hashes=={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in source_paths},'Source changed during acceptance'
    proof={'at':datetime.now(timezone.utc).isoformat(),'passed':True,'production':False,
        'source':'real scheduler/Excel/checkpoints/WAL/export; controlled local lookup responses',
        'file_count':20,'account_checks':10010,'service_checks':1010,'unique_completed_rows':len(processed),
        'file19_published_before_file15':True,'restored_completed_rows':100,'stale_result_ignored':True,
        'reports_owner_isolated':True,'completion_events':len(events),'reports':reports,'source_hashes':source_hashes,
        'builds_run':False,'files_deleted':False,'external_checks_performed':False,'cleanup_calls_retained':len(retained)}
    (root/'acceptance.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
    folder=ROOT/'retained-test-evidence';folder.mkdir(exist_ok=True)
    (folder/'cross_file_acceptance.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
    emit('acceptance_passed',files=20,completed=11020,evidence=str(root/'acceptance.json'))

if __name__=='__main__':main()
