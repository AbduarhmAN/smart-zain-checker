"""Offline Excel ingestion tests: memory workbooks, no workers or network."""
import unittest
import threading
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook
from domain.workbook import inspect_sheet_schema, extract_customer_records
from domain.workbook_validation import validate_sheet
from manager.orchestrator import Orchestrator
from server.web_api import WebApiHandler
from zain_checker.sheet_updater import update_raw_portfolio_sheet


def workbook(rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    return wb


class MinimalSheetTests(unittest.TestCase):
    def test_reported_379593_rows_with_337039_empty_helper_links(self):
        class TemplateSheet:
            title = 'Sheet'
            max_column = 4
            max_row = 379594

            def iter_rows(self, min_row=1, max_row=None, **kwargs):
                last = min(max_row or self.max_row, self.max_row)
                for row in range(min_row, last+1):
                    if row == 1:
                        yield ('اسم المحصل','رقم العقد','مبلغ العقد','رابط')
                    elif row <= 42555:
                        yield ('test','1000000001',0,None)
                    else:
                        yield (None,None,None,'https://zain.app/ar/contract-payment?contract=')

        sheet = TemplateSheet()
        analysis = inspect_sheet_schema(sheet)
        result = validate_sheet(sheet, analysis['letters'])
        self.assertEqual(analysis['estimated_rows'], 42554)
        self.assertTrue(result['ok'])
        self.assertEqual(result['rows'], 42554)
        self.assertEqual(result['invalid_rows'], 0)

    def test_formatted_tail_rows_are_not_customer_records(self):
        wb = workbook([['Number','Amount'],['1000000001',0],['2000000001',25]])
        wb.active.cell(1000, 1).number_format = '@'
        self.assertEqual(wb.active.max_row, 1000)
        analysis = inspect_sheet_schema(wb.active)
        self.assertEqual(analysis['estimated_rows'], 2)
        self.assertEqual(validate_sheet(wb.active, analysis['letters'])['rows'], 2)

    def test_empty_zain_links_are_not_customer_records(self):
        wb = workbook([['اسم المحصل','رقم العقد','مبلغ العقد','رابط'],
                       ['test','1000000001',0,'https://zain.app/ar/contract-payment?contract=1000000001'],
                       [None,None,None,'https://zain.app/ar/contract-payment?contract='],
                       [None,None,None,'https://app.sa.zain.com/ar/quickpay?account=']])
        analysis = inspect_sheet_schema(wb.active)
        result = validate_sheet(wb.active, analysis['letters'])
        self.assertEqual(analysis['estimated_rows'], 1)
        self.assertTrue(result['ok'])
        self.assertEqual(result['rows'], 1)
        self.assertEqual(len(extract_customer_records(wb.active, 2, 3, service_col=2)), 1)

    def test_incomplete_customer_rows_still_fail_validation(self):
        for row in ([None,'1000000001',None,None], [None,None,0,None],
                    ['customer',None,None,'https://zain.app/ar/contract-payment?contract='],
                    [None,None,None,'https://zain.app/ar/contract-payment?contract=1000000001']):
            wb = workbook([['اسم المحصل','رقم العقد','مبلغ العقد','رابط'], row])
            schema = inspect_sheet_schema(wb.active)
            result = validate_sheet(wb.active, schema['letters'])
            self.assertEqual(schema['estimated_rows'], 1)
            self.assertFalse(result['ok'])
            self.assertEqual(result['invalid_rows'], 1)

    def test_selected_link_is_validated_instead_of_ignored(self):
        wb = workbook([['Number','Amount'], ['https://zain.app/ar/contract-payment?contract=',None]])
        result = validate_sheet(wb.active, {'lookup_col':'A','amount_col':'B'})
        self.assertEqual(result['rows'], 1)
        self.assertFalse(result['ok'])

    def test_generic_headers_mixed_numbers_and_zero_amount(self):
        wb = workbook([['Number', 'Amount'], ['1000000001', 0], ['2000000001', 125.50]])
        analysis = inspect_sheet_schema(wb.active)
        self.assertEqual(analysis['letters']['lookup_col'], 'A')
        self.assertEqual(analysis['letters']['amount_col'], 'B')
        self.assertTrue(validate_sheet(wb.active, analysis['letters'])['ok'])
        customers = extract_customer_records(wb.active, 1, 2, service_col=1)
        self.assertEqual([c.expected_amount for c in customers], [0, 12550])
        self.assertEqual([Orchestrator._search_number(c) for c in customers], ['1000000001', '2000000001'])

    def test_unknown_headers_infer_only_unique_candidates(self):
        wb = workbook([['X', 'Y'], ['1000000001', 0], ['2000000001', 95]])
        self.assertEqual(inspect_sheet_schema(wb.active)['letters']['lookup_col'], 'A')
        ambiguous = workbook([['X', 'Y', 'Z'], ['1000000001', '2000000001', 0]])
        analysis = inspect_sheet_schema(ambiguous.active)
        self.assertIsNone(analysis['letters']['lookup_col'])
        self.assertEqual(len(analysis['columns']), 3)
        self.assertFalse(analysis['is_acceptable'])
        self.assertTrue(validate_sheet(ambiguous.active, {'lookup_col':'B', 'amount_col':'C'})['ok'])

    def test_headerless_preview_validation_and_extraction_keep_first_customer(self):
        wb = workbook([['1000000001', 0], ['2000000001', 125.50]])
        analysis = inspect_sheet_schema(wb.active, has_headers=False)
        self.assertEqual(analysis['estimated_rows'], 2)
        self.assertEqual(analysis['preview_rows'][0]['row_number'], 1)
        self.assertEqual(analysis['preview_rows'][0]['values'], ['1000000001', '0'])
        self.assertEqual([c['header'] for c in analysis['columns']], ['', ''])
        mapping = {**analysis['letters'], 'has_headers':False}
        result = validate_sheet(wb.active, mapping)
        self.assertTrue(result['ok'])
        self.assertEqual(result['rows'], 2)
        customers = extract_customer_records(wb.active, 1, 2, service_col=1, has_headers=False)
        self.assertEqual([c.row_number for c in customers], [1, 2])

    def test_row_specific_errors_and_wrong_mode(self):
        wb = workbook([['Number', 'Amount'], ['2000000001', None], ['bad', 'not money']])
        result = validate_sheet(wb.active, {'lookup_col':'A', 'service_col':'A', 'amount_col':'B'})
        self.assertFalse(result['ok'])
        self.assertEqual([error['row'] for error in result['row_errors']], [2, 3])
        self.assertIn('فارغ', result['row_errors'][0]['reasons'][0])
        self.assertEqual(len(result['row_errors'][1]['reasons']), 2)
        wb.active['B2'] = 0
        result = validate_sheet(wb.active, {'lookup_col':'A', 'amount_col':'B'}, 'account_only')
        self.assertIn('العقود والخدمات', result['row_errors'][0]['reasons'][0])

    def test_server_preserves_headerless_mapping_and_count(self):
        wb = workbook([['1000000001', 0], ['2000000001', 125.50]])
        handler = object.__new__(WebApiHandler)
        handler.send_json = Mock()
        with patch('openpyxl.load_workbook', return_value=wb):
            mapped, name, count = handler._validate_input_sheet(Path('fixture.xlsx'), 0,
                {'lookup_col':'A', 'amount_col':'B', 'has_headers':False}, 'smart_hybrid', 'remaining')
        self.assertFalse(mapped['has_headers'])
        self.assertEqual(count, 2)
        handler.send_json.assert_not_called()

    def test_results_do_not_replace_first_headerless_customer(self):
        wb = workbook([['1000000001', 0], ['2000000001', 125.50]])
        wb.save = Mock()
        records = [{'row':1, 'account':'1000000001', 'status':'match', 'live_sar':0, 'diff_sar':0},
                   {'row':2, 'service':'2000000001', 'status':'match', 'live_sar':125.50, 'diff_sar':0}]
        with patch('pathlib.Path.exists', return_value=True), patch('zain_checker.sheet_updater.load_workbook', return_value=wb):
            ok, _, _ = update_raw_portfolio_sheet(Path('fixture.xlsx'), records, has_headers=False)
        self.assertTrue(ok)
        self.assertEqual(wb.active['A1'].value, '1000000001')
        self.assertEqual(wb.active['B1'].value, 0)
        self.assertEqual(wb.active['C1'].value, 0)
        self.assertEqual(wb.active['C2'].value, 125.50)
        self.assertTrue(wb.save.called)

    def test_worker_loader_receives_headerless_flag(self):
        wb = workbook([['1000000001', 0], ['2000000001', 125.50]])
        orchestrator = object.__new__(Orchestrator)
        orchestrator.lock = threading.RLock()
        orchestrator.is_running = False
        with patch('openpyxl.load_workbook', return_value=wb), patch('manager.orchestrator.extract_customer_records', return_value=[]) as extract:
            self.assertFalse(orchestrator.start_session_from_config(Path('fixture.xlsx'), 0,
                {'lookup_col':'A', 'service_col':'A', 'amount_col':'B', 'has_headers':False}, amount_target='remaining'))
        self.assertFalse(extract.call_args.kwargs['has_headers'])
        self.assertFalse(orchestrator.source_has_headers)


if __name__ == '__main__':
    unittest.main()
