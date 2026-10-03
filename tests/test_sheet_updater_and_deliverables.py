"""Tests for SheetUpdater and Executive Deliverables in Smart Zain Checker.
Verifies that:
1. Master portfolios and error sheets are updated with live amounts, differences, and statuses.
2. Executive deliverables contain comprehensive master archives and never show empty sheets ('صفر، لا شيء').
3. Error maintenance sessions generate dedicated results workbooks with 100% of repaired records.
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl
import pytest

from domain.models import Customer
from zain_checker.executive_reporter import export_executive_workbook, ExactTemplateReporter
from zain_checker.sheet_updater import (
    update_errors_or_results_sheet,
    update_raw_portfolio_sheet,
    update_source_workbook_after_job,
)


def test_executive_reporter_all_records_tab(tmp_path: Path):
    """Test that all checked records are included in Tab 1, even if 100% match with 0 diff."""
    records = [
        {
            "row": 10,
            "name": "عميل مطابق 1",
            "account": "1001111111",
            "service": "2001111111",
            "contract": "1001111111",
            "expected_sar": 150.0,
            "live_sar": 150.0,
            "diff_sar": 0.0,
            "status": "match",
            "status_label": "تطابق تام",
            "main_status": "وعد سداد",
            "sub_status": "متابعة",
            "follow_notes": "لا توجد ملاحظات",
            "timestamp": "2026-10-03 01:00:00",
        },
        {
            "row": 20,
            "name": "عميل مطابق 2",
            "account": "1002222222",
            "service": "2002222222",
            "contract": "1002222222",
            "expected_sar": 300.0,
            "live_sar": 300.0,
            "diff_sar": 0.0,
            "status": "match",
            "status_label": "تطابق تام",
            "main_status": "متابعة",
            "sub_status": "لا يرد",
            "follow_notes": "تم الاتصال",
            "timestamp": "2026-10-03 01:05:00",
        },
    ]

    out_file = tmp_path / "نتائج_فحص_مطابقة.xlsx"
    export_executive_workbook(records, out_file, is_repair=False)

    assert out_file.exists()
    wb = openpyxl.load_workbook(out_file, data_only=True)
    assert "جميع السجلات المفحوصة" in wb.sheetnames

    ws1 = wb["جميع السجلات المفحوصة"]
    # 1 header + 2 data rows + 1 summary row = 4 rows
    assert ws1.max_row == 4
    assert ws1.cell(2, 4).value == "عميل مطابق 1"
    assert ws1.cell(2, 6).value == 150.0
    assert ws1.cell(2, 7).value == 150.0
    assert ws1.cell(2, 9).value == "تطابق تام"
    assert ws1.cell(3, 4).value == "عميل مطابق 2"


def test_executive_reporter_repair_mode(tmp_path: Path):
    """Test that maintenance sessions generate 'نتائج صيانة وتدقيق الأخطاء' with all rows."""
    repaired_records = [
        {
            "row": 54,
            "name": "احمد الشايع",
            "account": "1006659013",
            "service": "2006695911",
            "contract": "1006659013",
            "expected_sar": 143.33,
            "live_sar": 143.33,
            "diff_sar": 0.0,
            "status": "match",
            "status_label": "تطابق تام",
            "main_status": "عدم توصل",
            "sub_status": "لا يرد",
            "follow_notes": "تم إصلاح الخطأ",
            "timestamp": "2026-10-03 02:17:29",
        }
    ]

    out_file = tmp_path / "نتائج_صيانة_شيت.xlsx"
    export_executive_workbook(repaired_records, out_file, is_repair=True)

    assert out_file.exists()
    wb = openpyxl.load_workbook(out_file, data_only=True)
    assert "نتائج صيانة وتدقيق الأخطاء" in wb.sheetnames
    ws1 = wb["نتائج صيانة وتدقيق الأخطاء"]
    assert ws1.max_row == 3  # Header + 1 record + 1 summary
    assert ws1.cell(2, 4).value == "احمد الشايع"
    assert ws1.cell(2, 7).value == 143.33
    assert ws1.cell(2, 9).value == "تطابق تام"


def test_update_errors_sheet_in_place(tmp_path: Path):
    """Test updating an existing errors sheet with newly verified Zain data."""
    # Create sample error workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "الأخطاء والملاحظات للمحصل"
    ws.append([
        "رقم الصف في الشيت",
        "اسم العميل",
        "رقم الحساب ورقم الخدمة",
        "المبلغ في الشيت",
        "المشكلة باختصار",
        "رابط صفحة زين",
        "الحالة الرئيسية بالملف",
        "الحالة الفرعية بالملف",
    ])
    ws.append([
        "54",
        "احمد الشايع",
        "حساب: 1006659013 | خدمة: 2006695911",
        143.33,
        "تعذر استكمال المتصفح: TargetClosedError",
        "اضغط لفتح حساب العميل بزين",
        "عدم توصل",
        "لا يرد",
    ])
    src_file = tmp_path / "نتائج_سابقة.xlsx"
    wb.save(src_file)
    wb.close()

    records = [{
        "row": "54",
        "name": "احمد الشايع",
        "account": "1006659013",
        "service": "2006695911",
        "expected_sar": 143.33,
        "live_sar": 143.33,
        "diff_sar": 0.0,
        "status": "match",
        "status_label": "تطابق تام",
        "timestamp": "2026-10-03 02:17:29",
    }]

    ok, out_path, msg = update_errors_or_results_sheet(src_file, records)
    assert ok is True

    # Verify updated content in source workbook
    wb_updated = openpyxl.load_workbook(src_file, data_only=True)
    # 1. Matched record placed in 'السجلات المطابقة (تطابق تام)'
    assert "السجلات المطابقة (تطابق تام)" in wb_updated.sheetnames
    ws_match = wb_updated["السجلات المطابقة (تطابق تام)"]
    assert ws_match.cell(2, 4).value == "احمد الشايع"
    assert ws_match.cell(2, 6).value == 143.33  # Expected SAR
    assert ws_match.cell(2, 7).value == 143.33  # Live SAR
    # 2. Resolved error cleared from 'الأخطاء والملاحظات للمحصل'
    ws_up = wb_updated["الأخطاء والملاحظات للمحصل"]
    assert "✓" in str(ws_up.cell(2, 5).value)  # Success banner when all errors resolved
    wb_updated.close()


def test_update_results_sheet_diff_tabs(tmp_path: Path):
    """Test distributing positive and negative differences into exact respective tabs."""
    from zain_checker.executive_reporter import export_executive_workbook
    initial_records = [
        {
            "row": 10,
            "name": "عميل أول",
            "account": "1001111111",
            "service": "2001111111",
            "expected_sar": 200.0,
            "live_sar": 100.0,
            "diff_sar": 100.0,
            "status": "mismatch",
            "status_label": "فرق رصيد",
            "main_status": "متابعة",
            "sub_status": "لا يرد",
            "follow_notes": "ملاحظة",
            "timestamp": "2026-10-03 01:00:00",
        }
    ]
    results_file = tmp_path / "نتائج فحص زين - تجربة.xlsx"
    export_executive_workbook(initial_records, results_file, is_repair=False)

    # Now simulate a maintenance session with:
    # 1. Positive difference (Expected 500, Live 300 -> Diff 200)
    # 2. Negative difference (Expected 100, Live 300 -> Diff -200)
    # 3. Match (Expected 50, Live 50 -> Diff 0)
    new_records = [
        {
            "row": 11,
            "name": "عميل فرق إيجابي",
            "account": "1002222222",
            "service": "2002222222",
            "expected_sar": 500.0,
            "live_sar": 300.0,
            "diff_sar": 200.0,
            "status": "mismatch",
            "status_label": "سداد جزئي",
            "main_status": "متابعة",
            "sub_status": "وعد سداد",
            "follow_notes": "وعد بالسداد",
            "timestamp": "2026-10-03 02:00:00",
        },
        {
            "row": 12,
            "name": "عميل فرق سالب",
            "account": "1003333333",
            "service": "2003333333",
            "expected_sar": 100.0,
            "live_sar": 300.0,
            "diff_sar": -200.0,
            "status": "mismatch",
            "status_label": "زيادة مديونية",
            "main_status": "عدم توصل",
            "sub_status": "مغلق",
            "follow_notes": "مغلق",
            "timestamp": "2026-10-03 02:01:00",
        },
        {
            "row": 13,
            "name": "عميل مطابق تماما",
            "account": "1004444444",
            "service": "2004444444",
            "expected_sar": 50.0,
            "live_sar": 50.0,
            "diff_sar": 0.0,
            "status": "match",
            "status_label": "تطابق تام",
            "main_status": "تم السداد",
            "sub_status": "منتهي",
            "follow_notes": "مسدد",
            "timestamp": "2026-10-03 02:02:00",
        },
    ]

    ok, out_path, msg = update_errors_or_results_sheet(results_file, new_records)
    assert ok is True

    wb = openpyxl.load_workbook(results_file, data_only=True)
    # Check positive diff tab
    assert "الفروقات الايجابيه" in wb.sheetnames
    ws_pos = wb["الفروقات الايجابيه"]
    names_pos = [ws_pos.cell(r, 4).value for r in range(2, ws_pos.max_row)]
    assert "عميل فرق إيجابي" in names_pos

    # Check negative diff tab
    assert "فروقات سالبه" in wb.sheetnames
    ws_neg = wb["فروقات سالبه"]
    names_neg = [ws_neg.cell(r, 4).value for r in range(2, ws_neg.max_row)]
    assert "عميل فرق سالب" in names_neg

    # Check all diffs tab
    assert "فروقات شامله" in wb.sheetnames
    ws_all = wb["فروقات شامله"]
    names_all = [ws_all.cell(r, 4).value for r in range(2, ws_all.max_row)]
    assert "عميل فرق إيجابي" in names_all
    assert "عميل فرق سالب" in names_all

    # Check matched tab
    assert "السجلات المطابقة (تطابق تام)" in wb.sheetnames
    ws_match = wb["السجلات المطابقة (تطابق تام)"]
    names_match = [ws_match.cell(r, 4).value for r in range(2, ws_match.max_row)]
    assert "عميل مطابق تماما" in names_match

    wb.close()


def test_update_raw_portfolio_sheet(tmp_path: Path):
    """Test updating a raw customer portfolio (e.g. 2.xlsx) with audit columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "المحفظة"
    ws.append(["م", "اسم العميل", "رقم الحساب", "مبلغ المديونية"])
    ws.append([1, "سالم محمد", "1009998887", 500.0])
    ws.append([2, "خالد علي", "1007776665", 800.0])
    portfolio_file = tmp_path / "محفظة_عملاء.xlsx"
    wb.save(portfolio_file)
    wb.close()

    records = [
        {
            "row": 2,  # Row 2 in Excel (سالم محمد)
            "name": "سالم محمد",
            "account": "1009998887",
            "expected_sar": 500.0,
            "live_sar": 500.0,
            "diff_sar": 0.0,
            "status": "match",
            "status_label": "تطابق تام",
            "timestamp": "2026-10-03 03:00:00",
        },
        {
            "row": 3,  # Row 3 in Excel (خالد علي)
            "name": "خالد علي",
            "account": "1007776665",
            "expected_sar": 800.0,
            "live_sar": 750.0,
            "diff_sar": 50.0,
            "status": "mismatch",
            "status_label": "فرق رصيد",
            "timestamp": "2026-10-03 03:01:00",
        },
    ]

    ok, out_path, msg = update_raw_portfolio_sheet(portfolio_file, records)
    assert ok is True

    wb_res = openpyxl.load_workbook(portfolio_file, data_only=True)
    ws_res = wb_res["المحفظة"]
    headers = [ws_res.cell(1, c).value for c in range(1, ws_res.max_column + 1)]
    assert "المبلغ بموقع زين (ر.س)" in headers
    assert "فرق المديونية (ر.س)" in headers
    assert "نتيجة فحص زين" in headers

    # Row 2: سالم محمد
    assert ws_res.cell(2, 5).value == 500.0
    assert ws_res.cell(2, 6).value == 0.0
    assert ws_res.cell(2, 7).value == "تطابق تام"

    # Row 3: خالد علي
    assert ws_res.cell(3, 5).value == 750.0
    assert ws_res.cell(3, 6).value == 50.0
    assert ws_res.cell(3, 7).value == "فرق رصيد"

    wb_res.close()
