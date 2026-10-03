"""Offline, in-memory tests. Never contact Zain or Telegram or change a customer file."""
import threading
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from openpyxl import Workbook
from domain.workbook import get_sheet_header_columns
from domain.workbook_validation import validate_sheet, FIELDS
from manager.queue import QueueService
from manager.orchestrator import Orchestrator
from telegram_bot.bot import TelegramBotRunner
from telegram_bot.notifier import TelegramNotifier
from telegram_bot.sheet_workflow import columns_panel, review_panel
from server.web_api import WebApiHandler


def workbook(amount=0):
    wb = Workbook()
    wb.active.append(['اسم', 'مرجع', 'رصيد', None, 'رصيد'])
    wb.active.append(['عميل تجريبي', '2000000001', amount, 'إضافي', 50])
    return wb


class SheetValidationTests(unittest.TestCase):
    def test_all_columns_include_blank_and_duplicate_headers(self):
        columns = get_sheet_header_columns(workbook().active)
        self.assertEqual([c['col_letter'] for c in columns], ['A','B','C','D','E'])
        self.assertEqual(columns[3]['header'], '')

    def test_zero_amount_valid_and_blanks_or_invalid_money_rejected(self):
        mapping = {'service_col':'B','amount_col':'C'}
        self.assertTrue(validate_sheet(workbook(0).active,mapping,'service_only')['ok'])
        for value in (None,'oops',float('inf')):
            self.assertFalse(validate_sheet(workbook(value).active,mapping,'service_only')['ok'])

    def test_mapping_requires_number_amount_and_columns_in_bounds(self):
        for mapping in ({}, {'service_col':'B'}, {'service_col':'F','amount_col':'C'},
                        {'service_col':'C','amount_col':'C'}):
            self.assertFalse(validate_sheet(workbook().active,mapping)['ok'])

    def test_every_populated_row_checked_and_blank_rows_ignored(self):
        wb = workbook()
        wb.active.append([None]*5)
        wb.active.append(['incomplete',None,100])
        result = validate_sheet(wb.active,{'service_col':'B','amount_col':'C'})
        self.assertEqual((result['rows'],result['invalid_rows']),(2,1))
        self.assertIn('4',result['problems'][0])

    def test_mode_and_amount_target_are_checked(self):
        m={'service_col':'B','amount_col':'C'}
        self.assertFalse(validate_sheet(workbook().active,m,'account_only')['ok'])
        self.assertFalse(validate_sheet(workbook().active,m,amount_target='smart_dual')['ok'])
        m['amount_col_2']='E'
        self.assertTrue(validate_sheet(workbook().active,m,amount_target='smart_dual')['ok'])

    def test_exclusive_modes_do_not_validate_a_different_number_column(self):
        wb=workbook();wb.active['D2']='1000000001'
        self.assertFalse(validate_sheet(wb.active,{'lookup_col':'A','service_col':'D','amount_col':'C'},'account_only')['ok'])
        self.assertFalse(validate_sheet(wb.active,{'lookup_col':'B','service_col':'A','amount_col':'C'},'service_only')['ok'])

    def test_loader_respects_cleared_fields_and_exclusive_number_mode(self):
        wb=Workbook();wb.active.append(['رقم الحساب','رقم الخدمة','المبلغ','اسم العميل'])
        wb.active.append(['1000000001','2000000001',100,'test'])
        orchestrator=Orchestrator.__new__(Orchestrator)
        orchestrator.lock=threading.RLock();orchestrator.is_running=False
        mapping={'lookup_col':'A','service_col':'B','amount_col':'C','amount_col_2':None,'customer_col':None}
        for mode,lookup,service in (('account_only',1,None),('service_only',2,2)):
            with patch('openpyxl.load_workbook',return_value=wb), patch('manager.orchestrator.extract_customer_records',return_value=[]) as extract:
                self.assertFalse(orchestrator.start_session_from_config(Path('unused.xlsx'),sheet_index=0,mapping=mapping,mode=mode))
            self.assertEqual(extract.call_args.kwargs['lookup_col'],lookup)
            self.assertEqual(extract.call_args.kwargs['service_col'],service)
            self.assertIsNone(extract.call_args.kwargs['customer_col'])
            self.assertIsNone(extract.call_args.kwargs['amount_col_2'])

    def test_empty_sheet_cannot_be_queued(self):
        wb=Workbook();wb.active.append(['service','amount'])
        self.assertFalse(validate_sheet(wb.active,{'service_col':'A','amount_col':'B'})['ok'])

    def test_phone_number_column_is_not_accepted_as_a_zain_service_number(self):
        wb=workbook();wb.active['B2']='0500000001'
        self.assertFalse(validate_sheet(wb.active,{'service_col':'B','amount_col':'C'},'service_only')['ok'])

    def test_wide_columns_are_paginated_with_valid_callbacks(self):
        wb=Workbook();wb.active.append([None]*80+['آخر عمود'])
        draft={'columns':get_sheet_header_columns(wb.active)}
        panel=columns_panel(draft,'a'*12,'collector_col',10)
        self.assertIn('CC',panel.text)
        self.assertIn('آخر عمود',panel.text)
        self.assertLess(len(panel.text),4096)
        for row in panel.rows:
            for b in row:
                self.assertLessEqual(len(b['callback_data'].encode()),64)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        blocker=patch('urllib.request.urlopen',side_effect=AssertionError('Network forbidden'))
        blocker.start();self.addCleanup(blocker.stop)
        self.state={'workbook':'example.xlsx','running':False,'paused':False,'verifications':[],
                    'kpis':{'errors':2,'needs_review':1,'completed':3,'active_leases':0}}
        self.orchestrator=Mock(lock=threading.RLock(),session_token='session',current_job_id=None,is_running=False)
        self.orchestrator.get_live_status.side_effect=lambda:self.state.copy()
        self.queue=Mock();self.queue.get_jobs.return_value=[]
        self.bot=TelegramBotRunner(self.orchestrator,self.queue,Path.cwd(),bot_token='TEST',authorized_id=123)
        self.bot._api_call=Mock(return_value={'ok':True})

    def callback(self, action):
        self.bot._handle_callback({'id':'q','from':{'id':123},'message':{'chat':{'id':123},'message_id':7},'data':'zc:'+action})

    def open(self, wb=None):
        self.wb=wb or workbook()
        self.patcher=patch('telegram_bot.sheet_workflow.load_workbook',return_value=self.wb)
        self.patcher.start();self.addCleanup(self.patcher.stop)
        stamp=patch('telegram_bot.sheet_workflow.file_version',return_value=(10,20))
        stamp.start();self.addCleanup(stamp.stop)
        self.bot._open_sheet_review(Path.cwd()/'sample.xlsx',123)
        return next(iter(self.bot._sheet_drafts))

    def test_unsuitable_sheet_can_be_mapped_then_explicitly_queued_once(self):
        token=self.open()
        self.callback(f'map:{token}:save')
        self.queue.add_job.assert_not_called()
        self.callback(f'map:{token}:set:service_col:B')
        self.callback(f'map:{token}:set:amount_col:C')
        self.callback(f'map:{token}:mode:service_only')
        self.callback(f'map:{token}:save')
        self.callback(f'map:{token}:save')
        self.queue.add_job.assert_called_once()
        payload=self.queue.add_job.call_args.kwargs
        self.assertEqual(payload['column_mapping']['service_col'],'B')
        self.assertEqual(payload['column_mapping']['amount_col'],'C')
        self.assertEqual(payload['amount_target'],'remaining')
        self.assertEqual(payload['total_records'],1)
        self.assertNotIn('force_clean',payload)

    def test_old_upload_and_wrong_column_cannot_modify_current_draft(self):
        old=self.open()
        self.bot._open_sheet_review(Path.cwd()/'new.xlsx',123)
        current=next(key for key in self.bot._sheet_drafts if key != old)
        original_service_col = self.bot._sheet_drafts[current]['mapping']['service_col']
        self.callback(f'map:{old}:set:service_col:B')
        self.callback(f'map:{current}:set:service_col:ZZZ')
        self.assertEqual(self.bot._sheet_drafts[current]['mapping']['service_col'], original_service_col)
        self.queue.add_job.assert_not_called()

    def test_round_start_requires_valid_files_and_explicit_confirmation(self):
        wb=workbook()
        job={'id':'round_1','filename':'example.xlsx','sheet_index':0,'sheet_name':'Sheet',
             'status':'pending','mode':'service_only','amount_target':'remaining',
             'column_mapping':{'service_col':'B','amount_col':'C'}}
        self.queue.get_jobs.return_value=[job]
        with patch('openpyxl.load_workbook',return_value=wb), patch('telegram_bot.sheet_workflow.file_version',return_value=(10,20)):
            self.callback('start_queue')
            self.orchestrator.start_queue.assert_not_called()
            token=next(iter(self.bot._confirmations))
            self.callback('confirm:'+token);self.callback('confirm:'+token)
        self.orchestrator.start_queue.assert_called_once()

    def test_round_start_rejects_changed_queue_and_unsuitable_files(self):
        job={'id':'round_1','filename':'example.xlsx','sheet_index':0,'sheet_name':'Sheet',
             'status':'pending','mode':'service_only','amount_target':'remaining',
             'column_mapping':{'service_col':'B','amount_col':'C'}}
        self.queue.get_jobs.return_value=[job]
        with patch('openpyxl.load_workbook',return_value=workbook()),patch('telegram_bot.sheet_workflow.file_version',return_value=(10,20)):
            self.callback('start_queue');token=next(iter(self.bot._confirmations))
            job['column_mapping']['customer_col']='A'
            self.callback('confirm:'+token)
        self.orchestrator.start_queue.assert_not_called()
        with patch('openpyxl.load_workbook',return_value=workbook(None)),patch('telegram_bot.sheet_workflow.file_version',return_value=(10,20)):
            self.callback('start_queue')
        self.orchestrator.start_queue.assert_not_called()

    def test_results_stay_unavailable_until_session_and_round_are_complete(self):
        self.state['running']=True
        self.callback('results');self.assertFalse(self.bot._confirmations)
        self.state['running']=False
        self.queue.get_jobs.return_value=[{'id':'next','status':'pending'}]
        self.callback('results');self.assertFalse(self.bot._confirmations)
        self.queue.get_jobs.return_value=[]
        self.state['kpis']['total']=3
        self.callback('results');self.assertTrue(self.bot._confirmations)

    def test_multiple_uploaded_files_remain_accessible_for_review(self):
        first=self.open()
        self.bot._open_sheet_review(Path.cwd()/'second.xlsx',123)
        panel=self.bot._drafts_panel(123)
        self.assertEqual(len(self.bot._sheet_drafts),2)
        self.assertTrue(any(b['callback_data']==f'zc:map:{first}:review' for row in panel.rows for b in row))

    def test_simple_mapping_only_shows_number_amount_and_preserves_advanced_options(self):
        token=self.open();draft=self.bot._sheet_drafts[token]
        panel=review_panel(draft,token)
        self.assertIn('رقم الحساب أو الخدمة',panel.rows[0][0]['text'])
        self.callback(f'map:{token}:set:lookup_col:B')
        self.callback(f'map:{token}:set:amount_col:C')
        self.assertTrue(draft['validation']['ok'])
        self.callback(f'map:{token}:advanced')
        self.assertTrue(draft['advanced'])
        self.assertTrue(any('اسم العميل' in b['text'] for row in review_panel(draft,token).rows for b in row))

    def test_optional_column_clear_is_persisted_and_sheet_switch_works(self):
        wb=workbook();wb.create_sheet('Second').append(['Service','Amount'])
        token=self.open(wb)
        self.callback(f'map:{token}:set:customer_col:A')
        self.callback(f'map:{token}:set:customer_col:-')
        self.assertIsNone(self.bot._sheet_drafts[token]['mapping']['customer_col'])
        self.callback(f'map:{token}:sheet:1')
        self.assertEqual(self.bot._sheet_drafts[token]['sheet_name'],'Second')

    def test_expired_review_and_changed_workbook_cannot_queue(self):
        token=self.open()
        self.bot._sheet_drafts[token]['expires']=time.monotonic()-1
        self.callback(f'map:{token}:save')
        self.queue.add_job.assert_not_called()
        self.bot._sheet_drafts[token]['expires']=time.monotonic()+20
        with patch('telegram_bot.sheet_workflow.file_version',return_value=(11,21)):
            self.callback(f'map:{token}:save')
        self.queue.add_job.assert_not_called()

    def test_queue_removal_rechecks_active_status_and_is_single_use(self):
        job=SimpleNamespace(id='job_test',filename='test.xlsx',sheet_index=0,status='pending')
        self.queue.get_job.return_value=job
        self.callback('remove:job_test');token=next(iter(self.bot._confirmations))
        job.status='active';self.callback('confirm:'+token)
        self.queue.remove_job.assert_not_called()
        job.status='pending'
        self.callback('remove:job_test');token=next(iter(self.bot._confirmations))
        self.callback('confirm:'+token);self.callback('confirm:'+token)
        self.queue.remove_job.assert_called_once_with('job_test')

    def test_repair_confirm_is_session_bound_and_retries_current_errors_only(self):
        self.orchestrator.repair_errors.return_value={'status':'ok','message':'Started'}
        self.callback('repair');token=next(iter(self.bot._confirmations))
        self.orchestrator.repair_errors.assert_not_called()
        self.callback('confirm:'+token);self.callback('confirm:'+token)
        self.orchestrator.repair_errors.assert_called_once_with()

    def test_repair_does_not_interrupt_live_or_challenged_workers(self):
        for running,paused,leases,verification in ((True,False,0,[]),(True,True,1,[]),(False,False,0,[{}])):
            self.state.update(running=running,paused=paused,verifications=verification)
            self.state['kpis']['active_leases']=leases
            self.callback('repair')
            self.assertFalse(self.bot._confirmations)
        self.orchestrator.repair_errors.assert_not_called()

    def test_pending_job_edit_refused_when_it_starts_before_save(self):
        job=SimpleNamespace(id='job_test',filename='sample.xlsx',sheet_index=0,status='pending',completed=0,
                            mode='service_only',amount_target='remaining',column_mapping={'service_col':'B','amount_col':'C'})
        token=self.open();draft=self.bot._sheet_drafts[token]
        draft['job_id']=job.id;draft['mapping'].update(job.column_mapping)
        self.queue.update_pending_mapping.return_value=False
        self.callback(f'map:{token}:save')
        self.queue.add_job.assert_not_called()
        self.assertIn('تعذر',self.bot._api_call.call_args.args[1]['text'])

    def test_web_validation_gate_rejects_bad_rows_without_starting_session(self):
        handler=WebApiHandler.__new__(WebApiHandler);handler.send_json=Mock()
        with patch('openpyxl.load_workbook',return_value=workbook('invalid')):
            result=handler._validate_input_sheet(Path('unused'),0,{'service_col':'B','amount_col':'C'},'service_only','remaining')
        self.assertIsNone(result)
        self.assertEqual(handler.send_json.call_args.args[0],400)

    def test_web_queue_start_validates_every_file_before_starting(self):
        handler=WebApiHandler.__new__(WebApiHandler);handler.send_json=Mock()
        handler.orchestrator=self.orchestrator;handler.queue_service=self.queue;handler.project_root=Path.cwd()
        jobs=[{'id':str(i),'filename':f'{i}.xlsx','sheet_index':0,'status':'pending','column_mapping':{}} for i in range(2)]
        self.queue.get_jobs.return_value=jobs
        with patch.object(handler,'_validate_input_sheet',side_effect=[({},'Sheet',1),None]) as check:
            handler.handle_queue_start()
        self.assertEqual(check.call_count,2)
        self.orchestrator.start_queue.assert_not_called()
        with patch.object(handler,'_validate_input_sheet',return_value=({},'Sheet',1)):
            handler.handle_queue_start()
        self.orchestrator.start_queue.assert_called_once()

    def test_telegram_notifications_hold_all_result_files_until_round_finishes(self):
        notifier=TelegramNotifier(bot_token='TEST',chat_id=123)
        notifier.send_message=Mock();notifier.send_excel_document=Mock()
        notifier._on_session_completed('session_completed',{'workbook':'one.xlsx','round_pending':1})
        notifier.send_excel_document.assert_not_called()
        notifier._on_session_completed('session_completed',{'workbook':'two.xlsx','round_pending':0,
            'round_result_files':['one_result.xlsx','two_result.xlsx']})
        self.assertEqual(notifier.send_excel_document.call_count,2)
        self.assertEqual([call.args[0].name for call in notifier.send_excel_document.call_args_list],
                         ['one_result.xlsx','two_result.xlsx'])


class QueueAtomicTests(unittest.TestCase):
    def test_failed_start_does_not_leave_a_reserved_job_stuck_active(self):
        orchestrator=Orchestrator.__new__(Orchestrator)
        orchestrator.lock=threading.RLock();orchestrator.is_running=False;orchestrator.project_root=Path.cwd()
        orchestrator.queue_service=Mock()
        orchestrator.queue_service.get_jobs.return_value=[{'id':'one','status':'pending'}]
        orchestrator.queue_service.claim_next_pending_job.return_value=SimpleNamespace(id='one',filename='unused.xlsx',
            sheet_index=0,column_mapping={},mode='smart_hybrid',amount_target='remaining',completed=0)
        with patch.object(orchestrator,'start_session_from_config',return_value=False):
            self.assertFalse(orchestrator.start_queue())
        orchestrator.queue_service.mark_job_failed.assert_called_once()

    def service(self):
        queue=QueueService.__new__(QueueService)
        queue._lock=threading.RLock();queue._save_to_disk=Mock()
        queue._jobs=[SimpleNamespace(id='one',status='pending',completed=0,total_records=1,
                                    to_dict=lambda:{'id':'one'},remaining=1)]
        return queue

    def test_reserved_active_job_cannot_be_removed_or_remapped(self):
        queue=self.service()
        with patch('manager.queue.EVENT_BUS.publish'):
            job=queue.claim_next_pending_job()
            self.assertEqual(job.status,'active')
            self.assertFalse(queue.remove_job('one'))
            self.assertFalse(queue.update_pending_mapping('one',column_mapping={},mode='smart_hybrid',amount_target='remaining',total_records=2))
        self.assertEqual(len(queue._jobs),1)

    def test_pending_mapping_updates_without_resetting_or_removing_files(self):
        queue=self.service()
        with patch('manager.queue.EVENT_BUS.publish'),patch.object(Path,'unlink',side_effect=AssertionError('Deletion forbidden')):
            self.assertTrue(queue.update_pending_mapping('one',column_mapping={'service_col':'B'},mode='service_only',amount_target='remaining',total_records=4))
            self.assertEqual(queue._jobs[0].remaining,4)
            self.assertTrue(queue.remove_job('one'))


if __name__=='__main__':
    unittest.main()
