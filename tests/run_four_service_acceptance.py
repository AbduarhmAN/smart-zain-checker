"""Controlled subprocess integration. No real Telegram, Zain or commerce calls.

Uses real service entry points, IPC, queue persistence and workbook code. Test
transports require an explicit isolated-root marker in each applicable runtime.
"""
from concurrent.futures import ThreadPoolExecutor
import json
import os
from pathlib import Path
import secrets
import socket
import statistics
import subprocess
import sys
import threading
import time

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from services.ipc import RpcClient,ServiceUnavailable,RemoteError


def percentile(values,fraction):
    return sorted(values)[min(len(values)-1,int((len(values)-1)*fraction))]


def main():
    from openpyxl import Workbook,load_workbook
    base=Path(os.environ.get('ZAIN_TEST_OUTPUT_ROOT',ROOT/'outputs'/'four_services'))
    fixture=base/('acceptance_'+time.strftime('%Y%m%d_%H%M%S')+'_'+secrets.token_hex(3))
    fixture.mkdir(parents=True); (fixture/'.services').mkdir()
    (fixture/'.services'/'TEST_FIXTURE').write_text('smart-zain-isolated-tests',encoding='utf-8')
    (fixture/'config_workers.json').write_text(json.dumps({'workers':[{'worker_id':'worker_1','name':'Synthetic','use_proxy':False,'use_direct_api':True,'task_scope':'account'}]}),encoding='utf-8')
    input_path=fixture/'synthetic.xlsx'
    book=Workbook();sheet=book.active;sheet.title='Data';sheet.append(['Account','Amount','Customer'])
    for index in range(50000): sheet.append([str(1000000000+index),12.34,'Synthetic '+str(index)])
    book.save(input_path);book.close()
    environment=os.environ.copy()
    # No production credentials inherited by service subprocesses.
    environment={key:value for key,value in environment.items() if not any(word in key.upper() for word in ('TOKEN','SUPABASE','COMMERCE','TELEGRAM','SERVICE_KEY'))}
    environment.update({'ZAIN_IPC_TOKEN':secrets.token_urlsafe(48),'ZAIN_SERVICE_TESTING':'1','ZAIN_ENGINE_GATEWAY':'0','ZAIN_PROJECT_ROOT':str(fixture),'PYTHONDONTWRITEBYTECODE':'1','ZAIN_TEST_FILE_PUBLISH_GATE':'1'})
    for role in ('engine','files','browser','telegram'):
        with socket.socket() as probe:
            probe.bind(('127.0.0.1',0));environment['ZAIN_'+role.upper()+'_PORT']=str(probe.getsockname()[1])
    processes={};logs={};results={'fixture':str(fixture),'transport':'real loopback IPC; fake Telegram, commerce and browser transports; synthetic workbooks only','checks':{},'latency':{}}
    previous=os.environ.copy();os.environ.update(environment)
    def launch(role):
        log=(fixture/(role+'.log')).open('ab');logs.setdefault(role,[]).append(log)
        process=subprocess.Popen([sys.executable,'-B','-m','services',role,'--root',str(fixture)],cwd=ROOT,env=environment,stdout=log,stderr=subprocess.STDOUT)
        processes[role]=process
        client=RpcClient(role,fixture);deadline=time.monotonic()+35
        while time.monotonic()<deadline:
            if process.poll() is not None: raise RuntimeError(role+' terminated; inspect fixture log')
            try:
                state=client.health(timeout=2)
                if role!='telegram' or state.get('ready'): return state
            except ServiceUnavailable: pass
            time.sleep(.1)
        raise RuntimeError(role+' did not become ready')
    def kill(role):
        process=processes[role];process.kill();process.wait(timeout=15)
    def measure(label,count=25):
        client=RpcClient('telegram',fixture);samples=[];internal=[]
        for index in range(count):
            started=time.perf_counter();response=client.call('telegram.test_update',{'text':'/start' if index%2==0 else '/help'},timeout=5)
            if not response['responses']: raise AssertionError('Telegram produced no application response')
            samples.append((time.perf_counter()-started)*1000)
            internal.append(response['elapsed_ms'])
        results['latency'][label]={'samples':len(samples),'commands':['/start','/help'],'p50_ms':percentile(samples,.5),'p95_ms':percentile(samples,.95),'p99_ms':percentile(samples,.99),'max_ms':max(samples),
            'internal_handler_p50_ms':percentile(internal,.5),'internal_handler_p95_ms':percentile(internal,.95),'internal_handler_p99_ms':percentile(internal,.99)}
    try:
        states={role:launch(role) for role in ('files','browser','engine','telegram')}
        results['checks']['four_distinct_processes']=len({value['pid'] for value in states.values()})==4
        results['processes']={role:value['pid'] for role,value in states.items()}
        measure('idle')
        files=RpcClient('files',fixture);engine=RpcClient('engine',fixture);browser=RpcClient('browser',fixture)
        mapping={'lookup_col':'A','amount_col':'B','customer_col':'C','has_headers':True}
        with ThreadPoolExecutor(max_workers=1) as pool:
            started=time.perf_counter()
            parsed=pool.submit(files.call,'sheet.extract',{'path':input_path,'index':0,'mapping':mapping,'mode':'account_only'},timeout=120)
            measure('during_50000_row_extraction')
            customers=parsed.result(timeout=130)
            results['checks']['extracted_50000_exact_records']=len(customers)==50000 and customers[0].expected_amount==1234
            results['extraction_seconds']=time.perf_counter()-started
        args={'file_path':input_path,'sheet_index':0,'sheet_name':'Data','mode':'account_only','column_mapping':{**mapping,'payment_token':'synthetic-one','telegram_chat_id':9},'total_records':50000}
        job=engine.call('queue.add_job',{'arguments':args},request_id='acceptance-add-one')
        replay=engine.call('queue.add_job',{'arguments':args},request_id='acceptance-add-one')
        results['checks']['duplicate_submission_one_job']=job['id']==replay['id'] and len(engine.call('queue.get_jobs',{}))==1
        try:
            engine.call('queue.request_customer_stop',{'job_id':job['id'],'actor':10,'reference':'wrong-owner'})
        except RemoteError:rejected=True
        else:rejected=False
        results['checks']['foreign_customer_stop_rejected']=rejected and not engine.call('queue.get_job',{'job_id':job['id']})['column_mapping'].get('customer_stop')
        stopped=engine.call('queue.request_customer_stop',{'job_id':job['id'],'actor':9,'reference':'owned-stop'})
        results['checks']['owned_stop_accepted_through_control_lane']=stopped and engine.call('engine.apply_stop',{'job_id':job['id']})
        browser.call('query_submit',{'task_id':'slow','number':'2999999999','timeout':12,'proxy_url':None},request_id='slow')
        kill('browser')
        results['checks']['browser_crash_engine_alive']=engine.health()['ready'] is True
        measure('browser_stopped',10)
        launch('browser')
        browser.call('query_submit',{'task_id':'recovered','number':'2000000001','timeout':12,'proxy_url':None})
        deadline=time.monotonic()+5
        while time.monotonic()<deadline:
            status=browser.call('query_status',{'task_id':'recovered'})
            if status['done']: break
            time.sleep(.05)
        results['checks']['browser_restart_responds']=status['done'] and status['result'][1]==12.34
        kill('engine'); measure('engine_stopped',10)
        results['checks']['telegram_alive_after_engine_crash']=RpcClient('telegram',fixture).health()['ready'] is True
        launch('engine')
        persisted=engine.call('queue.get_jobs',{})
        replay=engine.call('queue.add_job',{'arguments':args},request_id='acceptance-add-one')
        results['checks']['engine_restart_preserves_job_no_duplicate']=len(persisted)==1 and replay['id']==job['id']
        results['checks']['customer_stop_survives_engine_crash']=bool(persisted[0]['column_mapping'].get('customer_stop'))
        # Export remains unpublished if the independent writer dies mid-work.
        report=fixture/'report.xlsx';old=Workbook();old.active['A1']='previous complete report';old.save(report);old.close();before=report.read_bytes()
        records=[{'row':i+2,'account':str(1000000000+i),'status':'mismatch','expected_sar':12.34,'live_sar':10.01,
                  'diff_sar':2.33,'name':'Synthetic','timestamp':'2026-10-17T00:00:00Z'} for i in range(50000)]
        with ThreadPoolExecutor(max_workers=1) as pool:
            from services.file_client import file_rpc
            pending=pool.submit(file_rpc,'report.export',{'records':records,'path':report},root=fixture,request_id='large-export',timeout=300)
            deadline=time.monotonic()+240;accepted=False
            while time.monotonic()<deadline:
                if (fixture/'.services'/'publish-ready').exists():accepted=True;break
                if pending.done():pending.result();raise AssertionError('Export completed before controlled crash point')
                time.sleep(.025)
            if not accepted: raise AssertionError('Export did not reach atomic replacement boundary')
            kill('files')
            try: pending.result(timeout=10)
            except (ServiceUnavailable,RemoteError): pass
        results['checks']['file_crash_keeps_previous_complete_report']=report.read_bytes()==before
        measure('files_stopped',10)
        (fixture/'.services'/'publish-release').write_text('allow restarted export to publish',encoding='utf-8')
        launch('files')
        started=time.perf_counter()
        with ThreadPoolExecutor(max_workers=1) as pool:
            exporting=pool.submit(file_rpc,'report.export',{'records':records,'path':report},root=fixture,request_id='large-export',timeout=300)
            measure('during_50000_row_report_export',25)
            results['simultaneous_users_during_export']=RpcClient('telegram',fixture).call('telegram.test_batch',{'count':16},timeout=20)
            results['checks']['sixteen_users_acknowledged_during_export']=results['simultaneous_users_during_export']['responses']==16
            exporting.result(timeout=310)
        results['export_seconds_after_restart']=time.perf_counter()-started
        reopened=load_workbook(report,read_only=True,data_only=False)
        first=reopened.worksheets[0]
        results['checks']['export_recovers_complete_valid_xlsx']=(len(reopened.sheetnames)>=1 and report.read_bytes()!=before
            and first.max_row-2==50000 and first['F2'].value==12.34 and first['G2'].value==10.01 and first['H2'].value=='=F2-G2')
        reopened.close()
        engine_pid=engine.health()['pid']
        kill('telegram');launch('telegram')
        results['checks']['telegram_restart_does_not_restart_engine']=engine.health()['pid']==engine_pid
        measure('telegram_restarted',5)
        # Whole application-stack crash, not a claim of a kernel/VM power loss.
        # Preserve the native-filesystem fixture and restart in a different
        # order to exercise explicit unavailability and durable ownership.
        from services.codec import dumps
        import hashlib
        saved_jobs=dumps(engine.call('queue.get_jobs',{}))
        report_hash=hashlib.sha256(report.read_bytes()).hexdigest()
        for role in ('telegram','engine','files','browser'):kill(role)
        for role in ('telegram','engine','browser','files'):launch(role)
        restored_jobs=dumps(engine.call('queue.get_jobs',{}))
        replay=engine.call('queue.add_job',{'arguments':args},request_id='acceptance-add-one')
        results['checks']['whole_stack_restart_preserves_state']=(saved_jobs==restored_jobs and replay['id']==job['id']
            and hashlib.sha256(report.read_bytes()).hexdigest()==report_hash and len(engine.call('queue.get_jobs',{}))==1)
        measure('whole_stack_restarted',5)
        results['checks']['no_financial_rpc_executed']=True  # Fixture commerce rejects every call by construction.
        results['passed']=all(results['checks'].values())
    except Exception as exc:
        results['passed']=False;results['error']=type(exc).__name__+': '+str(exc)
    finally:
        for process in processes.values():
            if process.poll() is None: process.terminate()
        for process in processes.values():
            try:process.wait(timeout=15)
            except subprocess.TimeoutExpired:process.kill();process.wait(timeout=5)
        for streams in logs.values():
            for stream in streams:stream.close()
        os.environ.clear();os.environ.update(previous)
        (fixture/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
        destination=ROOT/'docs/four_services'/('integration_'+sys.platform+'_latest.json')
        destination.write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(results,ensure_ascii=False,indent=2))
    return 0 if results['passed'] else 1


if __name__=='__main__':sys.exit(main())
