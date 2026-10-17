"""Synthetic workbooks prove split parsing, exact records and safe publication."""
import os
from pathlib import Path
import queue
import socket
import subprocess
import sys
import tempfile
import time
import threading
import unittest
from unittest.mock import patch
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from services.codec import dumps, loads
from services.file_client import RemoteWorkbook, atomic_workbook_save
from services.files import FileApplication


class FileServiceTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(prefix='zain-files-test-')
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name).resolve()
        (self.root/'.services').mkdir()
        (self.root/'.services'/'TEST_FIXTURE').write_text('smart-zain-isolated-tests',encoding='utf-8')
        self.app = FileApplication(self.root)
        self.addCleanup(self.app.stop)
        self.env = patch.dict(os.environ, {'ZAIN_SERVICE_MODE':'1','ZAIN_SERVICE_ROLE':'files',
            'ZAIN_PROJECT_ROOT':str(self.root),'ZAIN_SERVICE_TESTING':'1'})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.source = self.root/'sample.xlsx'
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Data'
        sheet.append(['Account','Amount','Customer','Notes','Existing formula'])
        sheet.append(['1000000001',12.34,'Synthetic customer','retain note','=B2*2'])
        sheet['E2'].font = Font(bold=True, color='FF00FF')
        workbook.create_sheet('Other').append(['retained'])
        workbook.save(self.source)
        workbook.close()

    def mapping(self):
        return {'lookup_col':'A','amount_col':'B','customer_col':'C','notes_col':'D','has_headers':True}

    def dispatch(self, operation, payload, **kwargs):
        with patch.dict(os.environ, {'ZAIN_SERVICE_ROLE':'files'}):
            # Cross the actual codec, preserving Customer integer halalas and
            # every metadata field instead of the lossy public to_dict().
            return loads(dumps(self.app.handlers()[operation](loads(dumps(payload)))))

    def test_remote_parse_validate_and_quote_preserve_native_rows_and_metadata(self):
        from domain.workbook import inspect_sheet_schema, extract_session_customers
        from domain.workbook_validation import validate_sheet
        with patch.dict(os.environ, {'ZAIN_SERVICE_ROLE':'telegram'}), patch('services.file_client.file_rpc', self.dispatch):
            workbook = RemoteWorkbook(self.source)
            self.assertEqual(workbook.sheetnames, ['Data','Other'])
            sheet = workbook.worksheets[0]
            self.assertTrue(inspect_sheet_schema(sheet)['columns'])
            self.assertTrue(validate_sheet(sheet,self.mapping(),'account_only')['ok'])
            customers = extract_session_customers(sheet,self.mapping(),'account_only')
            self.assertEqual(len(customers),1)
            self.assertEqual(customers[0].expected_amount,1234)
            self.assertEqual(customers[0].notes,'retain note')
            self.assertEqual(customers[0].row_number,2)
            self.assertEqual(sheet.rpc('quote_rows',mapping=self.mapping(),mode='account_only'),
                [{'row_number':2,'check_type':'account'}])

    def test_workspace_escape_and_frozen_write_are_rejected(self):
        for path in (self.root.parent/'outside.xlsx', self.root/'..'/'outside.xlsx'):
            with self.assertRaises(ValueError):self.app.names({'path':path})
        with self.assertRaises(ValueError):
            self.app.update({'path':self.root/'.commerce'/'inputs'/'abc.xlsx','records':[]})
        with self.assertRaises(ValueError):
            self.app.records({'path':self.source,'options':{'__import__':'os'}})

    def test_numeric_original_rows_never_gain_decimal_zero(self):
        from domain.workbook import extract_customer_records
        workbook=Workbook()
        sheet=workbook.active
        sheet.append(['Row','Account','Amount'])
        sheet.append([54.0,'1000000001',12.34])
        sheet.append(['79.0','1000000002',12.34])
        records=extract_customer_records(sheet,lookup_col=2,amount_col=3,source_row_col=1)
        self.assertEqual([record.row_number for record in records],[54,79])
        sheet['A2']=54.5
        with self.assertRaises(ValueError):extract_customer_records(sheet,lookup_col=2,amount_col=3,source_row_col=1)
        workbook.close()

    def test_atomic_replace_failure_retains_original_and_cleans_pending(self):
        before = self.source.read_bytes()
        replacement = Workbook()
        replacement.active['A1'] = 'unpublished'
        with patch('services.file_client.os.replace',side_effect=PermissionError('synthetic lock')):
            with self.assertRaises(PermissionError):atomic_workbook_save(replacement,self.source)
        replacement.close()
        self.assertEqual(self.source.read_bytes(),before)
        self.assertFalse(list(self.root.glob('*.pending.xlsx')))

    def test_report_export_is_remote_and_preserves_formula_styles(self):
        from zain_checker.executive_reporter import export_executive_workbook
        record = {'row':2,'name':'Synthetic','account':'1000000001','expected_sar':12.34,
                  'live_sar':10,'diff_sar':2.34,'status':'mismatch','main_status':'open'}
        destination=self.root/'report.xlsx'
        with patch.dict(os.environ, {'ZAIN_SERVICE_ROLE':'engine'}), patch('services.file_client.file_rpc',self.dispatch):
            result=export_executive_workbook([record],destination)
        self.assertEqual(result,destination)
        workbook=load_workbook(destination,data_only=False)
        self.addCleanup(workbook.close)
        self.assertTrue(any(cell.data_type=='f' for sheet in workbook for row in sheet for cell in row))
        self.assertTrue(workbook.worksheets[0]['A1'].font.bold)
        self.assertTrue(workbook.worksheets[0].sheet_view.rightToLeft)
        self.assertEqual(self.app.names({'path':self.source}),['Data','Other'])

    def test_source_update_preserves_unrelated_formulas_and_formatting(self):
        record={'row':2,'name':'Synthetic','account':'1000000001','expected_sar':12.34,
                'live_sar':10,'diff_sar':2.34,'status':'mismatch'}
        result=self.app.update({'path':self.source,'records':[record]})
        self.assertTrue(result[0])
        workbook=load_workbook(self.source,data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook['Data']['E2'].value,'=B2*2')
        self.assertTrue(workbook['Data']['E2'].font.bold)
        self.assertEqual(workbook['Other']['A1'].value,'retained')

    def test_preview_append_is_idempotent_across_file_service_restart(self):
        keys=['name','national_id','contract','account','service','customer_phones','collector','supervisor',
              'branch','expected_sar','live_sar','diff_sar','status_label','main_status','sub_status',
              'follow_notes','follow_date','timestamp']
        record={key:'' for key in keys}
        record.update(row=2,name='Synthetic',expected_sar=12.34,live_sar=10,diff_sar=2.34)
        payload={'path':self.root/'stream.xlsx','records':[record]}
        self.assertEqual(self.app.stream(payload)['rows'],1)
        self.assertEqual(self.app.stream(payload)['rows'],1)
        self.app.stop()
        self.assertEqual(self.app.stream(payload)['rows'],1)

    def test_stream_thread_calls_remote_writer_without_opening_workbook_locally(self):
        from zain_checker.async_pipeline_engine import AsyncExcelWriterThread
        writer=AsyncExcelWriterThread(queue.Queue(),self.root/'preview.xlsx',self.root/'engine-wal.jsonl',wal_already_persisted=True)
        writer.result_queue.put({'row':2})
        writer.result_queue.put('STOP_SENTINEL')
        calls=[]
        def capture(operation,payload,**kwargs):
            calls.append((operation,payload,kwargs))
            return {'rows':1}
        with patch.dict(os.environ, {'ZAIN_SERVICE_ROLE':'engine'}), patch('services.file_client.file_rpc',capture), \
                patch.object(writer,'_init_workbook_and_wal',side_effect=AssertionError('local Excel forbidden')):
            writer.run()
        self.assertEqual(calls[0][0],'stream.append')
        self.assertEqual(calls[0][1]['records'],[{'row':2}])
        self.assertTrue(calls[0][2]['request_id'])

    def test_fifty_thousand_records_extract_without_losing_row_identity(self):
        path=self.root/'large.xlsx'
        workbook=Workbook(write_only=True)
        sheet=workbook.create_sheet('Data')
        sheet.append(['Account','Amount'])
        for index in range(50000):sheet.append([str(1000000000+index),1.23])
        workbook.save(path)
        workbook.close()
        mapping={'lookup_col':'A','amount_col':'B','has_headers':True}
        result=self.app.extract({'path':path,'mapping':mapping,'mode':'account_only'})
        self.assertEqual(len(result),50000)
        self.assertEqual((result[0].row_number,result[-1].row_number),(2,50001))
        self.assertTrue(all(customer.expected_amount==123 for customer in result))

    def test_actual_file_process_serves_workbooks_across_authenticated_ipc(self):
        from services.ipc import RpcClient
        with socket.socket() as probe:
            probe.bind(('127.0.0.1',0))
            port=probe.getsockname()[1]
        environment=dict(os.environ,ZAIN_FILES_PORT=str(port),ZAIN_IPC_TOKEN='synthetic-file-test-token-'+'x'*32,
            ZAIN_SERVICE_ROLE='files')
        program=("import os,time; from pathlib import Path; from services.files import build; "
            "from services.ipc import RpcServer; app=build(Path(os.environ['ZAIN_PROJECT_ROOT'])); "
            "server=RpcServer('files',app.root,app.handlers(),health=app.health,workers=1); "
            "server.start(); time.sleep(60)")
        process=subprocess.Popen([sys.executable,'-B','-c',program],cwd=Path(__file__).resolve().parents[1],
            env=environment,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        try:
            with patch.dict(os.environ,{**environment,'ZAIN_SERVICE_ROLE':'engine'}):
                client=RpcClient('files',self.root)
                deadline=time.monotonic()+10
                while time.monotonic()<deadline:
                    try:
                        status=client.health(timeout=.5)
                        break
                    except Exception:
                        if process.poll() is not None:self.fail('Synthetic File service exited before readiness')
                        time.sleep(.05)
                else:self.fail('Synthetic File service did not become ready')
                self.assertEqual(status['pid'],process.pid)
                self.assertNotEqual(status['pid'],os.getpid())
                workbook=RemoteWorkbook(self.source)
                from domain.workbook import extract_session_customers
                records=extract_session_customers(workbook.worksheets[0],self.mapping(),'account_only')
                self.assertEqual(records[0].expected_amount,1234)
                self.assertEqual(records[0].notes,'retain note')
        finally:
            process.terminate()
            process.wait(timeout=10)

    def test_observation_timeout_reuses_request_identity_while_service_is_live(self):
        from services.file_client import file_rpc
        from services.ipc import ServiceUnavailable
        client=Mock()
        client.call.side_effect=[ServiceUnavailable('running'),{'saved':True}]
        with patch('services.ipc.RpcClient',return_value=client):
            self.assertEqual(file_rpc('report.export',{'path':self.root/'report.xlsx','records':[]}),{'saved':True})
        first,second=client.call.call_args_list
        self.assertEqual(first.kwargs['request_id'],second.kwargs['request_id'])
        client.health.assert_called_once()

    def test_unavailable_service_fails_without_starting_another_request(self):
        from services.file_client import file_rpc
        from services.ipc import ServiceUnavailable
        client=Mock()
        client.call.side_effect=ServiceUnavailable('offline')
        client.health.side_effect=ServiceUnavailable('offline')
        with patch('services.ipc.RpcClient',return_value=client):
            with self.assertRaises(ServiceUnavailable):file_rpc('report.export',{'path':self.root/'report.xlsx','records':[]})
        self.assertEqual(client.call.call_count,1)

    def test_memory_admission_defers_before_workbook_open_or_owner_lock(self):
        from services.ipc import RetryLater
        with patch('services.resources.admit',return_value=False), \
                patch('domain.workbook.load_fast_workbook',side_effect=AssertionError('must defer before parse')):
            with self.assertRaises(RetryLater):self.app.validate({'path':self.source,'mapping':self.mapping()})
        self.assertTrue(self.app.lock.acquire(blocking=False))
        self.app.lock.release()

    def test_concurrent_file_requests_execute_one_workbook_operation_at_a_time(self):
        counters={'active':0,'maximum':0,'finished':0}
        guard=threading.Lock()
        def consume(sheet):
            with guard:
                counters['active']+=1
                counters['maximum']=max(counters['maximum'],counters['active'])
            time.sleep(.03)
            with guard:
                counters['active']-=1
                counters['finished']+=1
        threads=[threading.Thread(target=self.app.sheet,args=({'path':self.source},consume)) for _ in range(6)]
        for thread in threads:thread.start()
        for thread in threads:thread.join(timeout=5)
        self.assertEqual(counters,{'active':0,'maximum':1,'finished':6})


if __name__=='__main__':unittest.main()
