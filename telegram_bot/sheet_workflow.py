"""Telegram sheet review: all columns, explicit mapping, and offline row validation."""
from __future__ import annotations

import secrets
import time
from pathlib import Path
from openpyxl import load_workbook
from domain.workbook import inspect_sheet_schema
from . import views

from domain.workbook_validation import FIELDS, MODES, validate_sheet


def file_version(path):
    stat = Path(path).stat()
    return stat.st_size, stat.st_mtime_ns


def review_panel(draft, token):
    result = draft["validation"]
    text = ("<b>📋 مراجعة الشيت</b>\n\n" + views.safe(draft["path"].name) +
            "\nالورقة: <b>" + views.safe(draft["sheet_name"]) + "</b>\n" +
            ("✓ مناسب للفحص" if result["ok"] else "⚠ يحتاج تعديل قبل إضافته") +
            "\nصفوف البيانات: " + views.count(result["rows"]))
    if result["problems"]:
        text += "\n\n" + "\n".join(views.safe(p, 180) for p in result["problems"])
    text += "\n\nاضغط على نوع البيانات لاختيار عموده. جميع الأعمدة متاحة، حتى بدون عنوان."
    if draft.get("advanced"):
        rows = [[views.button(f"{label}: {draft['mapping'].get(key) or 'غير محدد'}", f"map:{token}:field:{key}:0")]
                for key, label in FIELDS.items()]
    else:
        number_col = draft['mapping'].get('lookup_col') or draft['mapping'].get('service_col')
        rows = [[views.button(f"رقم الحساب أو الخدمة: {number_col or 'اختر العمود'}", f"map:{token}:field:lookup_col:0")],
                [views.button(f"المبلغ: {draft['mapping'].get('amount_col') or 'اختر العمود'}", f"map:{token}:field:amount_col:0")]]
    if draft.get("advanced"):
        rows += [[views.button(("✓ " if draft["mode"] == mode else "") + label, f"map:{token}:mode:{mode}")
              for mode, label in MODES.items()],
             [views.button(("✓ " if draft["amount_target"] == target else "") + label, f"map:{token}:target:{target}")
              for target, label in (("remaining", "مبلغ الشيت"), ("contract", "المبلغ الثاني"), ("smart_dual", "كلاهما"))],
              ]
    rows += [[views.button("كل الأعمدة", f"map:{token}:all:0"), *([] if draft.get("job_id") else [views.button("تغيير الورقة", f"map:{token}:sheets:0")])],
             [views.button("إعدادات أبسط" if draft.get("advanced") else "أعمدة إضافية / خيارات متقدمة", f"map:{token}:advanced")],
             [views.button("↻ التحقق من البيانات", f"map:{token}:review")],
             [views.button("✓ حفظ الأعمدة" if draft.get("job_id") else "✓ إضافة إلى الطابور", f"map:{token}:save", "success")],
             views.nav()]
    return views.Panel(text, rows)


def columns_panel(draft, token, field=None, page=0):
    columns = draft["columns"]
    page = max(0, min(page, max(0, (len(columns)-1)//8)))
    label = "رقم الحساب أو الخدمة" if field == "lookup_col" and not draft.get("advanced") else (FIELDS[field] if field else "كل الأعمدة")
    text = f"<b>📑 {views.safe(label)}</b>\n{len(columns)} عمود · صفحة {page+1} من {max(1,(len(columns)+7)//8)}\n"
    rows = []
    for col in columns[page*8:(page+1)*8]:
        letter = col["col_letter"]
        title = f"{letter} · {col['header'] or 'بدون عنوان'}"
        text += "\n" + views.safe(title, 100)
        if field:
            rows.append([views.button(title[:65], f"map:{token}:set:{field}:{letter}")])
    paging = []
    prefix = f"map:{token}:field:{field}" if field else f"map:{token}:all"
    if page:
        paging.append(views.button("‹ السابق", f"{prefix}:{page-1}"))
    if (page+1)*8 < len(columns):
        paging.append(views.button("التالي ›", f"{prefix}:{page+1}"))
    if paging:
        rows.append(paging)
    if field:
        rows.append([views.button("بدون عمود / إلغاء الاختيار", f"map:{token}:set:{field}:-")])
    rows.append([views.button("‹ مراجعة الشيت", f"map:{token}:review")])
    return views.Panel(text, rows)


class SheetWorkflowMixin:
    def _drafts_panel(self, chat_id, page=0):
        with self._ui_lock:
            drafts = [(token,draft) for token,draft in self._sheet_drafts.items()
                      if draft["chat"] == chat_id and draft["expires"] > time.monotonic()]
        page = max(0, min(page, max(0,(len(drafts)-1)//8)))
        text = "<b>ملفات تحتاج مراجعة</b>\nاختر ملفًا، راجع بياناته، ثم أضفه إلى الجولة.\n"
        rows = [[views.button(draft["path"].name[:65], f"map:{token}:review")]
                for token,draft in drafts[page*8:(page+1)*8]]
        if not drafts:
            text += "\nلا توجد ملفات تنتظر مراجعتك."
        if page:
            rows.append([views.button("‹ السابق", f"drafts:{page-1}")])
        if (page+1)*8 < len(drafts):
            rows.append([views.button("التالي ›", f"drafts:{page+1}")])
        rows += [[views.button("📎 إضافة ملف", "upload"), views.button("📁 الطابور", "queue:0")], views.nav()]
        return views.Panel(text, rows)

    def _load_review_sheet(self, draft, index):
        wb = load_workbook(draft["path"], read_only=True, data_only=True)
        try:
            if index < 0 or index >= len(wb.worksheets):
                raise ValueError("Unknown worksheet")
            sheet = wb.worksheets[index]
            has_headers = True
            if draft.get("job_id"):
                saved_job = self.queue_service.get_job(draft["job_id"])
                if saved_job:
                    has_headers = saved_job.column_mapping.get("has_headers", True)
            schema = inspect_sheet_schema(sheet, has_headers=has_headers)
            mapping = {key: schema["letters"].get(key) for key in FIELDS}
            mapping["has_headers"] = has_headers
            mode = "smart_hybrid"
            target = "remaining"
            if draft.get("job_id"):
                job = self.queue_service.get_job(draft["job_id"])
                if not job or job.status != "pending" or job.completed:
                    raise ValueError("Job is no longer editable")
                mapping.update({key: job.column_mapping[key] for key in FIELDS if key in job.column_mapping})
                mode, target = job.mode, job.amount_target
            draft.update(sheet_index=index, sheet_name=sheet.title, sheets=[s.title for s in wb.worksheets],
                         columns=schema["columns"], mapping=mapping, mode=mode, amount_target=target)
            # Preserve non-interactive metadata columns needed by the existing parser.
            draft["metadata"] = {k:v for k,v in schema["letters"].items() if k not in FIELDS}
            draft["validation"] = validate_sheet(sheet, mapping, mode, target)
        finally:
            wb.close()

    def _open_sheet_review(self, path, chat_id, message_id=None, job=None, show=True):
        token = secrets.token_hex(6)
        draft = {"path":Path(path), "chat":chat_id, "expires":time.monotonic()+1800,
                 "version":file_version(path), "job_id":job.id if job else None}
        self._load_review_sheet(draft, job.sheet_index if job else 0)
        with self._ui_lock:
            self._sheet_drafts = {k:v for k,v in self._sheet_drafts.items() if v["expires"] > time.monotonic()}
            if len(self._sheet_drafts) >= 32:
                self._sheet_drafts.pop(next(iter(self._sheet_drafts)))
            self._sheet_drafts[token] = draft
        panel = review_panel(draft, token)
        if show:
            self._show_panel(chat_id, panel, message_id)
        return panel

    def _sheet_callback(self, command, chat_id, message_id):
        parts = command.split(":")
        token = parts[1] if len(parts) > 1 else ""
        with self._ui_lock:
            draft = self._sheet_drafts.get(token)
        if not draft or draft["chat"] != chat_id or draft["expires"] <= time.monotonic():
            self._show_panel(chat_id, views.notice("مراجعة غير متاحة", "أرسل الملف مجددًا أو افتح أعمدة الشيت من الطابور."), message_id)
            return
        try:
            if draft["version"] != file_version(draft["path"]):
                raise ValueError("Workbook changed")
            action = parts[2]
            if action == "field" and len(parts) == 5 and parts[3] in FIELDS:
                panel = columns_panel(draft, token, parts[3], int(parts[4]))
            elif action == "all" and len(parts) == 4:
                panel = columns_panel(draft, token, page=int(parts[3]))
            elif action == "sheets" and len(parts) == 4:
                page = max(0, min(int(parts[3]), (len(draft["sheets"])-1)//8))
                rows = [[views.button(name[:60], f"map:{token}:sheet:{index}")]
                        for index,name in enumerate(draft["sheets"]) if page*8 <= index < (page+1)*8]
                if page:
                    rows.append([views.button("‹ السابق", f"map:{token}:sheets:{page-1}")])
                if (page+1)*8 < len(draft["sheets"]):
                    rows.append([views.button("التالي ›", f"map:{token}:sheets:{page+1}")])
                rows.append([views.button("‹ مراجعة الشيت", f"map:{token}:review")])
                panel = views.Panel("<b>اختر ورقة العمل</b>", rows)
            elif action == "sheet" and len(parts) == 4 and not draft.get("job_id"):
                self._load_review_sheet(draft, int(parts[3]))
                panel = review_panel(draft, token)
            elif action in {"set", "mode", "target", "save", "review", "advanced"}:
                if action == "advanced":
                    draft["advanced"] = not draft.get("advanced", False)
                    if not draft["advanced"]:
                        draft["amount_target"] = "remaining"
                if action == "set":
                    field, letter = parts[3:]
                    if field not in FIELDS or (letter != "-" and letter not in {c["col_letter"] for c in draft["columns"]}):
                        raise ValueError("Unknown column")
                    draft["mapping"][field] = None if letter == "-" else letter
                    if field == "lookup_col" and not draft.get("advanced"):
                        draft["mapping"]["service_col"] = draft["mapping"][field]
                elif action == "mode":
                    if parts[3] not in MODES:
                        raise ValueError("Unknown mode")
                    draft["mode"] = parts[3]
                elif action == "target":
                    if parts[3] not in {"remaining", "contract", "smart_dual"}:
                        raise ValueError("Unknown target")
                    draft["amount_target"] = parts[3]
                wb = load_workbook(draft["path"], read_only=True, data_only=True)
                try:
                    draft["validation"] = validate_sheet(wb.worksheets[draft["sheet_index"]], draft["mapping"], draft["mode"], draft["amount_target"])
                finally:
                    wb.close()
                if draft["version"] != file_version(draft["path"]):
                    raise ValueError("Workbook changed during review")
                if action == "save" and draft["validation"]["ok"]:
                    with self.orchestrator.lock:
                        if draft.get("job_id"):
                            saved = self.queue_service.update_pending_mapping(draft["job_id"],
                                column_mapping={**draft["metadata"], **draft["mapping"]}, mode=draft["mode"],
                                amount_target=draft["amount_target"], total_records=draft["validation"]["rows"])
                            if not saved:
                                raise ValueError("Job is no longer waiting")
                        else:
                            self.queue_service.add_job(file_path=draft["path"], sheet_index=draft["sheet_index"],
                                sheet_name=draft["sheet_name"], mode=draft["mode"], amount_target=draft["amount_target"],
                                column_mapping={**draft["metadata"], **draft["mapping"]}, total_records=draft["validation"]["rows"])
                    with self._ui_lock:
                        self._sheet_drafts.pop(token, None)
                    panel = views.Panel("<b>✓ حُفظت الأعمدة</b>\n" + views.safe(draft["sheet_name"]) + "\nالشيت في الطابور، وسيستخدم الأعمدة التي اخترتها.",
                                        [[views.button("▶ بدء الجولة", "start_queue", "success")],
                                         [views.button("📎 إضافة ملف آخر", "upload"), views.button("ملفات تحتاج مراجعة", "drafts:0")],
                                         [views.button("📁 عرض الطابور", "queue:0")], views.nav()])
                else:
                    panel = review_panel(draft, token)
            else:
                raise ValueError("Unknown action")
        except Exception:
            panel = views.notice("تعذر إكمال المراجعة", "قد يكون الملف أو حالة الشيت تغيرت. افتح مراجعة جديدة من الطابور أو أرسل الملف مجددًا.")
        self._show_panel(chat_id, panel, message_id)
