"""Web REST API Handler for Smart Zain Checker Dashboard.
Serves static assets, templates, and provides JSON endpoints.
"""
from __future__ import annotations

import json
import logging
import mimetypes
from http.server import BaseHTTPRequestHandler
from pathlib import Path
from typing import Any, Optional
from urllib.parse import urlparse

from domain.workbook import inspect_sheet_schema
from manager.orchestrator import Orchestrator
from manager.queue import QueueService

logger = logging.getLogger("WebAPI")


class WebApiHandler(BaseHTTPRequestHandler):
    orchestrator: Orchestrator
    queue_service: QueueService
    web_ui_dir: Path
    project_root: Path

    def log_message(self, format: str, *args: Any) -> None:
        # Suppress noisy standard HTTP access logs
        pass

    def send_json(self, status_code: int, data: dict[str, Any]) -> None:
        raw = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status_code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(raw)

    def send_file_response(self, file_path: Path) -> None:
        if not file_path.exists() or not file_path.is_file():
            self.send_error(404, "File Not Found")
            return
        mime_type, _ = mimetypes.guess_type(str(file_path))
        mime_type = mime_type or "application/octet-stream"
        raw = file_path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", f"{mime_type}; charset=utf-8" if "text" in mime_type or "javascript" in mime_type else mime_type)
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path

        # Static assets and template routing
        if path in ("/", "/index.html"):
            index_path = self.web_ui_dir / "index.html"
            self.send_file_response(index_path)
            return

        if path.startswith("/static/"):
            rel_path = path[len("/static/"):]
            target = self.web_ui_dir / "static" / rel_path
            self.send_file_response(target)
            return

        # API Endpoints
        if path == "/api/session-info":
            self.handle_session_info()
        elif path == "/api/workbook-sheets":
            self.handle_workbook_sheets(parsed)
        elif path == "/api/live-status":
            self.handle_live_status()
        elif path == "/api/service-verification-image":
            from urllib.parse import parse_qs
            from workers.stealth_service_engine import verification_image
            challenge_id = parse_qs(parsed.query).get("id", [""])[0]
            raw = verification_image(challenge_id)
            if raw is None:
                self.send_json(404, {"status": "error", "message": "انتهت صلاحية التحدي"})
            else:
                self.send_response(200)
                self.send_header("Content-Type", "image/png")
                self.send_header("Cache-Control", "no-store")
                self.send_header("Content-Length", str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)
        elif path == "/api/queue":
            self.handle_get_queue()
        elif path == "/api/telegram/status":
            self.handle_telegram_status()
        elif path == "/api/check-sheet-status":
            self.handle_check_sheet_status(parsed)
        elif path == "/api/detect-results-sheet":
            self.handle_detect_results_sheet(parsed)
        elif path == "/api/download-results":
            from urllib.parse import parse_qs, quote
            query_params = parse_qs(parsed.query)
            job_id = query_params.get("job_id", [None])[0]
            requested_file = query_params.get("file", [None])[0]

            res_file = None
            download_name = "نتائج فحص زين.xlsx"
            is_current = query_params.get("current", ["0"])[0] == "1"

            if is_current:
                res_file = self.orchestrator.result_path
                download_name = res_file.name
            elif job_id:
                jobs = self.queue_service.get_jobs()
                for j in jobs:
                    if j.get("id") == job_id:
                        rf_name = j.get("result_file")
                        if rf_name:
                            candidate = Path(self.project_root) / rf_name
                            if candidate.exists():
                                res_file = candidate
                                download_name = rf_name
                        break
            elif requested_file:
                candidate = Path(self.project_root) / requested_file
                if candidate.exists():
                    res_file = candidate
                    download_name = requested_file

            if not res_file or not res_file.exists():
                res_file = Path(self.project_root).resolve() / "نتائج فحص زين.xlsx"
                download_name = "نتائج فحص زين.xlsx"

            try:
                if (is_current or not job_id and not requested_file) and hasattr(self.orchestrator, "all_completed_records") and self.orchestrator.all_completed_records:
                    from zain_checker.executive_reporter import export_executive_workbook
                    export_executive_workbook(list(self.orchestrator.all_completed_records.values()), res_file)
            except Exception as e:
                logger.warning(f"On-demand export warning: {e}")

            if res_file.exists():
                raw = res_file.read_bytes()
                encoded_name = quote(download_name.encode('utf-8'))
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f"attachment; filename*=UTF-8''{encoded_name}")
                self.send_header("Content-Length", str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)
            else:
                self.send_json(404, {"status": "error", "message": "ملف النتائج غير موجود حالياً"})
            return
        else:
            self.send_json(404, {"status": "error", "message": f"Endpoint {path} not found"})

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path
        if path.startswith("/api/telegram/"):
            origin = self.headers.get("Origin")
            if origin and urlparse(origin).netloc != self.headers.get("Host"):
                self.send_json(403, {"status": "error", "message": "مصدر الطلب غير مسموح"})
                return

        content_length = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(content_length).decode("utf-8") if content_length > 0 else "{}"
        try:
            payload = json.loads(body) if body.strip() else {}
        except Exception:
            payload = {}

        if path == "/api/upload-workbook":
            self.handle_upload_workbook(payload)
        elif path == "/api/start-session":
            self.handle_start_session(payload)
        elif path == "/api/pause-session":
            self.orchestrator.pause_session()
            self.send_json(200, {"status": "ok", "message": "Session paused"})
        elif path == "/api/resume-session":
            from workers.stealth_service_engine import pending_verifications
            if pending_verifications():
                self.send_json(409, {"status": "error", "message": "أكمل التحقق الظاهر قبل استئناف الفحص"})
                return
            self.orchestrator.resume_session()
            self.send_json(200, {"status": "ok", "message": "Session resumed"})
        elif path == "/api/service-verification-submit":
            origin = self.headers.get("Origin")
            if origin and urlparse(origin).netloc != self.headers.get("Host"):
                self.send_json(403, {"status": "error", "message": "مصدر الطلب غير مسموح"})
                return
            from workers.stealth_service_engine import submit_verification
            accepted = submit_verification(str(payload.get("id", "")), payload.get("answer", ""))
            self.send_json(200 if accepted else 409, {"status": "ok" if accepted else "error",
                            "message": "تم إرسال الرمز؛ انتظر نتيجة زين" if accepted else "التحدي منتهٍ أو سبق إرسال الرمز"})
        elif path == "/api/cancel-session":
            self.orchestrator.cancel_session()
            self.send_json(200, {"status": "ok", "message": "Session cancelled"})
        elif path == "/api/inspect-sheet-columns":
            self.handle_inspect_sheet_columns(payload)
        elif path == "/api/validate-sheet":
            wb_name = str(payload.get("workbook", ""))
            try:
                checked = self._validate_input_sheet(self.project_root/wb_name, int(payload.get("sheet_index",0)),
                    payload.get("column_mapping") or {}, payload.get("mode","smart_hybrid"), payload.get("amount_target","remaining"))
            except (TypeError,ValueError):
                self.send_json(400,{"status":"error","message":"حدد ورقة عمل صحيحة."})
                return
            if checked:
                self.send_json(200,{"status":"ok","validation":{"ok":True,"rows":checked[2],"problems":[]}})
        elif path == "/api/queue/add":
            self.handle_queue_add(payload)
        elif path == "/api/queue/remove":
            self.handle_queue_remove(payload)
        elif path == "/api/queue/restart":
            self.handle_queue_restart(payload)
        elif path == "/api/queue/restart-all":
            self.handle_queue_restart_all(payload)
        elif path == "/api/queue/start":
            self.handle_queue_start()
        elif path == "/api/telegram/test-ping":
            self.handle_telegram_ping()
        elif path in ("/api/telegram/check-connection", "/api/telegram/toggle", "/api/telegram/send-results"):
            origin = self.headers.get("Origin")
            if origin and urlparse(origin).netloc != self.headers.get("Host"):
                self.send_json(403, {"status": "error", "message": "مصدر الطلب غير مسموح"})
                return
            self.handle_telegram_action(path, payload)
        elif path == "/api/repair-errors":
            row_arg = payload.get("row")
            try:
                row_val = int(row_arg) if row_arg is not None else None
            except (ValueError, TypeError):
                row_val = None
            wb_name = payload.get("workbook")
            wb_path = self.project_root / wb_name if wb_name else None
            sheet_idx = int(payload.get("sheet_index", 0))
            res = self.orchestrator.repair_errors(row=row_val, workbook_path=wb_path, sheet_index=sheet_idx)
            code = 200 if res.get("status") in ("ok", "info") else 400
            self.send_json(code, res)
        elif path == "/api/start-repair-session":
            self.handle_start_repair_session(payload)
        else:
            self.send_json(404, {"status": "error", "message": f"POST endpoint {path} not found"})

    def handle_session_info(self) -> None:
        # Discover all .xlsx files in the project directory (excluding hidden/stream/temp files)
        xlsx_files = [f.name for f in self.project_root.glob("*.xlsx")
                      if not f.name.startswith("~") and not f.name.startswith(".")]
        # Prioritize 2.xlsx and zain_data.xlsx
        xlsx_files.sort(key=lambda x: (x != "2.xlsx", x != "zain_data.xlsx", x))

        default_sheets = []
        if xlsx_files:
            try:
                from openpyxl import load_workbook
                wb = load_workbook(self.project_root / xlsx_files[0], read_only=True)
                default_sheets = wb.sheetnames
                wb.close()
            except Exception:
                pass

        self.send_json(200, {
            "status": "ok",
            "workbooks": xlsx_files,
            "default_sheets": default_sheets,
        })

    def handle_workbook_sheets(self, parsed: Any) -> None:
        from urllib.parse import parse_qs
        qs = parse_qs(parsed.query)
        wb_name = qs.get("workbook", ["2.xlsx"])[0]
        wb_path = self.project_root / Path(wb_name).name
        if not wb_path.exists():
            self.send_json(404, {"status": "error", "message": f"Workbook not found: {wb_name}"})
            return
        try:
            from openpyxl import load_workbook
            wb = load_workbook(wb_path, read_only=True)
            sheets = wb.sheetnames
            wb.close()
            self.send_json(200, {"status": "ok", "workbook": wb_name, "sheets": sheets})
        except Exception as exc:
            self.send_json(500, {"status": "error", "message": str(exc)})

    def handle_upload_workbook(self, payload: dict[str, Any]) -> None:
        filename = payload.get("filename", "").strip()
        data_base64 = payload.get("data_base64", "").strip()
        if not filename or not data_base64:
            self.send_json(400, {"status": "error", "message": "Missing filename or file data"})
            return

        import base64
        clean_name = Path(filename).name
        if not clean_name.lower().endswith((".xlsx", ".xls")):
            clean_name += ".xlsx"
        target_path = self.project_root / clean_name
        try:
            raw_bytes = base64.b64decode(data_base64)
            target_path.write_bytes(raw_bytes)
            from openpyxl import load_workbook
            wb = load_workbook(target_path, read_only=True)
            sheets = wb.sheetnames
            wb.close()
            self.send_json(200, {"status": "ok", "filename": clean_name, "sheets": sheets})
        except Exception as exc:
            self.send_json(500, {"status": "error", "message": str(exc)})

    def handle_live_status(self) -> None:
        status = self.orchestrator.get_live_status()
        self.send_json(200, status)

    def handle_get_queue(self) -> None:
        jobs = self.queue_service.get_jobs()
        self.send_json(200, {"status": "ok", "jobs": jobs})

    def handle_inspect_sheet_columns(self, payload: dict[str, Any]) -> None:
        wb_name = payload.get("workbook", "2.xlsx")
        wb_path = self.project_root / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        sheet_idx = int(payload.get("sheet_index", 0))

        if not wb_path.exists():
            self.send_json(404, {"status": "error", "message": f"File not found: {wb_name}"})
            return

        from openpyxl import load_workbook
        wb = load_workbook(wb_path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[sheet_idx] if 0 <= sheet_idx < len(wb.worksheets) else wb.worksheets[0]
            has_headers = payload.get("has_headers", True)
            if not isinstance(has_headers, bool):
                self.send_json(400, {"status":"error", "message":"حدد هل يحتوي الملف على صف عناوين."})
                return
            analysis = inspect_sheet_schema(ws, has_headers=has_headers)
            analysis["sheet_name"] = ws.title
            analysis["total_sheets"] = len(wb.worksheets)
            self.send_json(200, {"status": "ok", "analysis": analysis})
        finally:
            wb.close()

    def handle_start_session(self, payload: dict[str, Any]) -> None:
        wb_name = payload.get("workbook", "2.xlsx")
        wb_path = self.project_root / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        sheet_idx = int(payload.get("sheet_index", 0))
        mapping = payload.get("column_mapping") or payload.get("custom_config") or {}
        mode = payload.get("mode") or payload.get("scope") or "smart_hybrid"
        amount_target = payload.get("amount_target", "contract")
        target_url = payload.get("target_url", "https://business.zain.sa/dashboard/quick-pay")
        result_file = payload.get("result_file")
        force_restart = bool(payload.get("force_restart", False))

        # Check if queue has a saved job with full mapping for this workbook
        if not mapping:
            q_job = self.queue_service.get_job_by_file_and_sheet(wb_path.name, sheet_idx)
            if q_job and q_job.column_mapping:
                mapping = dict(q_job.column_mapping)
                mode = getattr(q_job, "mode", mode)
                amount_target = getattr(q_job, "amount_target", amount_target)
                if not result_file:
                    result_file = getattr(q_job, "result_file", None)

        checked = self._validate_input_sheet(wb_path, sheet_idx, mapping, mode, amount_target)
        if not checked:
            return
        mapping, _, _ = checked
        success = self.orchestrator.start_session_from_config(
            workbook_path=wb_path,
            sheet_index=sheet_idx,
            mapping=mapping,
            mode=mode,
            amount_target=amount_target,
            target_url=target_url,
            force_restart=force_restart,
            result_file=result_file or "نتائج فحص زين.xlsx",
        )
        if success:
            self.send_json(200, {"status": "ok", "message": "Session started successfully"})
        else:
            self.send_json(400, {"status": "error", "message": "Failed to start session. Already running or invalid file."})

    def handle_queue_add(self, payload: dict[str, Any]) -> None:
        wb_name = payload.get("workbook", "2.xlsx")
        wb_path = self.project_root / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        sheet_idx = int(payload.get("sheet_index", 0))
        sheet_name = str(payload.get("sheet_name", f"Sheet {sheet_idx+1}"))
        mode = str(payload.get("mode", "smart_hybrid"))
        amount_target = str(payload.get("amount_target", "contract"))
        col_mapping = payload.get("column_mapping") or {}
        total_records = int(payload.get("total_records", 0))
        target_url = str(payload.get("target_url", "https://business.zain.sa/dashboard/quick-pay"))
        force_clean = bool(payload.get("force_clean", False) or payload.get("force_restart", False))

        checked = self._validate_input_sheet(wb_path, sheet_idx, col_mapping, mode, amount_target)
        if not checked:
            return
        col_mapping, sheet_name, total_records = checked
        job = self.queue_service.add_job(
            file_path=wb_path,
            sheet_index=sheet_idx,
            sheet_name=sheet_name,
            mode=mode,
            amount_target=amount_target,
            target_url=target_url,
            column_mapping=col_mapping,
            total_records=total_records,
            force_clean=force_clean,
        )
        self.send_json(200, {"status": "ok", "job": job.to_dict(), "jobs": self.queue_service.get_jobs()})

    def _validate_input_sheet(self, path, index, mapping, mode, amount_target):
        from openpyxl import load_workbook
        from domain.workbook_validation import validate_sheet
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                if index < 0 or index >= len(wb.worksheets):
                    raise ValueError("Unknown worksheet")
                sheet = wb.worksheets[index]
                has_headers = mapping.get("has_headers", True)
                if not isinstance(has_headers, bool):
                    self.send_json(400, {"status":"error", "message":"حدد هل يحتوي الملف على صف عناوين."})
                    return None
                merged = {**inspect_sheet_schema(sheet, has_headers=has_headers, count_rows=False)["letters"], **mapping}
                result = validate_sheet(sheet, merged, mode, amount_target)
                if not result["ok"]:
                    self.send_json(400, {"status":"error", "message":"راجع أعمدة الشيت: " + " · ".join(result["problems"]), "validation":result})
                    return None
                return merged, sheet.title, result["rows"]
            finally:
                wb.close()
        except Exception:
            self.send_json(400, {"status":"error", "message":"تعذر قراءة ورقة العمل. راجع الملف والورقة المحددة."})
            return None

    def handle_queue_remove(self, payload: dict[str, Any]) -> None:
        job_id = str(payload.get("job_id", ""))
        with self.orchestrator.lock:
            if self.orchestrator.is_running and self.orchestrator.current_job_id == job_id:
                self.send_json(409, {"status": "error", "message": "لا يمكن إخراج الشيت الجاري من الطابور أثناء فحصه."})
                return
            removed = self.queue_service.remove_job(job_id)
        if not removed:
            self.send_json(409, {"status": "error", "message": "الشيت غير موجود أو قيد المعالجة. حدّث الطابور."})
            return
        self.send_json(200, {"status": "ok", "removed": removed, "jobs": self.queue_service.get_jobs()})

    def handle_queue_start(self) -> None:
        with self.orchestrator.lock:
            jobs = self.queue_service.get_jobs()
            if self.orchestrator.is_running or any(job.get("status") == "active" for job in jobs):
                self.send_json(409, {"status":"error","message":"هناك جولة نشطة. تابع الحالة أو استأنف الفحص المتوقف مؤقتًا."})
                return
            pending = [job for job in jobs if job.get("status") == "pending"]
            if not pending:
                self.send_json(400, {"status":"error","message":"تحقق من ملف وأضفه إلى الجولة أولًا."})
                return
            for job in pending:
                if not self._validate_input_sheet(self.project_root/job["filename"], job["sheet_index"],
                    job.get("column_mapping",{}), job.get("mode","smart_hybrid"), job.get("amount_target","remaining")):
                    return
            success = self.orchestrator.start_queue()
        self.send_json(200 if success else 409, {"status":"ok" if success else "error",
            "message":"بدأت الجولة. تُفحص الملفات بالتتابع والنتائج بعد اكتمالها." if success else "تعذر البدء. حدّث الحالة والطابور."})

    def handle_telegram_status(self) -> None:
        integration = getattr(self.orchestrator, "telegram_integration", None)
        if integration:
            self.send_json(200, integration.status())
        else:
            self.send_json(200, {"status": "ok", "enabled": False, "configured": False, "connected": None})

    def handle_telegram_ping(self) -> None:
        if not getattr(self.orchestrator, "telegram_enabled", False):
            self.send_json(409, {"status": "error", "message": "تكامل تلقرام غير مفعّل"})
            return
        try:
            integration = getattr(self.orchestrator, "telegram_integration", None)
            if integration is None:
                self.send_json(409, {"status": "error", "message": "أعد تشغيل البرنامج لتفعيل أدوات التلقرام"})
                return
            accepted = integration.notifier.send_message("تشيك: هذه رسالة تجريبية من لوحة المتابعة.", wait=True)
            self.send_json(200 if accepted else 502, {"status": "ok" if accepted else "error",
                "message": "أكد تلقرام استلام الرسالة التجريبية" if accepted else "لم يؤكد تلقرام إرسال الرسالة؛ راجع إعدادات البوت والحساب"})
        except Exception:
            self.send_json(502, {"status": "error", "message": "تعذر إرسال الرسالة إلى تلقرام"})

    def handle_telegram_action(self, path: str, payload: dict) -> None:
        integration = getattr(self.orchestrator, "telegram_integration", None)
        if integration is None:
            self.send_json(409, {"status": "error", "message": "أعد تشغيل البرنامج لتفعيل أدوات التلقرام"})
            return
        if path.endswith("/toggle"):
            enabled = payload.get("enabled")
            if not isinstance(enabled, bool):
                self.send_json(400, {"status": "error", "message": "اختر حالة تفعيل صحيحة"})
                return
            try:
                integration.set_enabled(enabled)
                self.send_json(200, {**integration.status(), "message": "فُعّلت تنبيهات التلقرام" if enabled else "أُوقفت تنبيهات التلقرام"})
            except Exception:
                self.send_json(502, {"status": "error", "message": "تعذر تغيير حالة التلقرام"})
        elif path.endswith("/check-connection"):
            result = integration.check_connection()
            self.send_json(200 if result["connected"] else 502, {**result, "status": "ok" if result["connected"] else "error"})
        elif path.endswith("/send-results"):
            if not self.orchestrator.telegram_enabled:
                self.send_json(409, {"status": "error", "message": "فعّل التلقرام قبل إرسال النتائج"})
                return
            with self.orchestrator.lock:
                records = list(self.orchestrator.all_completed_records.values())
                target = self.orchestrator.result_path
            if not records:
                self.send_json(409, {"status": "error", "message": "لا توجد نتائج محفوظة لهذه الجلسة لإرسالها"})
                return
            try:
                from zain_checker.executive_reporter import export_executive_workbook
                export_executive_workbook(records, target)
                accepted = integration.notifier.send_excel_document(target, caption="نتائج الجلسة من لوحة تشيك", wait=True)
                self.send_json(200 if accepted else 502, {"status": "ok" if accepted else "error",
                    "message": "أكد تلقرام استلام ملف النتائج" if accepted else "لم يؤكد تلقرام إرسال الملف"})
            except Exception:
                self.send_json(502, {"status": "error", "message": "تعذر تجهيز أو إرسال ملف النتائج"})

    def handle_check_sheet_status(self, parsed: Any) -> None:
        from urllib.parse import parse_qs
        qs = parse_qs(parsed.query)
        wb_name = qs.get("workbook", [""])[0]
        try:
            sheet_idx = int(qs.get("sheet_index", [0])[0])
        except (ValueError, TypeError):
            sheet_idx = 0

        wb_path = self.project_root / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        wb_stem = wb_path.stem

        job = self.queue_service.get_job_by_file_and_sheet(wb_path.name, sheet_idx)
        chk_path = self.project_root / f".checkpoint_{wb_stem}_{sheet_idx}.json"
        has_checkpoint = chk_path.exists()
        completed_count = 0
        matches = 0
        mismatches = 0
        errors = 0
        total_records = 0
        sheet_name = f"ورقة {sheet_idx+1}"
        is_completed = False

        if job:
            completed_count = job.completed
            matches = job.matches
            mismatches = job.mismatches
            errors = job.errors
            total_records = job.total_records
            sheet_name = job.sheet_name
            is_completed = (job.status == "completed" or (total_records > 0 and completed_count >= total_records))

        if has_checkpoint and completed_count == 0:
            try:
                raw_chk = json.loads(chk_path.read_text(encoding="utf-8"))
                completed_count = len(raw_chk.get("completed_indices", []))
                mismatches = len(raw_chk.get("mismatches", []))
                matches = max(0, completed_count - mismatches)
                if total_records > 0 and completed_count >= total_records:
                    is_completed = True
            except Exception:
                pass

        self.send_json(200, {
            "status": "ok",
            "is_completed": is_completed,
            "has_progress": (completed_count > 0 or has_checkpoint),
            "completed_count": completed_count,
            "total_records": total_records,
            "matches": matches,
            "mismatches": mismatches,
            "errors": errors,
            "filename": wb_path.name,
            "sheet_name": sheet_name,
            "job_id": job.id if job else None,
        })

    def handle_queue_restart(self, payload: dict[str, Any]) -> None:
        job_id = str(payload.get("job_id", "")).strip()
        auto_start = bool(payload.get("auto_start", True))
        if not job_id:
            self.send_json(400, {"status": "error", "message": "Missing job_id"})
            return

        success = self.orchestrator.restart_queue_job(job_id, auto_start=auto_start)
        if success:
            self.send_json(200, {
                "status": "ok",
                "message": "تمت إعادة تعيين الشيت بنجاح وبدء الفحص من الصفر",
                "jobs": self.queue_service.get_jobs(),
            })
        else:
            self.send_json(400, {
                "status": "error",
                "message": "تعذر إعادة تعيين الشيت. تأكد من أن الشيت ليس قيد الفحص حالياً."
            })

    def handle_queue_restart_all(self, payload: dict[str, Any]) -> None:
        auto_start = bool(payload.get("auto_start", True))
        success = self.orchestrator.restart_all_queue_jobs(auto_start=auto_start)
        if success:
            self.send_json(200, {
                "status": "ok",
                "message": "تمت إعادة تعيين كافة الشيتات في الطابور بنجاح وبدء الفحص من الصفر",
                "jobs": self.queue_service.get_jobs(),
            })
        else:
            self.send_json(400, {
                "status": "error",
                "message": "تعذر إعادة تعيين الطابور. تأكد من عدم وجود جلسة نشطة حالياً."
            })

    def handle_detect_results_sheet(self, parsed: Any) -> None:
        from urllib.parse import parse_qs
        qs = parse_qs(parsed.query)
        wb_name = qs.get("workbook", [""])[0]
        sheet_idx = int(qs.get("sheet_index", [0])[0])

        if not wb_name:
            self.send_json(400, {"status": "error", "message": "Missing workbook name"})
            return

        wb_path = self.project_root / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        if not wb_path.exists():
            self.send_json(404, {"status": "error", "message": f"Workbook not found: {wb_name}"})
            return

        from openpyxl import load_workbook
        from domain.workbook import inspect_sheet_schema, extract_customer_records

        wb = load_workbook(wb_path, read_only=True, data_only=True)
        try:
            target_ws = None
            target_idx = sheet_idx
            if 0 <= sheet_idx < len(wb.worksheets):
                target_ws = wb.worksheets[sheet_idx]

            if not target_ws:
                target_ws = wb.worksheets[0]
                target_idx = 0

            schema = inspect_sheet_schema(target_ws)
            doc_type = schema.get("document_type", "")
            is_special = ("أخطاء" in doc_type or "نتائج" in doc_type or "أخطاء" in target_ws.title or "فروقات" in target_ws.title)

            if not is_special:
                self.send_json(200, {
                    "status": "ok",
                    "is_results_or_errors_sheet": False,
                    "sheet_name": target_ws.title,
                    "sheet_index": target_idx,
                })
                return

            indices = schema["indices"]
            custs = extract_customer_records(
                sheet=target_ws,
                lookup_col=indices.get("lookup_col") or 3,
                amount_col=indices.get("amount_col") or 4,
                amount_col_2=indices.get("amount_col_2"),
                service_col=indices.get("service_col"),
                customer_col=indices.get("customer_col"),
                notes_col=indices.get("notes_col"),
                source_row_col=indices.get("source_row_col"),
                main_status_col=indices.get("main_status_col"),
                sub_status_col=indices.get("sub_status_col"),
                record_type="mixed",
            )

            # Limit preview rows to 200 max for instant UI rendering
            preview_rows = []
            for c in custs[:200]:
                preview_rows.append({
                    "row": c.row_number,
                    "type": c.record_type,
                    "number": c.lookup_number,
                    "name": c.customer_name,
                    "expected_sar": c.expected_amount / 100.0,
                    "live_sar": None,
                    "diff_sar": None,
                    "status": "error",
                    "status_label": "يحتاج صيانة",
                    "issue": c.error_or_review_details or "بانتظار فحص العامل",
                    "main_status": c.main_status,
                    "sub_status": c.sub_status,
                })

            self.send_json(200, {
                "status": "ok",
                "is_results_or_errors_sheet": True,
                "sheet_name": target_ws.title,
                "sheet_index": target_idx,
                "document_type": doc_type,
                "total_records": len(custs),
                "error_count": len(custs),
                "rows": preview_rows,
            })
        finally:
            wb.close()

    def handle_start_repair_session(self, payload: dict[str, Any]) -> None:
        wb_name = payload.get("workbook", "")
        sheet_idx = int(payload.get("sheet_index", 0))
        if not wb_name:
            self.send_json(400, {"status": "error", "message": "Missing workbook name"})
            return

        wb_path = self.project_root / wb_name if not Path(wb_name).is_absolute() else Path(wb_name)
        if not wb_path.exists():
            self.send_json(404, {"status": "error", "message": f"Workbook not found: {wb_name}"})
            return

        target_url = payload.get("target_url", "https://business.zain.sa/dashboard/quick-pay")
        res = self.orchestrator.start_repair_session_from_sheet(
            workbook_path=wb_path,
            sheet_index=sheet_idx,
            target_url=target_url,
        )
        code = 200 if res.get("status") == "ok" else 400
        self.send_json(code, res)

