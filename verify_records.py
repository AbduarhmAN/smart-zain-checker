import math
import threading
import time
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from main import BridgeServer
from zain_checker.config import ACCOUNTS_WORKSHEET_INDEX, PROJECT_DIRECTORY
from zain_checker.console import log
from zain_checker.money import money_to_halalas
from zain_checker.workbook import Customer, find_source_workbook


TARGET_ROWS = (
    25144,
    25176,
    25183,
    25186,
    25245,
    25265,
    25278,
    25314,
    25318,
    25357,
)
VERIFY_WAIT_SECONDS = 5


@dataclass(frozen=True)
class VerificationResult:
    customer: Customer
    initial_raw: object
    final_raw: object
    final_amount: int | None
    status: str


def load_target_customers() -> list[Customer]:
    workbook_path = find_source_workbook(PROJECT_DIRECTORY)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)

    try:
        sheet = workbook.worksheets[ACCOUNTS_WORKSHEET_INDEX]
        target_rows = set(TARGET_ROWS)
        found: dict[int, Customer] = {}
        rows = sheet.iter_rows(
            min_row=min(TARGET_ROWS),
            max_row=max(TARGET_ROWS),
            min_col=12,
            max_col=13,
            values_only=True,
        )
        for row_number, (account_value, amount_value) in enumerate(
            rows,
            start=min(TARGET_ROWS),
        ):
            if row_number not in target_rows:
                continue

            account = str(account_value or "").strip()
            if not account.isdigit():
                raise RuntimeError(
                    f"Excel row {row_number} has an invalid account number."
                )
            found[row_number] = Customer(
                account,
                money_to_halalas(amount_value),
                row_number,
            )

        missing_rows = [row for row in TARGET_ROWS if row not in found]
        if missing_rows:
            raise RuntimeError(f"Excel rows were not found: {missing_rows}")
        return [found[row] for row in TARGET_ROWS]
    finally:
        workbook.close()


class VerificationState:
    def __init__(self, customers: list[Customer]) -> None:
        self.customers = customers
        self.index = 0
        self.phase = "initial"
        self.initial_raw: object = None
        self.ready_at = 0.0
        self.results: list[VerificationResult] = []
        self.error: str | None = None
        self.finished = threading.Event()
        self.lock = threading.Lock()

    def get_task(self) -> dict[str, Any]:
        with self.lock:
            if self.error:
                return {"status": "error", "message": self.error}
            if self.index == len(self.customers):
                return {"status": "complete", "total": len(self.customers)}

            remaining = self.ready_at - time.monotonic()
            if self.phase == "settling" and remaining > 0:
                return {
                    "status": "waiting",
                    "retry_after_seconds": max(1, math.ceil(remaining)),
                    "checked": self.index,
                    "total": len(self.customers),
                }

            customer = self.customers[self.index]
            return {
                "status": "check",
                "row_number": customer.row_number,
                "contract": customer.contract,
                "sequence": self.index + 1,
                "total": len(self.customers),
            }

    def submit_result(self, payload: dict[str, Any]) -> dict[str, Any]:
        with self.lock:
            if self.error:
                return {"status": "error", "message": self.error}
            if self.index == len(self.customers):
                return {"status": "complete", "total": len(self.customers)}

            customer = self.customers[self.index]
            if payload.get("row_number") != customer.row_number:
                raise ValueError("The browser returned an unexpected Excel row.")
            if str(payload.get("contract", "")) != customer.contract:
                raise ValueError("The browser returned an unexpected account.")

            status = str(payload.get("status") or "")
            if status != "ok":
                detail = payload.get("message") or payload.get("support_id") or status
                self.results.append(
                    VerificationResult(customer, None, detail, None, status)
                )
                log(
                    f"Verification row {customer.row_number}, account "
                    f"{customer.contract}: website status {status!r}, "
                    f"detail={detail!r}."
                )
                return self._advance()

            website_raw = payload.get("website_amount")
            if self.phase == "initial":
                self.initial_raw = website_raw
                self.phase = "settling"
                self.ready_at = time.monotonic() + VERIFY_WAIT_SECONDS
                log(
                    f"Verification row {customer.row_number}, account "
                    f"{customer.contract}: initial website value={website_raw!r}; "
                    f"waiting {VERIFY_WAIT_SECONDS} seconds on the same page."
                )
                return {
                    "status": "waiting",
                    "retry_after_seconds": VERIFY_WAIT_SECONDS,
                    "checked": self.index,
                    "total": len(self.customers),
                }

            try:
                final_amount = money_to_halalas(website_raw)
            except ValueError as error:
                self.results.append(
                    VerificationResult(
                        customer,
                        self.initial_raw,
                        website_raw,
                        None,
                        f"invalid amount: {error}",
                    )
                )
                log(
                    f"Verification row {customer.row_number}, account "
                    f"{customer.contract}: invalid final value={website_raw!r}."
                )
                return self._advance()

            comparison = (
                "match" if final_amount == customer.expected_amount else "mismatch"
            )
            self.results.append(
                VerificationResult(
                    customer,
                    self.initial_raw,
                    website_raw,
                    final_amount,
                    comparison,
                )
            )
            log(
                f"Verified row {customer.row_number}, account {customer.contract}: "
                f"M={customer.expected_amount / 100:,.2f}, "
                f"initial={self.initial_raw!r}, final={website_raw!r}, "
                f"parsed={final_amount / 100:,.2f}: {comparison}."
            )
            return self._advance()

    def _advance(self) -> dict[str, Any]:
        self.index += 1
        self.phase = "initial"
        self.initial_raw = None
        self.ready_at = 0.0
        if self.index == len(self.customers):
            self.finished.set()
            return {"status": "complete", "total": len(self.customers)}
        return {
            "status": "waiting",
            "retry_after_seconds": 1,
            "checked": self.index,
            "total": len(self.customers),
        }


def run() -> None:
    customers = load_target_customers()
    state = VerificationState(customers)
    server = BridgeServer(state)
    server_thread = threading.Thread(target=server.serve_forever, daemon=True)
    server_thread.start()

    log(
        f"Verification bridge ready for {len(customers)} accounts. "
        "No checkpoint or Excel output will be changed. Attach the extension "
        "to a Zain tab and select Start."
    )
    try:
        state.finished.wait()
    finally:
        server.shutdown()
        server.server_close()

    log("Verification summary:")
    for result in state.results:
        expected = result.customer.expected_amount / 100
        if result.final_amount is None:
            final_text = repr(result.final_raw)
        else:
            final_text = f"{result.final_amount / 100:,.2f}"
        log(
            f"Row {result.customer.row_number}, account "
            f"{result.customer.contract}: M={expected:,.2f}, "
            f"website={final_text}, result={result.status}."
        )


if __name__ == "__main__":
    run()
